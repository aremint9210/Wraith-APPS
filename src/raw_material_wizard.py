import csv
import os
import shutil
import tkinter as tk
from tkinter import filedialog, messagebox, simpledialog
import customtkinter as ctk

try:
    from PIL import Image

    HAS_PIL = True
except ImportError:
    HAS_PIL = False

try:
    import openpyxl

    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

COLOR_CANVAS = "#080A0F"
COLOR_CARD = "#111520"
COLOR_CARD_HOVER = "#181E2E"
COLOR_INNER = "#0B0E14"
COLOR_BORDER = "#1E2638"
COLOR_CYAN = "#38BDF8"
COLOR_CYAN_HOVER = "#0284C7"
COLOR_AMBER = "#D97706"
COLOR_ROSE = "#F43F5E"
COLOR_EMERALD = "#10B981"
COLOR_TEXT_MAIN = "#F8FAFC"
COLOR_TEXT_MUTED = "#64748B"
COLOR_DISABLED_BG = "#1E293B"

FONT_FAMILY = "Segoe UI"
FONT_MONO = "Consolas"


class LocationPairRow(ctk.CTkFrame):
    def __init__(self, master, index, pair_info, on_row_select_cb, on_pick_direct_cb):
        super().__init__(
            master,
            fg_color=COLOR_CARD,
            corner_radius=12,
            border_width=1,
            border_color=COLOR_BORDER
        )
        self.index = index
        self.pair_info = pair_info
        self.on_row_select_cb = on_row_select_cb
        self.on_pick_direct_cb = on_pick_direct_cb

        self.vis_assigned = None
        self.ir_assigned = None
        self.is_selected = False

        self._create_ui()

    def _create_ui(self):
        top_bar = ctk.CTkFrame(self, fg_color="transparent")
        top_bar.pack(fill="x", padx=16, pady=(10, 4))

        self.idx_badge = ctk.CTkLabel(
            top_bar,
            text=f" #{self.index + 1:02d} ",
            font=ctk.CTkFont(family=FONT_FAMILY, size=11, weight="bold"),
            text_color="#000000",
            fg_color=COLOR_CYAN,
            corner_radius=6,
            padx=6, pady=2
        )
        self.idx_badge.pack(side="left", padx=(0, 10))

        self.loc_lbl = ctk.CTkLabel(
            top_bar,
            text=self.pair_info["name"],
            font=ctk.CTkFont(family=FONT_FAMILY, size=14, weight="bold"),
            text_color=COLOR_TEXT_MAIN,
            anchor="w"
        )
        self.loc_lbl.pack(side="left", fill="x", expand=True)

        pair_frame = ctk.CTkFrame(self, fg_color="transparent")
        pair_frame.pack(fill="x", padx=12, pady=(0, 10))
        pair_frame.columnconfigure(0, weight=1)
        pair_frame.columnconfigure(1, weight=1)

        # 1. VISUAL SLOT
        self.vis_card = ctk.CTkFrame(pair_frame, fg_color=COLOR_INNER, corner_radius=8, border_width=1,
                                     border_color=COLOR_BORDER)
        self.vis_card.grid(row=0, column=0, padx=4, pady=4, sticky="nsew")

        ctk.CTkLabel(
            self.vis_card, text="📷 VISUAL PHOTO",
            font=ctk.CTkFont(family=FONT_FAMILY, size=10, weight="bold"), text_color=COLOR_CYAN
        ).pack(anchor="w", padx=10, pady=(6, 2))

        self.vis_status_lbl = ctk.CTkLabel(
            self.vis_card, text="No File Linked",
            font=ctk.CTkFont(family=FONT_FAMILY, size=10), text_color=COLOR_TEXT_MUTED,
            anchor="w"
        )
        self.vis_status_lbl.pack(anchor="w", padx=10, pady=(0, 6))

        vis_btn_box = ctk.CTkFrame(self.vis_card, fg_color="transparent")
        vis_btn_box.pack(fill="x", padx=10, pady=(0, 6))

        self.vis_select_btn = ctk.CTkButton(
            vis_btn_box, text="Pick...", height=24, fg_color=COLOR_CYAN, text_color="#000000",
            hover_color=COLOR_CYAN_HOVER, font=ctk.CTkFont(family=FONT_FAMILY, size=11, weight="bold"),
            command=lambda: self.on_pick_direct_cb(self, "visual")
        )
        self.vis_select_btn.pack(side="left", padx=(0, 4))

        ctk.CTkButton(
            vis_btn_box, text="Skip", height=24, fg_color="transparent", border_width=1,
            border_color=COLOR_ROSE, text_color=COLOR_ROSE, width=38, command=lambda: self.set_status("visual", "NONE")
        ).pack(side="left", padx=2)

        ctk.CTkButton(
            vis_btn_box, text="KIV", height=24, fg_color="transparent", border_width=1,
            border_color=COLOR_AMBER, text_color=COLOR_AMBER, width=38, command=lambda: self.set_status("visual", "KIV")
        ).pack(side="left", padx=2)

        # 2. THERMAL IR SLOT
        self.ir_card = ctk.CTkFrame(pair_frame, fg_color=COLOR_INNER, corner_radius=8, border_width=1,
                                    border_color=COLOR_BORDER)
        self.ir_card.grid(row=0, column=1, padx=4, pady=4, sticky="nsew")

        ctk.CTkLabel(
            self.ir_card, text="🔥 THERMAL IR PHOTO",
            font=ctk.CTkFont(family=FONT_FAMILY, size=10, weight="bold"), text_color=COLOR_AMBER
        ).pack(anchor="w", padx=10, pady=(6, 2))

        self.ir_status_lbl = ctk.CTkLabel(
            self.ir_card, text="No File Linked",
            font=ctk.CTkFont(family=FONT_FAMILY, size=10), text_color=COLOR_TEXT_MUTED,
            anchor="w"
        )
        self.ir_status_lbl.pack(anchor="w", padx=10, pady=(0, 6))

        ir_btn_box = ctk.CTkFrame(self.ir_card, fg_color="transparent")
        ir_btn_box.pack(fill="x", padx=10, pady=(0, 6))

        self.ir_select_btn = ctk.CTkButton(
            ir_btn_box, text="Pick...", height=24, fg_color=COLOR_AMBER, text_color="#000000",
            hover_color="#B45309", font=ctk.CTkFont(family=FONT_FAMILY, size=11, weight="bold"),
            command=lambda: self.on_pick_direct_cb(self, "ir")
        )
        self.ir_select_btn.pack(side="left", padx=(0, 4))

        ctk.CTkButton(
            ir_btn_box, text="Skip", height=24, fg_color="transparent", border_width=1,
            border_color=COLOR_ROSE, text_color=COLOR_ROSE, width=38, command=lambda: self.set_status("ir", "NONE")
        ).pack(side="left", padx=2)

        ctk.CTkButton(
            ir_btn_box, text="KIV", height=24, fg_color="transparent", border_width=1,
            border_color=COLOR_AMBER, text_color=COLOR_AMBER, width=38, command=lambda: self.set_status("ir", "KIV")
        ).pack(side="left", padx=2)

        for widget in [self, top_bar, self.loc_lbl]:
            widget.bind("<Button-1>", lambda e: self.on_row_select_cb(self))

    def set_selected(self, state: bool):
        self.is_selected = state
        if state:
            self.configure(border_color=COLOR_CYAN, border_width=2, fg_color="#182338")
        else:
            self.configure(border_color=COLOR_BORDER, border_width=1, fg_color=COLOR_CARD)

    def set_status(self, kind, value):
        if kind == "visual":
            self.vis_assigned = value
            if os.path.isabs(value):
                self.vis_status_lbl.configure(text=f"✓ {os.path.basename(value)}", text_color=COLOR_CYAN)
                self.vis_card.configure(border_color=COLOR_CYAN)
                self.vis_select_btn.configure(text="✓ Linked", fg_color=COLOR_DISABLED_BG, text_color=COLOR_TEXT_MUTED)
            else:
                self.vis_status_lbl.configure(text=f"Status: {value}",
                                              text_color=COLOR_ROSE if value == "NONE" else COLOR_AMBER)
                self.vis_card.configure(border_color=COLOR_BORDER)
                self.vis_select_btn.configure(text="Pick...", fg_color=COLOR_CYAN, text_color="#000000")
        else:
            self.ir_assigned = value
            if os.path.isabs(value):
                self.ir_status_lbl.configure(text=f"✓ {os.path.basename(value)}", text_color=COLOR_AMBER)
                self.ir_card.configure(border_color=COLOR_AMBER)
                self.ir_select_btn.configure(text="✓ Linked", fg_color=COLOR_DISABLED_BG, text_color=COLOR_TEXT_MUTED)
            else:
                self.ir_status_lbl.configure(text=f"Status: {value}",
                                             text_color=COLOR_ROSE if value == "NONE" else COLOR_AMBER)
                self.ir_card.configure(border_color=COLOR_BORDER)
                self.ir_select_btn.configure(text="Pick...", fg_color=COLOR_AMBER, text_color="#000000")


class RawMaterialDeclarationWindow(ctk.CTkToplevel):
    def __init__(self, parent, substation_dir, mapping_dir, csv_file_name, master_list_path, output_img_dir):
        super().__init__(parent)
        self.title("RMS // Raw Material Declaration & Preview Studio")
        self.geometry("1400x900")
        self.minsize(1000, 650)
        self.resizable(True, True)
        self.configure(fg_color=COLOR_CANVAS)

        self.parent = parent
        self.substation_dir = substation_dir
        self.mapping_dir = mapping_dir
        self.csv_file_name = csv_file_name
        self.master_list_path = master_list_path
        self.output_img_dir = output_img_dir

        self.active_raw_dir = self._auto_detect_raw_dir()
        self.staging_dir = os.path.join(self.active_raw_dir, ".staged_raw_images")
        os.makedirs(self.staging_dir, exist_ok=True)
        self.staged_map = {}

        self.location_pairs = self._parse_paired_locations()
        self.rows = []
        self.active_selected_row = None
        self.current_preview_image = None

        self.protocol("WM_DELETE_WINDOW", self._on_close)
        self.transient(parent)
        self.grab_set()
        self._create_ui()
        self._refresh_file_browser()

    def _auto_detect_raw_dir(self):
        stem = os.path.splitext(self.csv_file_name)[0]
        m = re.match(r'^0*(\d+[a-zA-Z]?)', stem)
        clean_num = m.group(1).upper() if m else stem.upper()
        digits = "".join(c for c in stem if c.isdigit())
        clean_digits = digits.lstrip('0') or digits

        candidates = [
            os.path.join(self.substation_dir, "03. RAW MATERIAL"),
            os.path.join(self.substation_dir, "RAW_IMAGES"),
            self.substation_dir
        ]
        for base_path in candidates:
            if os.path.exists(base_path):
                for folder in os.listdir(base_path):
                    folder_path = os.path.join(base_path, folder)
                    if os.path.isdir(folder_path):
                        m_d = re.match(r'^0*(\d+[a-zA-Z]?)[\.\s\-_]', folder)
                        if m_d and m_d.group(1).upper() == clean_num:
                            return folder_path
                        if folder.startswith(f"{clean_num}.") or folder.startswith(f"{clean_digits}."):
                            return folder_path
                return base_path
        return self.substation_dir

    def stage_file(self, orig_path):
        if not os.path.exists(orig_path):
            return None
        file_name = os.path.basename(orig_path)
        staged_path = os.path.join(self.staging_dir, file_name)
        try:
            shutil.move(orig_path, staged_path)
            self.staged_map[staged_path] = orig_path
            return staged_path
        except Exception as e:
            messagebox.showerror("File Error", f"Could not temporarily move file:\n{e}")
            return None

    def unstage_file(self, staged_path):
        if staged_path in self.staged_map:
            orig_path = self.staged_map[staged_path]
            if os.path.exists(staged_path):
                try:
                    shutil.move(staged_path, orig_path)
                except Exception as e:
                    print(f"Error restoring file: {e}")
            del self.staged_map[staged_path]

    def _restore_all_staged_files(self):
        staged_items = list(self.staged_map.keys())
        for st in staged_items:
            self.unstage_file(st)
        if os.path.exists(self.staging_dir):
            try:
                os.rmdir(self.staging_dir)
            except Exception:
                pass

    def _detect_switchgear_type(self):
        csv_path = os.path.join(self.mapping_dir, self.csv_file_name)
        if os.path.exists(csv_path):
            try:
                with open(csv_path, mode="r", encoding="utf-8-sig") as f:
                    content = f.read().lower()
                    if "rmu" in content:
                        return "RMU"
                    elif "vcb" in content:
                        return "VCB"
            except Exception:
                pass
        return "RMU" if "rmu" in self.csv_file_name.lower() else "VCB"

    def _parse_paired_locations(self):
        raw_items = []
        sw_type = self._detect_switchgear_type()
        target_sheet = f"Image_Rename_{sw_type}"

        if HAS_OPENPYXL and os.path.exists(self.master_list_path):
            try:
                wb = openpyxl.load_workbook(self.master_list_path, data_only=True)
                sheet_to_use = None
                for sheet in wb.sheetnames:
                    if sheet.strip().lower() == target_sheet.lower():
                        sheet_to_use = sheet
                        break
                if not sheet_to_use:
                    for sheet in wb.sheetnames:
                        if sw_type.lower() in sheet.lower() and "rename" in sheet.lower():
                            sheet_to_use = sheet
                            break
                if sheet_to_use:
                    ws = wb[sheet_to_use]
                    for row in ws.iter_rows(min_row=2, values_only=True):
                        if row and len(row) >= 3:
                            col_b = str(row[1]).strip() if row[1] else ""
                            col_c = str(row[2]).strip() if row[2] else ""
                            col_d = str(row[3]).strip() if len(row) > 3 and row[3] else ""
                            full_loc = f"{col_b} {col_c}".strip()
                            if full_loc: raw_items.append((full_loc, col_d))
                wb.close()
            except Exception:
                pass

        if not raw_items:
            csv_path = os.path.join(self.mapping_dir, self.csv_file_name)
            if os.path.exists(csv_path):
                try:
                    with open(csv_path, mode="r", encoding="utf-8-sig") as f:
                        reader = csv.reader(f)
                        next(reader, None)
                        for row in reader:
                            if row and len(row) >= 3:
                                col_b = row[1].strip() if len(row) > 1 else ""
                                col_c = row[2].strip() if len(row) > 2 else ""
                                col_d = row[3].strip() if len(row) > 3 else ""
                                full_loc = f"{col_b} {col_c}".strip()
                                if full_loc: raw_items.append((full_loc, col_d))
                except Exception:
                    pass

        pairs = []
        i = 0
        while i < len(raw_items):
            vis_loc, _ = raw_items[i]
            ir_loc, _ = raw_items[i + 1] if (i + 1) < len(raw_items) else (vis_loc, "IR")
            clean_name = vis_loc.replace(" Visual", "").replace(" Normal", "").replace(" Thermal", "").replace(" IR",
                                                                                                               "").strip()
            pairs.append({"name": clean_name, "vis_target": vis_loc, "ir_target": ir_loc})
            i += 2

        return pairs if pairs else [
            {"name": "Default Panel 1", "vis_target": "Panel 1 Visual", "ir_target": "Panel 1 IR"}]

    def _create_ui(self):
        sw_type = self._detect_switchgear_type()

        header = ctk.CTkFrame(self, fg_color=COLOR_CARD, corner_radius=12, border_width=1, border_color=COLOR_BORDER)
        header.pack(fill="x", padx=20, pady=(16, 8))

        title_frame = ctk.CTkFrame(header, fg_color="transparent")
        title_frame.pack(fill="x", padx=16, pady=(12, 4))

        ctk.CTkLabel(
            title_frame, text=f"Raw Material Mapping Studio: {self.csv_file_name} [{sw_type}]",
            font=ctk.CTkFont(family=FONT_FAMILY, size=18, weight="bold"), text_color=COLOR_TEXT_MAIN
        ).pack(side="left")

        ctk.CTkButton(
            title_frame, text="⚙ Configure Count...", width=140, height=30,
            fg_color=COLOR_CYAN, hover_color=COLOR_CYAN_HOVER, text_color="#000000",
            font=ctk.CTkFont(family=FONT_FAMILY, size=12, weight="bold"), command=self._prompt_custom_counts
        ).pack(side="right")

        self.summary_lbl = ctk.CTkLabel(
            header, text="Total Pairs: 0 | Assigned: 0",
            font=ctk.CTkFont(family=FONT_FAMILY, size=12), text_color=COLOR_CYAN
        )
        self.summary_lbl.pack(anchor="w", padx=16, pady=(0, 12))

        split_body = ctk.CTkFrame(self, fg_color="transparent")
        split_body.pack(fill="both", expand=True, padx=20, pady=4)
        split_body.columnconfigure(0, weight=3)
        split_body.columnconfigure(1, weight=2)
        split_body.rowconfigure(0, weight=1)

        # LEFT COLUMN: Matrix List
        left_pane = ctk.CTkFrame(split_body, fg_color=COLOR_CARD, corner_radius=12, border_width=1,
                                 border_color=COLOR_BORDER)
        left_pane.grid(row=0, column=0, padx=(0, 10), sticky="nsew")

        ctk.CTkLabel(
            left_pane, text="1. SELECT TARGET SLOT FROM MATRIX",
            font=ctk.CTkFont(family=FONT_FAMILY, size=12, weight="bold"), text_color=COLOR_CYAN
        ).pack(anchor="w", padx=16, pady=(14, 6))

        self.matrix_scroll = ctk.CTkScrollableFrame(left_pane, fg_color=COLOR_INNER, corner_radius=10)
        self.matrix_scroll.pack(fill="both", expand=True, padx=14, pady=(0, 14))

        # RIGHT COLUMN: Image Browser & Live Preview Studio
        right_pane = ctk.CTkFrame(split_body, fg_color=COLOR_CARD, corner_radius=12, border_width=1,
                                  border_color=COLOR_BORDER)
        right_pane.grid(row=0, column=1, padx=(10, 0), sticky="nsew")

        ctk.CTkLabel(
            right_pane, text="2. IMAGE PREVIEW & FOLDER BROWSER",
            font=ctk.CTkFont(family=FONT_FAMILY, size=12, weight="bold"), text_color=COLOR_AMBER
        ).pack(anchor="w", padx=16, pady=(14, 4))

        self.target_indicator_lbl = ctk.CTkLabel(
            right_pane, text="Active Slot: None Selected",
            font=ctk.CTkFont(family=FONT_FAMILY, size=11, weight="bold"), text_color=COLOR_TEXT_MUTED
        )
        self.target_indicator_lbl.pack(anchor="w", padx=16, pady=(0, 8))

        self._rebuild_matrix_ui()

        self.preview_card = ctk.CTkFrame(right_pane, fg_color=COLOR_INNER, corner_radius=10, height=200, border_width=1,
                                         border_color=COLOR_BORDER)
        self.preview_card.pack(fill="x", padx=16, pady=(0, 10))
        self.preview_card.pack_propagate(False)

        self.preview_img_lbl = ctk.CTkLabel(self.preview_card, text="[ Select an image below to preview ]",
                                            text_color=COLOR_TEXT_MUTED)
        self.preview_img_lbl.pack(fill="both", expand=True)

        list_container = ctk.CTkFrame(right_pane, fg_color=COLOR_INNER, corner_radius=10, border_width=1,
                                      border_color=COLOR_BORDER)
        list_container.pack(fill="both", expand=True, padx=16, pady=(0, 10))

        scrollbar = tk.Scrollbar(list_container, orient="vertical")
        scrollbar.pack(side="right", fill="y", padx=(0, 4), pady=4)

        self.file_listbox = tk.Listbox(
            list_container, bg=COLOR_INNER, fg=COLOR_TEXT_MAIN,
            selectbackground=COLOR_CYAN, selectforeground="#000000",
            font=(FONT_MONO, 11), bd=0, highlightthickness=0,
            yscrollcommand=scrollbar.set, activestyle="none"
        )
        self.file_listbox.pack(side="left", fill="both", expand=True, padx=10, pady=10)
        scrollbar.config(command=self.file_listbox.yview)
        self.file_listbox.bind("<<ListboxSelect>>", self._on_file_selected_in_browser)

        assign_btn = ctk.CTkButton(
            right_pane, text="🔗 Assign Selected Image to Active Slot", height=40,
            fg_color=COLOR_EMERALD, hover_color="#059669", text_color="#000000",
            font=ctk.CTkFont(family=FONT_FAMILY, size=13, weight="bold"),
            command=self._assign_browser_image_to_active_slot
        )
        assign_btn.pack(fill="x", padx=16, pady=(0, 14))

        # Bottom Action Bar
        bottom_bar = ctk.CTkFrame(self, fg_color=COLOR_CARD, corner_radius=12, border_width=1,
                                  border_color=COLOR_BORDER)
        bottom_bar.pack(fill="x", padx=20, pady=(8, 16))

        b_left = ctk.CTkFrame(bottom_bar, fg_color="transparent")
        b_left.pack(side="left", padx=16, pady=12)

        ctk.CTkButton(
            b_left, text="Cancel", fg_color=COLOR_INNER, hover_color=COLOR_CARD_HOVER,
            text_color=COLOR_TEXT_MUTED, font=ctk.CTkFont(family=FONT_FAMILY, size=13),
            border_width=1, border_color=COLOR_BORDER, command=self._on_close
        ).pack(side="left", padx=(0, 10))

        self.missing_badge = ctk.CTkLabel(
            b_left, text="Pending: 0 slots", font=ctk.CTkFont(family=FONT_FAMILY, size=12, weight="bold"),
            text_color="#000000", fg_color=COLOR_AMBER, corner_radius=6, padx=10, pady=4
        )
        self.missing_badge.pack(side="left")

        finish_btn = ctk.CTkButton(
            bottom_bar, text="Finish & Process Renaming  ✓", height=40, width=220,
            fg_color=COLOR_CYAN, hover_color=COLOR_CYAN_HOVER, text_color="#000000",
            font=ctk.CTkFont(family=FONT_FAMILY, size=14, weight="bold"), command=self._finish_and_copy
        )
        finish_btn.pack(side="right", padx=16, pady=12)

        self._update_summary()

    def _rebuild_matrix_ui(self):
        for widget in self.matrix_scroll.winfo_children():
            widget.destroy()

        self.rows = []
        for idx, pair in enumerate(self.location_pairs):
            row_widget = LocationPairRow(
                self.matrix_scroll, idx, pair, self._on_row_clicked, self._on_pick_direct_clicked
            )
            row_widget.pack(fill="x", expand=True, padx=6, pady=6)
            self.rows.append(row_widget)

        if self.rows and hasattr(self, 'target_indicator_lbl'):
            self._on_row_clicked(self.rows[0])

    def _on_row_clicked(self, row_widget):
        for r in self.rows:
            r.set_selected(False)
        row_widget.set_selected(True)
        self.active_selected_row = row_widget

        if hasattr(self, 'target_indicator_lbl'):
            self.target_indicator_lbl.configure(
                text=f"Active Slot: #{row_widget.index + 1:02d} [{row_widget.pair_info['name']}]",
                text_color=COLOR_CYAN
            )

    def _on_pick_direct_clicked(self, row_widget, kind):
        self._on_row_clicked(row_widget)

        try:
            self.grab_release()
        except Exception:
            pass

        file_path = filedialog.askopenfilename(
            initialdir=self.active_raw_dir,
            title=f"Select Image for #{row_widget.index + 1} [{kind.upper()}]",
            filetypes=[("Image Files", "*.jpg *.jpeg *.png *.bmp")]
        )

        try:
            self.grab_set()
        except Exception:
            pass

        if file_path:
            self.active_raw_dir = os.path.dirname(file_path)
            staged_path = self.stage_file(file_path)
            if staged_path:
                old_val = row_widget.vis_assigned if kind == "visual" else row_widget.ir_assigned
                if old_val and os.path.isabs(old_val):
                    self.unstage_file(old_val)

                row_widget.set_status(kind, staged_path)
                self._refresh_file_browser()
                self._update_summary()

    def _refresh_file_browser(self):
        self.file_listbox.delete(0, tk.END)
        if os.path.exists(self.active_raw_dir):
            files = sorted([
                f for f in os.listdir(self.active_raw_dir)
                if f.lower().endswith((".jpg", ".jpeg", ".png", ".bmp"))
            ])
            for f in files:
                self.file_listbox.insert(tk.END, f"  📷  {f}")

    def _on_file_selected_in_browser(self, event):
        selection = self.file_listbox.curselection()
        if not selection:
            return
        line_text = self.file_listbox.get(selection[0]).strip()
        file_name = line_text.replace("📷", "").strip()
        full_path = os.path.join(self.active_raw_dir, file_name)

        if HAS_PIL and os.path.exists(full_path):
            try:
                pil_img = Image.open(full_path)
                # Use ctk.CTkImage to eliminate HighDPI PhotoImage warnings
                self.current_preview_image = ctk.CTkImage(light_image=pil_img, dark_image=pil_img, size=(240, 160))
                self.preview_img_lbl.configure(image=self.current_preview_image, text="")
            except Exception:
                self.preview_img_lbl.configure(image="", text="[ Preview Unavailable ]")
        else:
            self.preview_img_lbl.configure(image="", text=f"File: {file_name}")

    def _assign_browser_image_to_active_slot(self):
        if not self.active_selected_row:
            messagebox.showwarning("Select Slot", "Please click on a matrix row slot on the left first.")
            return

        selection = self.file_listbox.curselection()
        if not selection:
            messagebox.showwarning("Select Image", "Please select an image file from the right browser list.")
            return

        line_text = self.file_listbox.get(selection[0]).strip()
        file_name = line_text.replace("📷", "").strip()
        orig_path = os.path.join(self.active_raw_dir, file_name)

        staged_path = self.stage_file(orig_path)
        if staged_path:
            row = self.active_selected_row
            kind = "visual"
            if row.vis_assigned and not row.ir_assigned:
                kind = "ir"
            elif row.vis_assigned and row.ir_assigned:
                kind = "visual"

            old_val = row.vis_assigned if kind == "visual" else row.ir_assigned
            if old_val and os.path.isabs(old_val):
                self.unstage_file(old_val)

            row.set_status(kind, staged_path)
            self._refresh_file_browser()
            self._update_summary()

            curr_idx = self.rows.index(row)
            if curr_idx + 1 < len(self.rows):
                self._on_row_clicked(self.rows[curr_idx + 1])

    def _prompt_custom_counts(self):
        current_count = len(self.location_pairs)
        res = simpledialog.askinteger(
            "Configure Count", f"Current items: {current_count}\nEnter desired total location slots:",
            initialvalue=current_count, minvalue=1, maxvalue=200
        )
        if res and res != current_count:
            new_pairs = []
            for i in range(res):
                if i < len(self.location_pairs):
                    new_pairs.append(self.location_pairs[i])
                else:
                    new_pairs.append({"name": f"Item #{i + 1}", "vis_target": f"Item {i + 1} Visual",
                                      "ir_target": f"Item {i + 1} IR"})
            self.location_pairs = new_pairs
            self._rebuild_matrix_ui()
            self._update_summary()

    def _update_summary(self):
        total_slots = len(self.location_pairs) * 2
        assigned_slots = sum(1 for r in self.rows if r.vis_assigned and os.path.isabs(r.vis_assigned))
        assigned_slots += sum(1 for r in self.rows if r.ir_assigned and os.path.isabs(r.ir_assigned))

        missing_count = total_slots - assigned_slots
        self.summary_lbl.configure(
            text=f"Total Pairs: {len(self.location_pairs)} (Slots: {total_slots}) | Assigned: {assigned_slots}/{total_slots}")

        if missing_count == 0:
            self.missing_badge.configure(text="All Ready ✓", fg_color=COLOR_EMERALD)
        else:
            self.missing_badge.configure(text=f"Pending: {missing_count} slots", fg_color=COLOR_AMBER)

    def _finish_and_copy(self):
        os.makedirs(self.output_img_dir, exist_ok=True)
        copied_count = 0

        for row in self.rows:
            pair = row.pair_info
            for kind, target_key in [("visual", "vis_target"), ("ir", "ir_target")]:
                val = row.vis_assigned if kind == "visual" else row.ir_assigned
                if val and os.path.isabs(val) and os.path.exists(val):
                    ext = os.path.splitext(val)[1]
                    clean_name = "".join(c for c in pair[target_key] if c.isalnum() or c in (" ", "_", "-")).strip()
                    dest = os.path.join(self.output_img_dir, f"{clean_name}{ext}")
                    try:
                        shutil.copy2(val, dest)
                        copied_count += 1
                    except Exception as e:
                        print(f"Error copying: {e}")

        self._restore_all_staged_files()
        messagebox.showinfo("Declaration Complete",
                            f"Successfully processed raw material images!\nCopied: {copied_count} file(s).")
        self.destroy()

    def _on_close(self):
        self._restore_all_staged_files()
        self.destroy()


if __name__ == "__main__":
    root = ctk.CTk()
    root.withdraw()
    app = RawMaterialDeclarationWindow(
        root,
        substation_dir=r"C:\Users\aremi\Desktop\REPORT CBM",
        mapping_dir=r"C:\Users\aremi\Desktop\REPORT CBM\SCRIPT",
        csv_file_name="175.csv",
        master_list_path="",
        output_img_dir=r"C:\Users\aremi\Desktop\REPORT CBM\RENAMED_IMAGES"
    )
    app.mainloop()