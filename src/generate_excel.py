import os
import re
import sys
import time
import shutil
import warnings
import openpyxl
import subprocess
try:
    import fitz
except ImportError:
    fitz = None
from copy import copy
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime

# Suppress openpyxl's data validation user warning
warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')

# Reconfigure stdout for instant real-time line buffering
if sys.platform.startswith("win"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace", line_buffering=True)

# ── Paths ──────────────────────────────────────────────────────────────────────
DEFAULT_BASE = os.getcwd()
BASE = os.environ.get("BASE_DIR", DEFAULT_BASE)
FR_DIR = os.path.join(BASE, "01. FULL REPORT")
QR_DIR = os.path.join(BASE, "02. QUICK REPORT")
RM_DIR = os.path.join(BASE, "03. RAW MATERIAL")
XLSX_DIR = os.path.join(BASE, "04. TESTSHEET")
OUTPUT_DIR = os.path.join(BASE, "07. OUTPUT FILE")
OUT = os.path.join(BASE, "SUBSTATION_CHECKLIST.xlsx")
MASTER_LIST_PATH = os.path.join(BASE, "06. EXCELL MAPPING", "Master List.xlsx")


def create_timestamped_backup(file_path):
    if not os.path.exists(file_path):
        return
    try:
        backup_dir = os.path.join(os.path.dirname(file_path), "_backups")
        os.makedirs(backup_dir, exist_ok=True)
        base_name, ext = os.path.splitext(os.path.basename(file_path))
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup_file = os.path.join(backup_dir, f"{base_name}_{timestamp}{ext}")
        shutil.copy2(file_path, backup_file)
        print(f"  [BACKUP] Saved safety snapshot: {os.path.basename(backup_file)}", flush=True)
        backups = sorted([os.path.join(backup_dir, f) for f in os.listdir(backup_dir) if f.startswith(base_name)], key=os.path.getmtime)
        while len(backups) > 10:
            os.remove(backups.pop(0))
    except Exception as e:
        print(f"  [BACKUP WARNING] Could not create backup: {e}", flush=True)


def parse_substation_filter(input_str):
    if not input_str:
        return None, "ALL"
    s = str(input_str).strip()
    if s.upper() in ["ALL", "*", ""]:
        return None, "ALL"

    parts = [p.strip() for p in s.split(",") if p.strip()]
    ranges = []
    singles = set()

    for part in parts:
        part_upper = part.upper()
        if part_upper in ["ALL", "*"]:
            return None, "ALL"
        m_range = re.match(r"^(\d+)\s*-\s*(\d+)$", part)
        if m_range:
            low, high = int(m_range.group(1)), int(m_range.group(2))
            if low > high:
                low, high = high, low
            ranges.append((low, high))
            continue
        m_plus = re.match(r"^(?:>=|>)?\s*(\d+)\+?$", part)
        if "+" in part or ">=" in part:
            low = int(m_plus.group(1))
            ranges.append((low, 999999))
            continue
        m_single = re.match(r"^(?:PE\s*)?(\d+)([A-Z]?)$", part, re.IGNORECASE)
        if m_single:
            singles.add(int(m_single.group(1)))
            continue

    if not ranges and not singles:
        return None, "ALL"

    def matches(num):
        m = re.match(r"^\d+", str(num).strip())
        if not m:
            return False
        n = int(m.group(0))
        if n in singles:
            return True
        for low, high in ranges:
            if low <= n <= high:
                return True
        return False

    label_parts = []
    if singles:
        if len(singles) == 1:
            label_parts.append(f"PE {list(singles)[0]} only")
        else:
            label_parts.append(f"PE {', '.join(map(str, sorted(singles)))}")
    for low, high in ranges:
        if high == 999999:
            label_parts.append(f"PE {low} onwards")
        else:
            label_parts.append(f"PE {low}-{high}")

    return matches, " & ".join(label_parts)


KNOWN_SG_MANUFACTURERS = [
    ('INDKOM', ['INDKOM', 'INDROM', 'INDKO']),
    ('EPE', ['EPE', 'E.P.E']),
    ('TAMCO', ['TAMCO', 'TAMC']),
    ('LUCY ELECTRIC', ['LUCY', 'LUCY ELECTRIC']),
    ('SCHNEIDER', ['SCHNEIDER', 'MERLIN GERIN']),
    ('ABB', ['ABB']),
    ('SIEMENS', ['SIEMENS']),
    ('HEAG', ['HEAG']),
    ('SAREL', ['SAREL']),
    ('VOLTAMP', ['VOLTAMP']),
    ('TOSHIBA', ['TOSHIBA']),
    ('SSE', ['SSE']),
    ('SCHALTANLAGE', ['SCHALTANLAGE']),
    ('ITEM', ['ITEM']),
    ('ORMAZABAL', ['ORMAZABAL']),
]

PS_OCR_SCRIPT = '''param([string]$ImagePath)
[System.Reflection.Assembly]::LoadWithPartialName("System.Runtime.WindowsRuntime") | Out-Null
$null = [Windows.Media.Ocr.OcrEngine, Windows.Foundation.UniversalApiContract, ContentType=WindowsRuntime]
$null = [Windows.Graphics.Imaging.BitmapDecoder, Windows.Foundation.UniversalApiContract, ContentType=WindowsRuntime]
$null = [Windows.Storage.StorageFile, Windows.Foundation.UniversalApiContract, ContentType=WindowsRuntime]

function AwaitTask($winRtAsyncOp, $resultType) {
    $asTaskGen = [System.WindowsRuntimeSystemExtensions].GetMethods() | Where-Object { $_.Name -eq 'AsTask' -and $_.IsGenericMethod -and $_.GetParameters().Count -eq 1 } | Select-Object -First 1
    $asTask = $asTaskGen.MakeGenericMethod($resultType)
    $task = $asTask.Invoke($null, @($winRtAsyncOp))
    $task.Wait()
    return $task.Result
}

try {
    $file = AwaitTask ([Windows.Storage.StorageFile]::GetFileFromPathAsync($ImagePath)) ([Windows.Storage.StorageFile])
    $stream = AwaitTask ($file.OpenAsync([Windows.Storage.FileAccessMode]::Read)) ([Windows.Storage.Streams.IRandomAccessStream])
    $decoder = AwaitTask ([Windows.Graphics.Imaging.BitmapDecoder]::CreateAsync($stream)) ([Windows.Graphics.Imaging.BitmapDecoder])
    $bitmap = AwaitTask ($decoder.GetSoftwareBitmapAsync()) ([Windows.Graphics.Imaging.SoftwareBitmap])
    $engine = [Windows.Media.Ocr.OcrEngine]::TryCreateFromUserProfileLanguages()
    $ocrResult = AwaitTask ($engine.RecognizeAsync($bitmap)) ([Windows.Media.Ocr.OcrResult])
    Write-Output $ocrResult.Text
} catch {}
'''

def run_windows_ocr(image_path):
    ps_file = os.path.join(os.path.dirname(__file__), "temp_ocr.ps1")
    try:
        with open(ps_file, "w", encoding="utf-8") as f:
            f.write(PS_OCR_SCRIPT)
        cmd = ['powershell', '-ExecutionPolicy', 'Bypass', '-File', ps_file, '-ImagePath', image_path]
        res = subprocess.run(cmd, capture_output=True, encoding='utf-8', errors='ignore')
        return res.stdout.strip() if res and res.stdout else ""
    except Exception:
        return ""

def detect_sg_manufacturer_from_dg(substation_num, rm_dir):
    if not rm_dir or not os.path.exists(rm_dir):
        return ""

    val_str = str(substation_num).strip()
    m = re.match(r'^0*(\d+)([a-zA-Z]*)', val_str)
    clean_num = (m.group(1) + m.group(2).upper()) if m else val_str
    base_digits = m.group(1) if m else val_str

    matched = []
    for d in os.listdir(rm_dir):
        if not os.path.isdir(os.path.join(rm_dir, d)):
            continue
        m_d = re.match(r'^0*(\d+[a-zA-Z]?)[\.\s\-_]', d)
        if m_d and m_d.group(1).upper() == clean_num:
            matched.append(d)
        elif (d == val_str or d == clean_num or 
              d.startswith(f"{clean_num}.") or d.startswith(f"{clean_num} ") or 
              d.startswith(f"{val_str}.") or d.startswith(f"{val_str} ") or
              d.startswith(f"{base_digits}.") or d.startswith(f"{base_digits} ")):
            matched.append(d)

    if not matched:
        return ""

    sub_folder = os.path.join(rm_dir, matched[0])
    dg_folder = os.path.join(sub_folder, "DG")
    if not os.path.exists(dg_folder):
        return ""

    img_files = [f for f in os.listdir(dg_folder) if f.lower().endswith(('.jpg', '.jpeg', '.png'))]
    for img in img_files:
        img_path = os.path.join(dg_folder, img)
        text = run_windows_ocr(img_path)
        if text:
            text_upper = text.upper()
            for std_name, keywords in KNOWN_SG_MANUFACTURERS:
                if any(kw in text_upper for kw in keywords):
                    print(f"  [OCR AUTO-DETECT] Substation {substation_num}: Detected '{std_name}' from photo '{img}' in DG folder!")
                    return std_name

    return ""


# ── Helper: extract leading number ────────────────────────────────────────────
def extract_num(name):
    m = re.match(r'^(\d+)', name)
    return int(m.group(1)) if m else None


# ── Helper: validate substation name ──────────────────────────────────────────
def is_valid_substation_name(val):
    if not val:
        return False
    s = str(val).strip()
    if not s or s.isdigit():
        return False
    if s.lower() in ['none', 'null', 'nan', '-', 'n/a', '.', '/']:
        return False
    if re.match(r'^(?:substation\s*)?\d+[a-z]?$', s, re.IGNORECASE):
        return False
    return True


# ── Helper: clean substation name ──────────────────────────────────────────────
def clean_substation_name(name):
    name = re.sub(r'^\d+[a-zA-Z]?[\.\s\-]+', '', name)
    name = re.sub(r'^\d{8}\s*', '', name)
    name = re.sub(r'\s*\(.*?\)', '', name)
    name = re.sub(r'\.(xlsx|docx|doc|pdf)$', '', name, flags=re.IGNORECASE)
    cleaned = name.strip()
    return cleaned if is_valid_substation_name(cleaned) else ""



def scan_files(folder):
    result = set()
    if not os.path.exists(folder):
        return result
    for root, dirs, files in os.walk(folder):
        for f in files:
            num = extract_num(f)
            if num is not None:
                result.add(num)
    return result


def scan_full_report(folder):
    done_substations = set()
    if not os.path.exists(folder):
        return done_substations

    for root, dirs, files in os.walk(folder):
        rel = os.path.relpath(root, folder)
        if "NOT DONE" in rel.upper():
            continue

        prefix_files = {}
        for f in files:
            if f.startswith("~$"):
                continue
            m = re.match(r'^0*(\d+)', f)
            if m:
                num = int(m.group(1))
                prefix_files.setdefault(num, []).append(f)

        for num, f_list in prefix_files.items():
            has_g0_docx = any("group 0.docx" in f.lower() for f in f_list)
            has_g0_pdf  = any("group 0.pdf" in f.lower() for f in f_list)
            has_g1_docx = any("group 1.docx" in f.lower() for f in f_list)
            has_g1_pdf  = any("group 1.pdf" in f.lower() for f in f_list)
            has_main_pdf = any(f.lower().endswith(".pdf") and "group 0" not in f.lower() and "group 1" not in f.lower() for f in f_list)
            if (has_g0_docx and has_g0_pdf and has_g1_docx and has_g1_pdf and has_main_pdf) or has_main_pdf:
                done_substations.add(num)

    return done_substations


# ── Helper: clean empty-like values ───────────────────────────────────────────
def clean_val(val):
    if val is None:
        return ''
    val = str(val).strip()
    if val in ['-', 'N/A', '', '.', '/'] or val.lower() == 'n/a':
        return ''
    return val


def clean_sn_val(val):
    if val is None:
        return ''
    s = str(val).strip()
    if s.endswith('.0'):
        s = s[:-2]
    try:
        f = float(s)
        if f.is_integer():
            s = str(int(f))
    except ValueError:
        pass
    if s in ['-', 'N/A', '', '.', '/'] or s.lower() in ['n/a', 'nan', 'none', 'null']:
        return ''
    return s


def format_heater_amp(val):
    if val is None:
        return ''
    s = str(val).strip()
    if not s or s in ['-', 'N/A', '', '.', '/'] or s.lower() in ['n/a', 'nan', 'none', 'null']:
        return ''
    m = re.search(r'[-+]?\d*\.?\d+', s)
    if m:
        try:
            f = float(m.group(0))
            return f"{f:.2f}A"
        except ValueError:
            pass
    return s


def format_temp(val):
    if val is None:
        return ''
    s = str(val).strip()
    if not s or s in ['-', 'N/A', '.', '/']:
        return ''
    try:
        f = float(s)
        return f"{f:.1f}"
    except ValueError:
        return s


def normalize_date(date_val):
    if not date_val:
        return ''
    if hasattr(date_val, 'strftime'):
        return date_val.strftime('%d/%m/%Y')
    s = str(date_val).strip()
    m_iso = re.match(r'^(\d{4})-(\d{2})-(\d{2})', s)
    if m_iso:
        return f"{m_iso.group(3)}/{m_iso.group(2)}/{m_iso.group(1)}"
    m = re.match(r'^(\d{1,2})[-/](\d{1,2})[-/](\d{2,4})$', s)
    if m:
        d, mth, yr = m.group(1), m.group(2), m.group(3)
        if len(yr) == 2:
            yr = f"20{yr}"
        return f"{int(d):02d}/{int(mth):02d}/{yr}"
    return s


def format_time_val(time_raw):
    if not time_raw:
        return ''
    if hasattr(time_raw, 'strftime'):
        return time_raw.strftime('%I:%M %p').lstrip('0')
    time_str = str(time_raw).replace('.0', '').strip()
    m_ampm = re.match(r'^(\d{1,2}):(\d{2})\s*(AM|PM)?$', time_str, re.IGNORECASE)
    if m_ampm:
        hh, mm, ampm = int(m_ampm.group(1)), m_ampm.group(2), m_ampm.group(3)
        if not ampm:
            ampm = 'AM' if hh < 12 else 'PM'
            hh = hh if hh <= 12 else hh - 12
            if hh == 0:
                hh = 12
        return f"{hh}:{mm} {ampm.upper()}"
    m_digits = re.match(r'^(\d{3,4})$', time_str)
    if m_digits:
        time_str = time_str.zfill(4)
        hh, mm = int(time_str[:2]), int(time_str[2:])
        period = "AM" if hh < 12 else "PM"
        hh12 = hh if hh <= 12 else hh - 12
        if hh12 == 0:
            hh12 = 12
        return f"{hh12:02d}:{mm:02d} {period}"
    return time_str


def parse_sg_mfr_model(text):
    if not text:
        return '', ''
    s = str(text).strip()
    if s in ['-', 'N/A', '', '.', '/'] or s.lower() in ['none', 'null', 'nan']:
        return '', ''
    s_upper = s.upper()
    for std_name, keywords in KNOWN_SG_MANUFACTURERS:
        for kw in keywords:
            if s_upper.startswith(kw):
                rest = s[len(kw):].strip()
                rest = re.sub(r'^[\-_:\s]+', '', rest).strip()
                return std_name, rest
    parts = s.split(None, 1)
    if len(parts) == 1:
        return parts[0], ''
    return parts[0], parts[1]


def determine_fp_lvdb_types(has_tx1, has_tx2, fp1_data, fp2_data):
    """
    Implements naming conventions for Feeder Pillar (FP) and LVDB based on Transformer (TX) setup:
    - 1 TX + 2 FP/LVDB: Use sequential numbering (FP 1/FP 2, LVDB 1/LVDB 2) or distinct types (FP/LVDB).
    - 2 TX + 1 FP/LVDB: Append TX association (FP TX1/FP TX2 or LVDB TX1/LVDB TX2).
    - 2 TX + 2 FP/LVDB: Match each FP/LVDB to its corresponding TX (FP TX1/FP TX2, LVDB TX1/LVDB TX2, or mixed LVDB TX1/FP TX2).
    """
    def _is_present(d):
        if not d:
            return False
        if isinstance(d, dict):
            mfr = d.get('manufacturer')
            sn = d.get('serial')
            rtg = d.get('rating')
            name = d.get('name')
            for v in [mfr, sn, rtg]:
                if v and str(v).strip() not in ['', '-', 'N/A', 'n/a', 'None', 'nan', '.', '/']:
                    return True
            if name and str(name).strip() not in ['', '-', 'N/A', 'n/a', 'None', 'nan', '.', '/', 'LVDB / FP 1', 'LVDB / FP 2']:
                return True
            return False
        s = str(d).strip()
        return s not in ['', '-', 'N/A', 'n/a', 'None', 'nan', '.', '/', 'LVDB / FP 1', 'LVDB / FP 2']

    def _get_kind(d):
        text = ''
        if isinstance(d, dict):
            parts = []
            for k in ['name', 'model', 'type', 'manufacturer']:
                val = d.get(k)
                if val and str(val).strip() not in ['LVDB / FP 1', 'LVDB / FP 2']:
                    parts.append(str(val))
            text = ' '.join(parts)
        else:
            text = str(d)
        text_upper = text.upper()
        if 'LVDB' in text_upper or 'DISTRIBUTION BOARD' in text_upper or 'LV BOARD' in text_upper or re.search(r'\bLV\s*DB\b', text_upper):
            return 'LVDB'
        return 'FP'

    has_fp1 = _is_present(fp1_data)
    has_fp2 = _is_present(fp2_data)

    if not has_fp1 and not has_fp2:
        return '', ''

    kind1 = _get_kind(fp1_data) if has_fp1 else ''
    kind2 = _get_kind(fp2_data) if has_fp2 else ''

    tx_count = (1 if has_tx1 else 0) + (1 if has_tx2 else 0)
    if tx_count == 0:
        tx_count = 1
        has_tx1 = True

    # 1. 1 TX + 2 FP/LVDB: Use sequential numbering (FP 1/FP 2, LVDB 1/LVDB 2) or distinct types (FP/LVDB).
    if tx_count == 1 and (has_fp1 and has_fp2):
        if kind1 == kind2:
            return f'{kind1} 1', f'{kind2} 2'
        else:
            return kind1, kind2

    # 2. 2 TX + 1 FP/LVDB: Append TX association (FP TX1/FP TX2 or LVDB TX1/LVDB TX2).
    if tx_count >= 2 and (has_fp1 ^ has_fp2):
        if has_fp1:
            raw_text = ' '.join([str(v) for v in fp1_data.values() if v]).upper() if isinstance(fp1_data, dict) else str(fp1_data).upper()
            assoc = 'TX2' if ('TX2' in raw_text or 'TX 2' in raw_text) else 'TX1'
            return f'{kind1} {assoc}', ''
        else:
            raw_text = ' '.join([str(v) for v in fp2_data.values() if v]).upper() if isinstance(fp2_data, dict) else str(fp2_data).upper()
            assoc = 'TX1' if ('TX1' in raw_text or 'TX 1' in raw_text) else 'TX2'
            return '', f'{kind2} {assoc}'

    # 3. 2 TX + 2 FP/LVDB: Match each FP/LVDB to its corresponding TX (FP TX1/FP TX2, LVDB TX1/LVDB TX2, or mixed LVDB TX1/FP TX2).
    if tx_count >= 2 and (has_fp1 and has_fp2):
        return f'{kind1} TX1', f'{kind2} TX2'

    # 4. 1 TX + 1 FP/LVDB: Single setup
    if tx_count == 1:
        if has_fp1:
            raw_text = ' '.join([str(v) for v in fp1_data.values() if v]).upper() if isinstance(fp1_data, dict) else str(fp1_data).upper()
            if '1' in raw_text and 'TX' not in raw_text:
                return f'{kind1} 1', ''
            elif 'TX' in raw_text:
                return f'{kind1} TX1', ''
            return kind1, ''
        else:
            raw_text = ' '.join([str(v) for v in fp2_data.values() if v]).upper() if isinstance(fp2_data, dict) else str(fp2_data).upper()
            if '2' in raw_text and 'TX' not in raw_text:
                return '', f'{kind2} 2'
            elif 'TX' in raw_text:
                return '', f'{kind2} TX2'
            return '', kind2

    t1 = f'{kind1} 1' if has_fp1 else ''
    t2 = f'{kind2} 2' if has_fp2 else ''
    return t1, t2


def detect_testsheet_format(wb):
    forced = os.environ.get("TESTSHEET_FORMAT", "").upper().strip()
    if forced in ["EET", "TIANU"]:
        return forced

    sheet_names_lower = [s.lower() for s in wb.sheetnames]
    if 'pce testsheet' in sheet_names_lower or 'pce vi' in sheet_names_lower:
        return 'EET'
    if 'substation data' in sheet_names_lower or 'vi' in sheet_names_lower:
        return 'TIANU'
    for s in wb.sheetnames:
        if re.match(r'^PE\s*\d+(\s*\(VI\))?$', s, re.IGNORECASE):
            return 'TIANU'
    return 'EET'


# ── Format Reader: EET (.xlsx) ────────────────────────────────────────────────
def _read_ts_eet(wb, fields):
    # 1. First read PCE VI to get switchgear type and main serial number
    if 'PCE VI' in wb.sheetnames:
        ws_vi = wb['PCE VI']

        vi_date_val = ws_vi['C5'].value or ws_vi['D5'].value
        if vi_date_val:
            fields["vi_date"] = normalize_date(vi_date_val)

        pe_erms = ws_vi['C6'].value
        if pe_erms:
            fields["pe_name_erms"] = str(pe_erms).strip()

        fl_erms = ws_vi['K6'].value
        if fl_erms:
            fields["fl_no_erms"] = str(fl_erms).strip()

        pe_site = ws_vi['C7'].value
        if pe_site:
            fields["pe_name_site"] = str(pe_site).strip()

        fl_site = ws_vi['K7'].value
        if fl_site:
            fields["fl_no_site"] = str(fl_site).strip()

        # Ensure raw values are stripped
        fields["pe_name_erms"] = clean_val(fields["pe_name_erms"])
        fields["fl_no_erms"]   = clean_val(fields["fl_no_erms"])
        fields["pe_name_site"] = clean_val(fields["pe_name_site"])
        fields["fl_no_site"]   = clean_val(fields["fl_no_site"])

        # Rule #5: Strict Data Integrity - blank / 0 / dash F/L set explicitly to N/A
        if not fields["fl_no_erms"] or fields["fl_no_erms"] in ['0', '0.0']:
            fields["fl_no_erms"] = "N/A"
        if not fields["fl_no_site"] or fields["fl_no_site"] in ['0', '0.0']:
            fields["fl_no_site"] = "N/A"

        gps_val = ws_vi['C8'].value
        if gps_val:
            fields["gps_coordinate"] = str(gps_val).strip()

        CHECK_MARKS = ['✅', '✓', '✔', 'x', 'X', '1', 'v', 'V', 'TRUE', 'YES']
        def is_checked(val):
            if val is None: return False
            s = str(val).strip()
            return any(c in s for c in CHECK_MARKS) or s in ['1', '1.0', 'T', 't', 'Y', 'y', 'X', 'x']

        type_cells = [(3, 4), (6, 7), (8, 9), (10, 11), (12, 13), (14, 15)]
        selected_types = []
        for type_col, check_col in type_cells:
            type_val = ws_vi.cell(row=9, column=type_col).value
            check_val = ws_vi.cell(row=9, column=check_col).value
            if type_val and is_checked(check_val):
                selected_types.append(str(type_val).strip())
        if selected_types:
            fields["substation_type"] = ' | '.join(selected_types)

        sg_types = []
        for row_num in [11, 14]:
            found_type = ""
            for col in range(2, 40):
                cell = ws_vi.cell(row=row_num, column=col)
                if is_checked(cell.value):
                    left_cell = ws_vi.cell(row=row_num, column=col - 1)
                    if left_cell.value:
                        found_type = str(left_cell.value).strip()
                    break
            sg_types.append(found_type)
        if any(sg_types):
            fields["sg_type"] = ' | '.join([t for t in sg_types if t])
            fields["sg_types_list"] = sg_types

        sg_manufs, sg_models = [], []
        for manuf_row, model_row in [(12, 12), (15, 15)]:
            manuf_val = ws_vi.cell(row=manuf_row, column=3).value
            model_val = ws_vi.cell(row=model_row, column=6).value
            if manuf_val:
                sg_manufs.append(clean_val(str(manuf_val)))
            if model_val:
                sg_models.append(clean_val(str(model_val)))
        fields["sg_manufacturer"] = ' | '.join([m for m in sg_manufs if m])
        fields["sg_model"] = ' | '.join([m for m in sg_models if m])
        fields["sg_manufs_list"] = sg_manufs
        fields["sg_models_list"] = sg_models

        sg_ratings, sg_serials = [], []
        for info_row in [13, 16]:
            rating_val = ws_vi.cell(row=info_row, column=10).value
            serial_val = ws_vi.cell(row=info_row, column=15).value
            if rating_val:
                sg_ratings.append(clean_val(str(rating_val)))
            if serial_val:
                sg_serials.append(clean_sn_val(serial_val))
        fields["sg_rating"] = ' | '.join([r for r in sg_ratings if r])
        fields["sg_serial"] = ' | '.join([s for s in sg_serials if s])
        fields["sg_ratings_list"] = sg_ratings
        fields["sg_serials_list"] = sg_serials
        fields["sg_types_list"] = sg_types

        # If switchgear type was not checked in Row 11, infer accurately from Model / Manufacturer
        if not fields.get("sg_type"):
            sg_mfr = fields.get("sg_manufacturer", "").upper().strip()
            sg_mdl = fields.get("sg_model", "").upper().strip()
            lbl_m13 = str(ws_vi.cell(row=13, column=13).value or '').upper()

            # 1. VCB
            if any(k in sg_mdl or k in sg_mfr for k in ["8BK", "VCB", "VACUUM", "NXAIR", "VHIH", "VK10", "VK 10", "VFT"]):
                fields["sg_type"] = "VCB"
            # 2. MRMU (Modular RMU modules like 2S+2F, 2S+1F, ZZT, 2L+1T, 8DJH, GA630)
            elif any(k in sg_mdl for k in ["MRMU", "2S+2F", "2S+1F", "2S1F", "2S2F", "2L+1T", "2L+2T", "2L1T", "ID2L", "ZZT", "ZZTT", "ZZZT", "2LBS", "8DJH", "8DJ20", "GA630"]):
                fields["sg_type"] = "MRMU"
            # 3. RMU SF6 (Standard compact units)
            elif any(k in sg_mdl or k in sg_mfr for k in ["INS24", "INS 24", "INS23", "RM6", "RM 6", "AEGIS", "VRN", "FLUSARC", "RINGMASTER", "FBX", "CCV", "GR1", "INDKOM", "SCHNEIDER", "TAMCO", "SSE", "LUCY", "ORMAZABAL", "EPE"]):
                fields["sg_type"] = "RMU SF6"
            elif "RMU" in lbl_m13:
                fields["sg_type"] = "RMU SF6"

            if fields.get("sg_type"):
                fields["sg_types_list"] = [fields["sg_type"]]

        for tx_row, prefix in [(18, "tx1"), (19, "tx2")]:
            tx_type = ws_vi.cell(row=tx_row, column=4).value
            tx_kva = ws_vi.cell(row=tx_row, column=6).value
            tx_mfr = ws_vi.cell(row=tx_row, column=11).value
            tx_sn = ws_vi.cell(row=tx_row, column=15).value
            if tx_type:
                tx_type_str = clean_val(str(tx_type))
                if tx_type_str.upper() in ["H/S", "H/S."]:
                    tx_type_str = "HERMETICALLY SEAL"
                fields[f"{prefix}_model"] = tx_type_str
            if tx_kva:
                kva_str = clean_val(str(tx_kva))
                if kva_str and 'kva' not in kva_str.lower():
                    kva_str = kva_str + 'kVA'
                fields[f"{prefix}_rating"] = kva_str
            if tx_mfr:
                fields[f"{prefix}_manufacturer"] = clean_val(str(tx_mfr))
            if tx_sn:
                fields[f"{prefix}_serial"] = clean_sn_val(tx_sn)

    # 2. Next read PCE Testsheet for panels & FP data
    if 'PCE Testsheet' in wb.sheetnames:
        ws = wb['PCE Testsheet']

        date_raw = ws['P4'].value
        if date_raw is not None:
            fields["date"] = normalize_date(date_raw)

        time_raw = ws['P5'].value
        if time_raw is not None:
            fields["time_in"] = format_time_val(time_raw)

        time_out_raw = ws['S5'].value
        if time_out_raw is not None:
            fields["time_out"] = format_time_val(time_out_raw)

        humidity_val = ws['S6'].value
        if humidity_val is not None:
            humidity_str = str(humidity_val)
            if humidity_str.endswith('.0'):
                humidity_str = humidity_str[:-2]
            fields["humidity"] = humidity_str

        ambient_raw = ws['W6'].value
        if ambient_raw is not None:
            m = re.search(r':\s*([\d\.]+)', str(ambient_raw))
            if m:
                ambient_val = m.group(1)
                try:
                    ambient_val = f"{float(ambient_val):.1f}"
                except ValueError:
                    pass
                fields["ambient"] = ambient_val

        # FP1 Mappings (User rule: FP1 & FP2 model always left blank)
        fields["lv1_model"] = ""

        r48_val = clean_val(str(ws['R48'].value)) if ws['R48'].value else ""
        t48_val = clean_val(str(ws['T48'].value)) if ws['T48'].value else ""
        fields["lv1_name"] = f"{r48_val} {t48_val}".strip()
        fields["lv1_type"] = fields["lv1_name"]

        lv1_manuf = ws['V49'].value
        if lv1_manuf:
            fields["lv1_manufacturer"] = clean_val(str(lv1_manuf))
        lv1_serial = ws['V50'].value
        if lv1_serial:
            fields["lv1_serial"] = clean_sn_val(lv1_serial)
        lv1_rating = ws['V51'].value
        if lv1_rating:
            fields["lv1_rating"] = clean_val(str(lv1_rating)).replace(' ', '')

        # FP2 Mappings (User rule: FP1 & FP2 model always left blank)
        fields["lv2_model"] = ""

        r52_val = clean_val(str(ws['R52'].value)) if ws['R52'].value else ""
        t52_val = clean_val(str(ws['T52'].value)) if ws['T52'].value else ""
        fields["lv2_name"] = f"{r52_val} {t52_val}".strip()
        fields["lv2_type"] = fields["lv2_name"]

        lv2_manuf = ws['V53'].value
        if lv2_manuf:
            fields["lv2_manufacturer"] = clean_val(str(lv2_manuf))
        lv2_serial = ws['V54'].value
        if lv2_serial:
            fields["lv2_serial"] = clean_sn_val(lv2_serial)
        lv2_rating = ws['V55'].value
        if lv2_rating:
            fields["lv2_rating"] = clean_val(str(lv2_rating)).replace(' ', '')

        # Determine FP / LVDB convention types based on TX setup
        has_tx1 = bool(fields.get("tx1_manufacturer") or fields.get("tx1_serial") or fields.get("tx1_rating"))
        has_tx2 = bool(fields.get("tx2_manufacturer") or fields.get("tx2_serial") or fields.get("tx2_rating"))
        fp1_d = {"manufacturer": fields.get("lv1_manufacturer"), "model": fields.get("lv1_model"), "serial": fields.get("lv1_serial"), "rating": fields.get("lv1_rating"), "name": fields.get("lv1_name")}
        fp2_d = {"manufacturer": fields.get("lv2_manufacturer"), "model": fields.get("lv2_model"), "serial": fields.get("lv2_serial"), "rating": fields.get("lv2_rating"), "name": fields.get("lv2_name")}
        c_t1, c_t2 = determine_fp_lvdb_types(has_tx1, has_tx2, fp1_d, fp2_d)
        if c_t1: fields["lv1_type"] = c_t1
        if c_t2: fields["lv2_type"] = c_t2

        bat_rows = {59: "bat1", 60: "bat2", 61: "bat3", 62: "bat4"}
        for row_num, prefix in bat_rows.items():
            b_val = ws.cell(row=row_num, column=2).value
            if b_val and ('battery' in str(b_val).lower() or 'black box' in str(b_val).lower()):
                j_val = ws.cell(row=row_num, column=10).value
                if j_val:
                    bat_str = str(j_val).strip()
                    bat_str = re.sub(r'^BC\d+:\s*', '', bat_str)
                    m_manuf = re.match(r'^(.*?)\s+M:', bat_str)
                    m_model = re.search(r'M:(.*?)(?:\s+SN:|$)', bat_str)
                    m_serial = re.search(r'SN:(.*)', bat_str)
                    if m_manuf:
                        fields[f"{prefix}_manufacturer"] = clean_val(m_manuf.group(1))
                    if m_model:
                        fields[f"{prefix}_model"] = clean_val(m_model.group(1))
                    if m_serial:
                        fields[f"{prefix}_serial"] = clean_sn_val(m_serial.group(1))

                    # Fallback parsing for comma-separated values without M:/SN: labels
                    if not fields.get(f"{prefix}_manufacturer") and ',' in bat_str:
                        parts = [p.strip() for p in bat_str.split(',')]
                        if len(parts) >= 1:
                            fields[f"{prefix}_manufacturer"] = clean_val(parts[0])
                        if len(parts) >= 2:
                            fields[f"{prefix}_model"] = clean_val(parts[1])
                        if len(parts) >= 3:
                            fields[f"{prefix}_serial"] = clean_sn_val(parts[2])

        target_rows = [10, 14, 18, 22]  # Panel 1, 2, 3, 4
        sg_type_clean = fields["sg_type"].upper()
        is_vcb = "VCB" in sg_type_clean

        sn_1 = fields["sg_serials_list"][0] if len(fields["sg_serials_list"]) > 0 else fields["sg_serial"]
        sn_2 = fields["sg_serials_list"][1] if len(fields["sg_serials_list"]) > 1 else fields["sg_serial"]

        def parse_panels(ws_sheet, default_sn):
            p_list = []
            for row in target_rows:
                desc_raw = ws_sheet.cell(row=row, column=1).value
                link_raw = ws_sheet.cell(row=row, column=2).value
                name_raw = ws_sheet.cell(row=row, column=3).value
                status_raw = ws_sheet.cell(row=row, column=5).value
                cable_raw = ws_sheet.cell(row=row, column=7).value
                heater_raw = ws_sheet.cell(row=row, column=8).value
                col9_sn = clean_sn_val(ws_sheet.cell(row=row, column=9).value)

                desc_val = clean_sn_val(desc_raw)
                link_val = clean_sn_val(link_raw)
                name_val = clean_val(name_raw)
                status_val = clean_val(status_raw)
                cable_val = clean_val(cable_raw)
                heater_val = format_heater_amp(heater_raw)

                # A valid panel must have at least a link number, panel name, breaker status, or serial number
                if not (link_val or name_val or status_val or col9_sn):
                    continue

                name_upper = name_val.upper().strip()
                desc_upper = desc_val.upper().strip()

                # Rule: VCB + TRANSITION PANEL / TRANSITION -> cable type is '-', link number / panel number is '-'
                is_transition = (
                    name_upper in ["TRANSITION", "TRANSITION PANEL", "TRANSITION-PANEL", "TRANSITION_PANEL"]
                    or bool(re.match(r'^TRANSITION(\s+PANEL)?$', name_upper))
                    or "TRANSITION" in name_upper
                    or "TRANSITION" in desc_upper
                )

                # Rule: BUS SECTION -> cable type is '-'
                is_bus_section = (
                    name_upper in ["BUS SECTION", "BUS-SECTION", "BUSSECTION", "BUS SECTION PANEL"]
                    or bool(re.match(r'^BUS\s+SECTION(\s+PANEL)?$', name_upper))
                    or "BUS SECTION" in name_upper
                )

                if is_vcb and is_transition:
                    cable_val = "-"
                    link_val = "-"
                elif is_bus_section:
                    cable_val = "-"

                if is_vcb:
                    # VCB panels have individual serial numbers in col 9; fallback to master SN if blank
                    panel_sn = col9_sn if col9_sn else default_sn
                else:
                    # Non-VCB (RMU, MRMU, etc.) or unticked type inherits master switchgear serial number
                    panel_sn = default_sn if default_sn else col9_sn

                panel_data = {
                    "description": desc_val,
                    "link_number": link_val,
                    "name": name_val,
                    "cable_type": cable_val,
                    "breaker_status": status_val,
                    "busbar_position": "-",
                    "load_amp": clean_val(ws_sheet.cell(row=row, column=6).value),
                    "heater_amp": heater_val,
                    "serial_number": panel_sn
                }
                p_list.append(panel_data)
            return p_list

        # Discover all testsheet tabs in order: PCE Testsheet, PCE Testsheet (1), PCE Testsheet (2), etc.
        ts_sheets = []
        if 'PCE Testsheet' in wb.sheetnames:
            ts_sheets.append('PCE Testsheet')
        extra_sheets = [s for s in wb.sheetnames if s != 'PCE Testsheet' and re.match(r'^PCE\s*Testsheet(?:\s*\(\d+\))?$', s, re.IGNORECASE)]
        extra_sheets.sort(key=lambda s: int(re.search(r'\d+', s).group(0)) if re.search(r'\d+', s) else 0)
        ts_sheets.extend(extra_sheets)

        all_panels = []
        for sname in ts_sheets:
            all_panels.extend(parse_panels(wb[sname], sn_1))

        fields["panels"] = all_panels
        fields["panels_1"] = all_panels
        if 'PCE Testsheet (2)' in wb.sheetnames and len(fields.get("sg_serials_list", [])) > 1:
            fields["panels_2"] = parse_panels(wb['PCE Testsheet (2)'], sn_2)
        else:
            fields["panels_2"] = []


# ── Format Reader: TIANU (.xlsx) ──────────────────────────────────────────────
def _read_ts_tianu(wb, fields, sub_num=None):
    # 1. Locate VI Sheet (Visual Inspection Checklist)
    ws_vi = None
    if 'VI' in wb.sheetnames:
        ws_vi = wb['VI']
    elif sub_num and f'PE {sub_num} (VI)' in wb.sheetnames:
        ws_vi = wb[f'PE {sub_num} (VI)']
    else:
        for s in wb.sheetnames:
            if '(VI)' in s or s.upper().endswith('VI'):
                ws_vi = wb[s]
                break

    # 2. Locate Substation Data Sheet (CBA Test Sheet)
    ws_data = None
    if 'Substation Data' in wb.sheetnames:
        ws_data = wb['Substation Data']
    elif sub_num and f'PE {sub_num}' in wb.sheetnames:
        ws_data = wb[f'PE {sub_num}']
    else:
        for s in wb.sheetnames:
            if s != (ws_vi.title if ws_vi else '') and ('Data' in s or re.match(r'^PE\s*\d+$', s, re.IGNORECASE)):
                ws_data = wb[s]
                break

    CHECK_MARKS = {'✅', '✓', '✔', 'x', 'X', '1', '1.0', 'v', 'V', 'TRUE', 'YES', '/', '\\'}
    def is_checked(val):
        if val is None:
            return False
        s = str(val).strip()
        return s in CHECK_MARKS or s.upper() in CHECK_MARKS

    # Parse VI Sheet
    if ws_vi is not None:
        vi_date_val = ws_vi['C5'].value or ws_vi['D5'].value or ws_vi['C51'].value
        if vi_date_val:
            fields['vi_date'] = normalize_date(vi_date_val)

        pe_erms = clean_val(ws_vi['C6'].value)
        if pe_erms:
            fields['pe_name_erms'] = pe_erms

        fl_erms = clean_val(ws_vi['K6'].value)
        fields['fl_no_erms'] = fl_erms if (fl_erms and fl_erms not in ['0', '0.0']) else 'N/A'

        pe_site = clean_val(ws_vi['C7'].value)
        if pe_site:
            fields['pe_name_site'] = pe_site

        fl_site = clean_val(ws_vi['K7'].value)
        fields['fl_no_site'] = fl_site if (fl_site and fl_site not in ['0', '0.0']) else 'N/A'

        gps_val = clean_val(ws_vi['C8'].value)
        if gps_val:
            fields['gps_coordinate'] = gps_val

        # Substation Type (Row 9)
        type_cells = [(3, 4), (6, 7), (8, 9), (10, 11), (12, 13), (14, 15)]
        selected_types = []
        for type_col, check_col in type_cells:
            t_val = ws_vi.cell(row=9, column=type_col).value
            c_val = ws_vi.cell(row=9, column=check_col).value
            if t_val and is_checked(c_val):
                selected_types.append(str(t_val).strip())
        if selected_types:
            fields['substation_type'] = ' | '.join(selected_types)

        # Switchgear Types (Row 11 & Row 14)
        sg_types = []
        for r_num in [11, 14]:
            found_type = ''
            for c_idx in [4, 7, 9, 11, 13, 15]:
                chk = ws_vi.cell(row=r_num, column=c_idx).value
                if is_checked(chk):
                    for left_off in [1, 2]:
                        lbl = ws_vi.cell(row=r_num, column=c_idx - left_off).value
                        if lbl and clean_val(lbl):
                            found_type = str(lbl).strip()
                            break
                    if found_type:
                        break
            sg_types.append(found_type)

        sg_manufs, sg_models = [], []
        for r_mfr in [12, 15]:
            mfr_cell = ws_vi.cell(row=r_mfr, column=3).value
            if mfr_cell and clean_val(mfr_cell):
                mfr_parsed, mdl_parsed = parse_sg_mfr_model(mfr_cell)
                if mfr_parsed:
                    sg_manufs.append(mfr_parsed)
                    sg_models.append(mdl_parsed)

        if sg_manufs:
            fields['sg_manufacturer'] = ' | '.join(sg_manufs)
            fields['sg_model'] = ' | '.join([m for m in sg_models if m])
            fields['sg_manufs_list'] = sg_manufs
            fields['sg_models_list'] = sg_models

        sg_ratings, sg_serials = [], []
        for r_info in [13, 16]:
            rtg = clean_val(ws_vi.cell(row=r_info, column=10).value)
            sn = clean_sn_val(ws_vi.cell(row=r_info, column=15).value)
            if rtg:
                sg_ratings.append(rtg)
            if sn:
                sg_serials.append(sn)

        if sg_ratings:
            fields['sg_rating'] = ' | '.join(sg_ratings)
            fields['sg_ratings_list'] = sg_ratings
        if sg_serials:
            fields['sg_serial'] = ' | '.join(sg_serials)
            fields['sg_serials_list'] = sg_serials

        if any(sg_types):
            fields['sg_type'] = ' | '.join([t for t in sg_types if t])
            fields['sg_types_list'] = [t for t in sg_types if t]

        # Infer Switchgear Type if not checked in Row 11
        if not fields.get('sg_type'):
            sg_mfr = fields.get('sg_manufacturer', '').upper()
            sg_mdl = fields.get('sg_model', '').upper()
            if any(k in sg_mdl or k in sg_mfr for k in ['8BK', 'VCB', 'VACUUM', 'NXAIR', 'VHIH', 'VK10', 'VK 10', 'VFT', 'HV12']):
                fields['sg_type'] = 'VCB'
            elif any(k in sg_mdl for k in ['MRMU', '2S+2F', '2S+1F', '2S1F', '2S2F', '2L+1T', '2L+2T', '2L1T', 'ID2L', 'ZZT', 'ZZTT', 'ZZZT', '2LBS', '8DJH', '8DJ20', 'GA630', 'FALCON']):
                fields['sg_type'] = 'MRMU'
            elif any(k in sg_mdl or k in sg_mfr for k in ['INS24', 'INS 24', 'INS23', 'RM6', 'RM 6', 'AEGIS', 'VRN', 'FLUSARC', 'RINGMASTER', 'FBX', 'CCV', 'GR1', 'INDKOM', 'SCHNEIDER', 'TAMCO', 'SSE', 'LUCY', 'ORMAZABAL', 'EPE', 'SCHALTANLAGE']):
                fields['sg_type'] = 'RMU SF6'
            if fields.get('sg_type'):
                fields['sg_types_list'] = [fields['sg_type']]

        # Transformers (TX 1 - TX 4)
        for tx_row, prefix in [(18, 'tx1'), (19, 'tx2'), (20, 'tx3'), (21, 'tx4')]:
            tx_type = ws_vi.cell(row=tx_row, column=4).value
            tx_kva = ws_vi.cell(row=tx_row, column=6).value
            tx_mfr = ws_vi.cell(row=tx_row, column=11).value
            tx_sn = ws_vi.cell(row=tx_row, column=15).value
            if tx_type:
                t_str = clean_val(str(tx_type))
                if t_str.upper() in ['H/S', 'H/S.', 'H/SEAL']:
                    t_str = 'HERMETICALLY SEAL'
                fields[f'{prefix}_model'] = t_str
            if tx_kva:
                k_str = clean_val(str(tx_kva))
                if k_str and 'kva' not in k_str.lower():
                    k_str += 'kVA'
                fields[f'{prefix}_rating'] = k_str
            if tx_mfr:
                fields[f'{prefix}_manufacturer'] = clean_val(str(tx_mfr))
            if tx_sn:
                fields[f'{prefix}_serial'] = clean_sn_val(tx_sn)

    # Parse Substation Data Sheet
    if ws_data is not None:
        d_val = ws_data['P2'].value
        if d_val:
            fields['date'] = normalize_date(d_val)
        elif fields.get('vi_date'):
            fields['date'] = fields['vi_date']

        t_in = ws_data['P3'].value
        if t_in:
            fields['time_in'] = format_time_val(t_in)
        t_out = ws_data['S3'].value
        if t_out:
            fields['time_out'] = format_time_val(t_out)

        hum_val = ws_data['S4'].value
        if hum_val is not None:
            if isinstance(hum_val, float) and 0 < hum_val <= 1.0:
                fields['humidity'] = str(int(round(hum_val * 100)))
            else:
                s_hum = str(hum_val).strip()
                if s_hum.endswith('.0'):
                    s_hum = s_hum[:-2]
                fields['humidity'] = s_hum

        amb_val = ws_data['Y4'].value or ws_data['W4'].value
        if amb_val is not None:
            fields['ambient'] = format_temp(amb_val)

        # Fallback Substation ERMS Name from C3 if missing from VI
        if not fields.get('pe_name_erms'):
            c3_val = clean_val(ws_data['C3'].value)
            if c3_val:
                fields['pe_name_erms'] = c3_val
                if not fields.get('pe_name_site'):
                    fields['pe_name_site'] = c3_val

        # Fallback FL Number from W3 if missing from VI
        w3_val = clean_val(ws_data['W3'].value)
        if w3_val and fields.get('fl_no_erms', 'N/A') == 'N/A':
            fields['fl_no_erms'] = w3_val
            if fields.get('fl_no_site', 'N/A') == 'N/A':
                fields['fl_no_site'] = w3_val

        # Fallback SG Serial from W8 if not present
        w8_sn = clean_sn_val(ws_data['W8'].value)
        if w8_sn and not fields.get('sg_serial'):
            fields['sg_serial'] = w8_sn
            if not fields.get('sg_serials_list'):
                fields['sg_serials_list'] = [w8_sn]

        # Feeder Pillars (FP1 & FP2)
        lv1_mfr = clean_val(ws_data['X48'].value)
        lv1_sn = clean_sn_val(ws_data['X49'].value)
        lv1_rtg = clean_val(ws_data['X50'].value).replace(' ', '')
        if lv1_mfr:
            fields['lv1_manufacturer'] = lv1_mfr
        if lv1_sn:
            fields['lv1_serial'] = lv1_sn
        if lv1_rtg:
            fields['lv1_rating'] = lv1_rtg
        fields['lv1_cable_type'] = clean_val(ws_data['D44'].value or 'XLPE')
        fields['lv1_model'] = ""  # User rule: FP1 & FP2 model always left blank
        fields['lv1_name'] = ""

        lv2_mfr = clean_val(ws_data['X52'].value)
        lv2_sn = clean_sn_val(ws_data['X53'].value)
        lv2_rtg = clean_val(ws_data['X54'].value).replace(' ', '')
        if lv2_mfr:
            fields['lv2_manufacturer'] = lv2_mfr
        if lv2_sn:
            fields['lv2_serial'] = lv2_sn
        if lv2_rtg:
            fields['lv2_rating'] = lv2_rtg
        fields['lv2_cable_type'] = clean_val(ws_data['D46'].value or 'XLPE')
        fields['lv2_model'] = ""  # User rule: FP1 & FP2 model always left blank
        fields['lv2_name'] = ""

        # Batteries (Row 58 = bat1, Row 59 = bat2, Row 60 = bat3)
        for r_bat, pfx in [(58, 'bat1'), (59, 'bat2'), (60, 'bat3')]:
            b_mfr = clean_val(ws_data.cell(row=r_bat, column=10).value)
            b_mdl = clean_val(ws_data.cell(row=r_bat, column=12).value)
            b_sn = clean_sn_val(ws_data.cell(row=r_bat, column=13).value)

            # Rule #1 Tag/Comma Parsing Fallback if combined in single column
            if b_mfr and ('M:' in b_mfr or ',' in b_mfr):
                bat_str = b_mfr
                m_manuf = re.match(r'^(.*?)\s+M:', bat_str)
                m_model = re.search(r'M:(.*?)(?:\s+SN:|$)', bat_str)
                m_serial = re.search(r'SN:(.*)', bat_str)
                if m_manuf:
                    b_mfr = clean_val(m_manuf.group(1))
                if m_model:
                    b_mdl = clean_val(m_model.group(1))
                if m_serial:
                    b_sn = clean_sn_val(m_serial.group(1))
                if not b_mdl and ',' in bat_str:
                    parts = [p.strip() for p in bat_str.split(',')]
                    if len(parts) >= 1: b_mfr = clean_val(parts[0])
                    if len(parts) >= 2: b_mdl = clean_val(parts[1])
                    if len(parts) >= 3: b_sn = clean_sn_val(parts[2])

            if b_mfr:
                fields[f'{pfx}_manufacturer'] = b_mfr
            if b_mdl:
                fields[f'{pfx}_model'] = b_mdl
            if b_sn:
                fields[f'{pfx}_serial'] = b_sn

        # ── Bottom Summary Block Fallbacks (Rows 175-275 in Substation Data) ──
        # Switchgear 1 Type fallback (F177)
        if not fields.get('sg_type'):
            f177_val = clean_val(ws_data['F177'].value)
            if f177_val and f177_val not in ['-', 'N/A']:
                fields['sg_type'] = f177_val
                if not fields.get('sg_types_list'):
                    fields['sg_types_list'] = [f177_val]

        # Switchgear 1 Serial (C186) & Manufacturer/Model (C187)
        c186_sn = clean_sn_val(ws_data['C186'].value)
        if c186_sn and not fields.get('sg_serial'):
            fields['sg_serial'] = c186_sn
            fields['sg_serials_list'] = [c186_sn]

        c187_val = clean_val(ws_data['C187'].value)
        if c187_val and not fields.get('sg_manufacturer'):
            m_parsed, mdl_parsed = parse_sg_mfr_model(c187_val)
            if m_parsed:
                fields['sg_manufacturer'] = m_parsed
                fields['sg_model'] = mdl_parsed
                fields['sg_manufs_list'] = [m_parsed]
                fields['sg_models_list'] = [mdl_parsed]

        # Switchgear 2 Serial (C198) & Manufacturer/Model (C199)
        c198_sn = clean_sn_val(ws_data['C198'].value)
        c199_val = clean_val(ws_data['C199'].value)
        if c198_sn or c199_val:
            if c198_sn and len(fields.get('sg_serials_list', [])) < 2:
                fields['sg_serials_list'] = fields.get('sg_serials_list', []) + [c198_sn]
                fields['sg_serial'] = ' | '.join(fields['sg_serials_list'])
            if c199_val and len(fields.get('sg_manufs_list', [])) < 2:
                m2_parsed, mdl2_parsed = parse_sg_mfr_model(c199_val)
                if m2_parsed:
                    fields['sg_manufs_list'] = fields.get('sg_manufs_list', []) + [m2_parsed]
                    fields['sg_models_list'] = fields.get('sg_models_list', []) + [mdl2_parsed]
                    fields['sg_manufacturer'] = ' | '.join(fields['sg_manufs_list'])
                    fields['sg_model'] = ' | '.join([m for m in fields['sg_models_list'] if m])

        # Transformers TX 1 - TX 4 fallbacks from Rows 223-242
        tx_bottom_map = [
            ('tx1', 223, 224, 225, 227),
            ('tx2', 228, 229, 230, 232),
            ('tx3', 233, 234, 235, 237),
            ('tx4', 238, 239, 240, 242),
        ]
        for tx_pfx, r_mfr, r_kva, r_mdl, r_sn in tx_bottom_map:
            if not fields.get(f'{tx_pfx}_manufacturer'):
                b_mfr = clean_val(ws_data.cell(row=r_mfr, column=3).value)
                if b_mfr: fields[f'{tx_pfx}_manufacturer'] = b_mfr
            if not fields.get(f'{tx_pfx}_rating'):
                b_kva = clean_val(ws_data.cell(row=r_kva, column=3).value)
                if b_kva:
                    if 'kva' not in b_kva.lower(): b_kva += ' kVA'
                    fields[f'{tx_pfx}_rating'] = b_kva
            if not fields.get(f'{tx_pfx}_model'):
                b_mdl = clean_val(ws_data.cell(row=r_mdl, column=3).value)
                if b_mdl:
                    if b_mdl.upper() in ['H/S', 'H/S.', 'H/SEAL']: b_mdl = 'HERMETICALLY SEAL'
                    fields[f'{tx_pfx}_model'] = b_mdl
            if not fields.get(f'{tx_pfx}_serial'):
                b_sn = clean_sn_val(ws_data.cell(row=r_sn, column=3).value)
                if b_sn: fields[f'{tx_pfx}_serial'] = b_sn

        # Feeder Pillars fallbacks from Rows 246-253
        fp1_name = clean_val(ws_data['C246'].value) if ws_data['C246'].value else ""
        if fp1_name and fp1_name not in ['-', 'N/A']:
            fields['lv1_name'] = fp1_name
        if not fields.get('lv1_manufacturer'):
            fp1_mfr = clean_val(ws_data['C247'].value)
            if fp1_mfr: fields['lv1_manufacturer'] = fp1_mfr
        if not fields.get('lv1_serial'):
            fp1_sn = clean_sn_val(ws_data['C248'].value)
            if fp1_sn: fields['lv1_serial'] = fp1_sn
        if not fields.get('lv1_rating'):
            fp1_rtg = clean_val(ws_data['C249'].value).replace(' ', '')
            if fp1_rtg: fields['lv1_rating'] = fp1_rtg

        fp2_name = clean_val(ws_data['C250'].value) if ws_data['C250'].value else ""
        if fp2_name and fp2_name not in ['-', 'N/A']:
            fields['lv2_name'] = fp2_name
        if not fields.get('lv2_manufacturer'):
            fp2_mfr = clean_val(ws_data['C251'].value)
            if fp2_mfr: fields['lv2_manufacturer'] = fp2_mfr
        if not fields.get('lv2_serial'):
            fp2_sn = clean_sn_val(ws_data['C252'].value)
            if fp2_sn: fields['lv2_serial'] = fp2_sn
        if not fields.get('lv2_rating'):
            fp2_rtg = clean_val(ws_data['C253'].value).replace(' ', '')
            if fp2_rtg: fields['lv2_rating'] = fp2_rtg

        # Determine FP / LVDB convention types based on TX setup
        has_tx1 = bool(fields.get("tx1_manufacturer") or fields.get("tx1_serial") or fields.get("tx1_rating"))
        has_tx2 = bool(fields.get("tx2_manufacturer") or fields.get("tx2_serial") or fields.get("tx2_rating"))
        fp1_d = {"manufacturer": fields.get("lv1_manufacturer"), "model": fields.get("lv1_model"), "serial": fields.get("lv1_serial"), "rating": fields.get("lv1_rating"), "name": fields.get("lv1_name")}
        fp2_d = {"manufacturer": fields.get("lv2_manufacturer"), "model": fields.get("lv2_model"), "serial": fields.get("lv2_serial"), "rating": fields.get("lv2_rating"), "name": fields.get("lv2_name")}
        c_t1, c_t2 = determine_fp_lvdb_types(has_tx1, has_tx2, fp1_d, fp2_d)
        if c_t1: fields["lv1_type"] = c_t1
        if c_t2: fields["lv2_type"] = c_t2

        # Batteries fallbacks from Rows 264-275
        bat_bottom_map = [
            ('bat1', 264, 265, 266),
            ('bat2', 268, 269, 270),
            ('bat3', 272, 273, 274),
        ]
        for bat_pfx, r_bmfr, r_bmdl, r_bsn in bat_bottom_map:
            if not fields.get(f'{bat_pfx}_manufacturer'):
                bm = clean_val(ws_data.cell(row=r_bmfr, column=3).value)
                if bm: fields[f'{bat_pfx}_manufacturer'] = bm
            if not fields.get(f'{bat_pfx}_model'):
                bmd = clean_val(ws_data.cell(row=r_bmdl, column=3).value)
                if bmd: fields[f'{bat_pfx}_model'] = bmd
            if not fields.get(f'{bat_pfx}_serial'):
                bsn = clean_sn_val(ws_data.cell(row=r_bsn, column=3).value)
                if bsn: fields[f'{bat_pfx}_serial'] = bsn

        # Panels (Rows 8, 12, 16, 20...)
        sg_type_clean = fields['sg_type'].upper()
        is_vcb = 'VCB' in sg_type_clean
        default_sn = fields['sg_serials_list'][0] if len(fields.get('sg_serials_list', [])) > 0 else fields['sg_serial']

        panels = []
        for r_pnl in range(8, 32, 4):
            c1_val = ws_data.cell(row=r_pnl, column=1).value
            if c1_val in ['OVERVIEW', 'LOCAL TRANSFORMER']:
                break
            desc_raw = clean_sn_val(c1_val)
            link_raw = clean_sn_val(ws_data.cell(row=r_pnl, column=2).value)
            name_raw = clean_val(ws_data.cell(row=r_pnl, column=3).value)
            status_raw = clean_val(ws_data.cell(row=r_pnl, column=5).value)
            load_raw = clean_val(ws_data.cell(row=r_pnl, column=6).value)
            cable_raw = clean_val(ws_data.cell(row=r_pnl, column=7).value)
            heater_raw = format_heater_amp(ws_data.cell(row=r_pnl, column=8).value)
            col9_sn = clean_sn_val(ws_data.cell(row=r_pnl, column=9).value)

            if not (link_raw or name_raw or status_raw or col9_sn):
                continue
            if name_raw in ['-', 'N/A', 'None', ''] and link_raw in ['-', 'N/A', 'None', '']:
                continue

            status_val = '-'
            if 'CLOSE' in status_raw.upper():
                status_val = 'CLOSE'
            elif 'TRIP' in status_raw.upper():
                status_val = 'TRIP'
            elif status_raw:
                status_val = status_raw

            cable_val = cable_raw if cable_raw else 'XLPE'

            name_upper = name_raw.upper().strip()
            desc_upper = desc_raw.upper().strip()
            is_transition = ('TRANSITION' in name_upper or 'TRANSITION' in desc_upper)
            is_bus_section = ('BUS SECTION' in name_upper or 'BUS-SECTION' in name_upper)

            if is_vcb and is_transition:
                cable_val = '-'
                link_raw = '-'
            elif is_bus_section:
                cable_val = '-'

            if is_vcb:
                panel_sn = col9_sn if col9_sn else default_sn
            else:
                panel_sn = default_sn if default_sn else col9_sn

            panel_data = {
                'description': desc_raw,
                'link_number': link_raw,
                'name': name_raw,
                'cable_type': cable_val,
                'breaker_status': status_val,
                'busbar_position': '-',
                'load_amp': load_raw,
                'heater_amp': heater_raw,
                'serial_number': panel_sn
            }
            panels.append(panel_data)

        fields['panels'] = panels
        fields['panels_1'] = panels
        fields['panels_2'] = []


# ── Helper: extract ALL fields from an Excel testsheet (.xlsx) ────────────────
def read_ts_fields_excel(xlsx_path, sub_num=None):
    fields = {
        "date": "", "ambient": "", "humidity": "", "time_in": "", "time_out": "",
        "sg_manufacturer": "", "sg_model": "", "sg_rating": "", "sg_serial": "", "sg_type": "",
        "tx1_manufacturer": "", "tx1_model": "", "tx1_rating": "", "tx1_serial": "",
        "tx2_manufacturer": "", "tx2_model": "", "tx2_rating": "", "tx2_serial": "",
        "lv1_manufacturer": "", "lv1_model": "", "lv1_serial": "", "lv1_rating": "", "lv1_type": "",
        "lv2_manufacturer": "", "lv2_model": "", "lv2_serial": "", "lv2_rating": "", "lv2_type": "",
        "bat1_manufacturer": "", "bat1_model": "", "bat1_serial": "",
        "bat2_manufacturer": "", "bat2_model": "", "bat2_serial": "",
        "bat3_manufacturer": "", "bat3_model": "", "bat3_serial": "",
        "bat4_manufacturer": "", "bat4_model": "", "bat4_serial": "",
        "pe_name_erms": "", "pe_name_site": "",
        "fl_no_erms": "", "fl_no_site": "",
        "gps_coordinate": "", "vi_date": "",
        "substation_type": "",
        "sg_manufs_list": [], "sg_models_list": [],
        "sg_ratings_list": [], "sg_serials_list": [], "sg_types_list": [],
        "panels": [], "panels_1": [], "panels_2": [],
        "_format": "EET"
    }
    try:
        wb = openpyxl.load_workbook(xlsx_path, data_only=True)
        fmt = detect_testsheet_format(wb)
        fields["_format"] = fmt
        if fmt == "TIANU":
            _read_ts_tianu(wb, fields, sub_num)
        else:
            _read_ts_eet(wb, fields)
        wb.close()

        # Enforce FP / LVDB convention types based on TX setup
        has_tx1 = bool(fields.get("tx1_manufacturer") or fields.get("tx1_serial") or fields.get("tx1_rating"))
        has_tx2 = bool(fields.get("tx2_manufacturer") or fields.get("tx2_serial") or fields.get("tx2_rating"))
        fp1_d = {"manufacturer": fields.get("lv1_manufacturer"), "model": fields.get("lv1_model"), "serial": fields.get("lv1_serial"), "rating": fields.get("lv1_rating"), "name": fields.get("lv1_name"), "type": fields.get("lv1_type")}
        fp2_d = {"manufacturer": fields.get("lv2_manufacturer"), "model": fields.get("lv2_model"), "serial": fields.get("lv2_serial"), "rating": fields.get("lv2_rating"), "name": fields.get("lv2_name"), "type": fields.get("lv2_type")}
        c_t1, c_t2 = determine_fp_lvdb_types(has_tx1, has_tx2, fp1_d, fp2_d)
        if c_t1: fields["lv1_type"] = c_t1
        if c_t2: fields["lv2_type"] = c_t2
    except Exception:
        pass
    return fields


# ── Helper: extract ALL fields from a PDF testsheet (.pdf fallback) ───────────
def read_ts_fields_pdf(pdf_path):
    fields = {
        "date": "", "ambient": "", "humidity": "", "time_in": "", "time_out": "",
        "sg_manufacturer": "", "sg_model": "", "sg_rating": "", "sg_serial": "", "sg_type": "",
        "tx1_manufacturer": "", "tx1_model": "", "tx1_rating": "", "tx1_serial": "",
        "tx2_manufacturer": "", "tx2_model": "", "tx2_rating": "", "tx2_serial": "",
        "lv1_manufacturer": "", "lv1_model": "", "lv1_serial": "", "lv1_rating": "", "lv1_type": "",
        "lv2_manufacturer": "", "lv2_model": "", "lv2_serial": "", "lv2_rating": "", "lv2_type": "",
        "bat1_manufacturer": "", "bat1_model": "", "bat1_serial": "",
        "bat2_manufacturer": "", "bat2_model": "", "bat2_serial": "",
        "bat3_manufacturer": "", "bat3_model": "", "bat3_serial": "",
        "bat4_manufacturer": "", "bat4_model": "", "bat4_serial": "",
        "pe_name_erms": "", "pe_name_site": "",
        "fl_no_erms": "", "fl_no_site": "",
        "gps_coordinate": "", "vi_date": "",
        "substation_type": "",
        "sg_manufs_list": [],
        "sg_models_list": [],
        "sg_ratings_list": [],
        "sg_serials_list": [],
        "sg_types_list": [],
        "panels": [],
        "panels_1": [],
        "panels_2": []
    }

    if not fitz:
        print(f"  [PyMuPDF Missing] fitz is not installed; unable to parse PDF: {pdf_path}", flush=True)
        return fields

    try:
        doc = fitz.open(pdf_path)
        if len(doc) == 0:
            return fields

        # 1. Parse VI Page (Visual Inspection Checklist, always last page)
        vi_page = doc[-1]
        vi_words = sorted(vi_page.get_text('words'), key=lambda w: (round(w[1] / 3) * 3, w[0]))
        vi_lines = {}
        for w in vi_words:
            y = round(w[1] / 3) * 3
            vi_lines.setdefault(y, []).append(w[4])

        # Metadata block: Look for blocks near x=[110, 140], y=[110, 140]
        for b in vi_page.get_text('blocks'):
            if 110 <= b[0] <= 140 and 110 <= b[1] <= 140:
                lines = [l.strip() for l in b[4].splitlines() if l.strip()]
                if len(lines) >= 1:
                    fields["vi_date"] = normalize_date(lines[0])
                if len(lines) >= 2:
                    fields["pe_name_erms"] = clean_val(lines[1])
                if len(lines) >= 3:
                    fields["fl_no_erms"] = clean_val(lines[2])
                if len(lines) >= 4:
                    fields["pe_name_site"] = clean_val(lines[3])
                if len(lines) >= 5:
                    fields["fl_no_site"] = clean_val(lines[4])
                break

        # GPS coordinate
        for b in vi_page.get_text('blocks'):
            m_gps = re.search(r'(-?\d+\.\d+,\s*-?\d+\.\d+)', b[4])
            if m_gps:
                fields["gps_coordinate"] = m_gps.group(1).replace(' ', '')
                break

        # Substation Type: check for ticked type (e.g. INDOOR / OUTDOOR)
        for y, lwords in vi_lines.items():
            if 150 <= y <= 165:
                lstr = ' '.join(lwords)
                if '/' in lstr:
                    parts = lstr.split('/')
                    left_w = parts[0].strip().split()
                    if left_w:
                        for st in ["INDOOR", "OUTDOOR", "COMPACT", "ATTACH BUILDING", "PAT"]:
                            if st in left_w[-1].upper():
                                fields["substation_type"] = st
                                break

        # Check switchgear type
        sg_type_val = ""
        vi_full_txt = vi_page.get_text()
        for y, lwords in vi_lines.items():
            if 165 <= y <= 180:
                lstr = ' '.join(lwords)
                if '/' in lstr:
                    parts = lstr.split('/')
                    left_p = parts[0].strip().split()
                    if left_p:
                        candidate = ' '.join(left_p[-2:]) if len(left_p) >= 2 else left_p[-1]
                        for st in ["RMU SF6", "RMU OIL", "MRMU", "VCB", "OCB"]:
                            if st in candidate.upper():
                                sg_type_val = st
                                break

        sg_mfr = ""
        sg_mdl = ""
        sg_rating = ""
        sg_serial = ""

        # Extract Manufacturer and Model from the specific block at y ≈ 172..184
        for b in vi_page.get_text('blocks'):
            if 115 <= b[0] <= 145 and 172 <= b[1] <= 184:
                first_l = b[4].strip().splitlines()[0] if b[4].strip() else ""
                if first_l and first_l not in ['-', 'N/A']:
                    m_parts = first_l.split()
                    if len(m_parts) >= 2:
                        sg_mfr = m_parts[0]
                        sg_mdl = ' '.join(m_parts[1:])
                    elif len(m_parts) == 1:
                        sg_mfr = m_parts[0]

            # Switchgear Serial block at b[0] > 550, 175 <= b[1] <= 200
            elif b[0] > 550 and 175 <= b[1] <= 200:
                cand = b[4].strip()
                if not any(k in cand.upper() for k in ['SWITCGEAR', 'SWITCHGEAR', 'TYPE', 'RATING', 'DATE', 'MANUFAC']):
                    sn_cand = cand.replace('S/N', '').replace(':', '').strip()
                    if sn_cand:
                        sg_serial = clean_sn_val(sn_cand.splitlines()[0])

            # Rating in blocks
            if not sg_rating and re.search(r'\b\d{1,2}\s*kV\b', b[4], re.IGNORECASE):
                m_rtg = re.search(r'\b(\d{1,2}\s*kV)\b', b[4], re.IGNORECASE)
                if m_rtg:
                    sg_rating = m_rtg.group(1).replace(' ', '')

        if not sg_mfr:
            for y, lwords in vi_lines.items():
                if 170 <= y <= 190:
                    lstr = ' '.join(lwords)
                    for std_name, aliases in KNOWN_SG_MANUFACTURERS:
                        if any(al in lstr.upper() for al in aliases):
                            sg_mfr = std_name
                            m_rest = re.search(rf'{std_name}\s+([A-Za-z0-9\-_]+)', lstr, re.IGNORECASE)
                            if m_rest:
                                sg_mdl = m_rest.group(1)
                            break
                    if sg_mfr:
                        break

        if not sg_serial:
            for y, lwords in vi_lines.items():
                if 180 <= y <= 195:
                    lstr = ' '.join(lwords)
                    m_sn = re.search(r'(?:S/N\s*:\s*|OIL\s*S/N\s*:\s*)([A-Za-z0-9\-_/]+)', lstr)
                    if m_sn:
                        cand = m_sn.group(1)
                        if not any(k in cand.upper() for k in ['SWITCGEAR', 'SWITCHGEAR', 'TYPE', 'RATING']):
                            sg_serial = clean_sn_val(cand)

        if not sg_type_val:
            sg_comb = f"{sg_mfr} {sg_mdl}".upper()
            if any(k in sg_comb for k in ["8BK", "VCB", "VACUUM", "NXAIR", "VHIH", "VK10", "VFT", "HV12"]):
                sg_type_val = "VCB"
            elif any(k in sg_comb for k in ["MRMU", "2S+2F", "2S+1F", "2S1F", "2S2F", "2L+1T", "8DJH", "GA630"]):
                sg_type_val = "MRMU"
            elif any(k in sg_comb for k in ["INS24", "INS 24", "RM6", "VRN", "FLUSARC", "RINGMASTER", "GR1", "FNG"]):
                sg_type_val = "RMU SF6"
            elif sg_mfr:
                sg_type_val = "RMU SF6"

        fields["sg_type"] = sg_type_val
        fields["sg_manufacturer"] = clean_val(sg_mfr)
        fields["sg_model"] = clean_val(sg_mdl)
        fields["sg_rating"] = clean_val(sg_rating)
        fields["sg_serial"] = clean_sn_val(sg_serial)

        fields["sg_manufs_list"] = [fields["sg_manufacturer"]] if fields["sg_manufacturer"] else []
        fields["sg_models_list"] = [fields["sg_model"]] if fields["sg_model"] else []
        fields["sg_ratings_list"] = [fields["sg_rating"]] if fields["sg_rating"] else []
        fields["sg_serials_list"] = [fields["sg_serial"]] if fields["sg_serial"] else []
        fields["sg_types_list"] = [fields["sg_type"]] if fields["sg_type"] else []

        # Transformers 1 & 2 from VI page
        def get_text_in_rect(page, rect):
            words = [w[4] for w in page.get_text('words') if rect[0] <= w[0] <= rect[2] and rect[1] <= w[1] <= rect[3]]
            return ' '.join(words).strip()

        def clean_tok(s):
            toks = [t.strip(':').strip('-').strip() for t in s.split() if t.strip(':').strip('-').strip()]
            return toks[0] if toks else ''

        mfr1 = clean_tok(get_text_in_rect(vi_page, (430, 222, 580, 234)).replace('Manufacturer', ''))
        kva1 = clean_tok(get_text_in_rect(vi_page, (225, 222, 310, 234)).replace('KVA', ''))
        sn1  = clean_tok(get_text_in_rect(vi_page, (600, 222, 750, 234)).replace('S/N', ''))
        typ1 = get_text_in_rect(vi_page, (140, 222, 225, 234))
        if any(k in typ1.upper() for k in ['H/SEAL', 'H/S', 'HERMETICALLY']):
            fields["tx1_model"] = "HERMETICALLY SEAL"
        if kva1 and kva1.isdigit():
            fields["tx1_rating"] = f"{kva1}kVA"
        if mfr1:
            fields["tx1_manufacturer"] = clean_val(mfr1)
        if sn1:
            fields["tx1_serial"] = clean_sn_val(sn1)

        mfr2 = clean_tok(get_text_in_rect(vi_page, (430, 233, 580, 243)).replace('Manufacturer', ''))
        kva2 = clean_tok(get_text_in_rect(vi_page, (225, 233, 310, 243)).replace('KVA', ''))
        sn2  = clean_tok(get_text_in_rect(vi_page, (600, 233, 750, 243)).replace('S/N', ''))
        typ2 = get_text_in_rect(vi_page, (140, 233, 225, 243))
        if any(k in typ2.upper() for k in ['H/SEAL', 'H/S', 'HERMETICALLY']):
            fields["tx2_model"] = "HERMETICALLY SEAL"
        if kva2 and kva2.isdigit():
            fields["tx2_rating"] = f"{kva2}kVA"
        if mfr2:
            fields["tx2_manufacturer"] = clean_val(mfr2)
        if sn2:
            fields["tx2_serial"] = clean_sn_val(sn2)

        # Fallback Construction block for Transformer Manufacturer/Serial if missing
        for b in vi_page.get_text('blocks'):
            if b[0] > 300 and 220 <= b[1] <= 255 and 'Construction' in b[4]:
                tx_b_lines = [l.strip() for l in b[4].splitlines() if l.strip()]
                items = [item for item in tx_b_lines if 'Construction' not in item]
                if len(items) >= 2:
                    if not fields["tx1_manufacturer"] or fields["tx1_manufacturer"].upper() in ['S', 'S/N']:
                        fields["tx1_manufacturer"] = clean_val(items[0])
                    if not fields["tx1_serial"]:
                        fields["tx1_serial"] = clean_sn_val(items[1])
                if len(items) >= 4:
                    if not fields["tx2_manufacturer"] or fields["tx2_manufacturer"].upper() in ['S', 'S/N']:
                        fields["tx2_manufacturer"] = clean_val(items[2])
                    if not fields["tx2_serial"]:
                        fields["tx2_serial"] = clean_sn_val(items[3])

        # 2. Parse CBA Test Sheet Page(s)
        cba_page = doc[0]
        words0 = [w[4] for w in sorted(cba_page.get_text('words'), key=lambda w: (round(w[1] / 3) * 3, w[0])) if w[1] < 70]
        txt0 = ' '.join(words0)

        m_cd = re.search(r'DATE\s+(\d{1,2}[-/]\d{1,2}[-/]\d{2,4})', txt0)
        if m_cd:
            fields["date"] = normalize_date(m_cd.group(1))
        elif fields.get("vi_date"):
            fields["date"] = fields["vi_date"]

        times = re.findall(r'\b(\d{1,2}:\d{2}\s*(?:AM|PM)?)\b', txt0, re.IGNORECASE)
        if len(times) >= 1:
            fields["time_in"] = times[0]
        if len(times) >= 2:
            fields["time_out"] = times[1]

        m_a = re.search(r'(\d{2}(?:\.\d+)?)\s*°?\s*C', txt0, re.IGNORECASE) or re.search(r'AMBIENT\s+TEMP\s*:\s*(\d{2}(?:\.\d+)?)', txt0, re.IGNORECASE)
        if m_a:
            try:
                fields["ambient"] = f"{float(m_a.group(1)):.1f}"
            except ValueError:
                fields["ambient"] = m_a.group(1)

        m_h = re.search(r'(\d{2})%', txt0)
        if m_h:
            fields["humidity"] = m_h.group(1)

        # FP1 / FP2
        cba_words = sorted(cba_page.get_text('words'), key=lambda w: (round(w[1] / 3) * 3, w[0]))
        cba_lines = {}
        for w in cba_words:
            y = round(w[1] / 3) * 3
            cba_lines.setdefault(y, []).append(w[4])

        cur_fp = None
        for y in sorted(cba_lines.keys()):
            lstr = ' '.join(cba_lines[y])
            if 'FEEDER PILLAR 1' in lstr:
                cur_fp = "lv1"
            elif 'FEEDER PILLAR 2' in lstr:
                cur_fp = "lv2"

            if cur_fp and 'Manufacturer' in lstr:
                m_mfr = re.search(r'Manufacturer\s*:\s*(.*)', lstr)
                if m_mfr and m_mfr.group(1).strip():
                    fields[f"{cur_fp}_manufacturer"] = clean_val(m_mfr.group(1).strip())
            elif cur_fp and 'Serial' in lstr:
                m_sn = re.search(r'Serial\s*No\.?\s*:\s*(.*)', lstr)
                if m_sn and m_sn.group(1).strip():
                    fields[f"{cur_fp}_serial"] = clean_sn_val(m_sn.group(1).strip())
            elif cur_fp and 'Rating' in lstr:
                m_rtg = re.search(r'Rating\s*:\s*(.*)', lstr)
                if m_rtg and m_rtg.group(1).strip():
                    fields[f"{cur_fp}_rating"] = clean_val(m_rtg.group(1).strip().replace(' ', ''))

        # Battery Banks
        bat_map = [
            ("Battery Bank", "bat1"),
            ("Black Box 1", "bat2"),
            ("Black Box 2", "bat3")
        ]
        for y in sorted(cba_lines.keys()):
            lstr = ' '.join(cba_lines[y])
            for b_label, b_prefix in bat_map:
                if b_label in lstr:
                    b_rest = lstr.replace(b_label, '').replace('240V', '').strip()
                    b_rest = re.sub(r'^[-\s]+', '', b_rest)
                    if b_rest and not all(c in '- ' for c in b_rest):
                        parts = b_rest.split()
                        if len(parts) == 1:
                            fields[f"{b_prefix}_manufacturer"] = clean_val(parts[0])
                        elif len(parts) == 2:
                            fields[f"{b_prefix}_manufacturer"] = clean_val(parts[0])
                            fields[f"{b_prefix}_model"] = clean_val(parts[1])
                        elif len(parts) >= 3:
                            if parts[0].upper() == "BERKAT" and parts[1].upper() == "INSAF":
                                fields[f"{b_prefix}_manufacturer"] = "BERKAT INSAF"
                                fields[f"{b_prefix}_model"] = clean_val(parts[2])
                                fields[f"{b_prefix}_serial"] = clean_sn_val(' '.join(parts[3:]))
                            else:
                                fields[f"{b_prefix}_manufacturer"] = clean_val(parts[0])
                                fields[f"{b_prefix}_model"] = clean_val(parts[1])
                                fields[f"{b_prefix}_serial"] = clean_sn_val(' '.join(parts[2:]))

        # 3. Parse Circuit Panels
        all_panels = []
        is_vcb = "VCB" in fields["sg_type"].upper()
        default_sn = fields["sg_serial"]

        for pno in range(len(doc) - 1):
            page_p = doc[pno]
            pw = sorted(page_p.get_text('words'), key=lambda w: (round(w[1] / 3) * 3, w[0]))
            plines = {}
            for w in pw:
                y = round(w[1] / 3) * 3
                plines.setdefault(y, []).append(w[4])

            for y in sorted(plines.keys()):
                if 85 <= y <= 225:
                    lstr = ' '.join(plines[y])
                    m = re.match(r'^\s*(\d+)\s+([A-Za-z0-9\-_/]+)\s+(.*)', lstr)
                    if m:
                        p_num = m.group(1)
                        link_no = m.group(2)
                        rest = m.group(3)

                        if 'TOP PANEL' in rest:
                            rest = rest.replace('TOP PANEL', '').strip()

                        tokens = rest.split()
                        status = ''
                        load_val = ''
                        cable_val = ''
                        heater_val = ''
                        panel_sn = default_sn

                        status_idx = -1
                        for idx, tok in enumerate(tokens):
                            if tok.upper() in ['CLOSE', 'TRIP', 'OPEN']:
                                status = tok.upper()
                                status_idx = idx
                                break

                        if status_idx != -1:
                            name_val = ' '.join(tokens[:status_idx]).strip()
                            after_tokens = tokens[status_idx + 1:]
                            if len(after_tokens) >= 1:
                                load_val = clean_val(after_tokens[0])
                            if len(after_tokens) >= 2:
                                cable_val = clean_val(after_tokens[1])
                            if len(after_tokens) >= 3:
                                heater_val = format_heater_amp(after_tokens[2])
                            if len(after_tokens) >= 4:
                                panel_sn = clean_sn_val(' '.join(after_tokens[3:])) or default_sn
                        else:
                            name_val = rest

                        if link_no in ['-', 'N/A', ''] and (not name_val or all(c in '- ' for c in name_val)):
                            continue

                        name_upper = name_val.upper().strip()
                        if is_vcb and "TRANSITION" in name_upper:
                            cable_val = "-"
                            link_no = "-"
                        elif "BUS SECTION" in name_upper:
                            cable_val = "-"

                        if not is_vcb:
                            panel_sn = default_sn

                        all_panels.append({
                            "description": p_num,
                            "link_number": link_no,
                            "name": name_val,
                            "cable_type": cable_val,
                            "breaker_status": status,
                            "busbar_position": "-",
                            "load_amp": load_val,
                            "heater_amp": heater_val,
                            "serial_number": panel_sn
                        })

        fields["panels"] = all_panels
        fields["panels_1"] = all_panels

        # Enforce FP / LVDB convention types based on TX setup
        has_tx1 = bool(fields.get("tx1_manufacturer") or fields.get("tx1_serial") or fields.get("tx1_rating"))
        has_tx2 = bool(fields.get("tx2_manufacturer") or fields.get("tx2_serial") or fields.get("tx2_rating"))
        fp1_d = {"manufacturer": fields.get("lv1_manufacturer"), "model": fields.get("lv1_model"), "serial": fields.get("lv1_serial"), "rating": fields.get("lv1_rating"), "name": fields.get("lv1_name"), "type": fields.get("lv1_type")}
        fp2_d = {"manufacturer": fields.get("lv2_manufacturer"), "model": fields.get("lv2_model"), "serial": fields.get("lv2_serial"), "rating": fields.get("lv2_rating"), "name": fields.get("lv2_name"), "type": fields.get("lv2_type")}
        c_t1, c_t2 = determine_fp_lvdb_types(has_tx1, has_tx2, fp1_d, fp2_d)
        if c_t1: fields["lv1_type"] = c_t1
        if c_t2: fields["lv2_type"] = c_t2

        doc.close()
    except Exception as e:
        print(f"  [PDF Extraction Error] {e}", flush=True)

    return fields


# ── 1. Raw Material master list (subfolders) ───────────────────────────────────
start_overall_time = time.time()
print("=" * 60, flush=True)
print("  STEP 1: SCANNING WORKSPACE DIRECTORIES", flush=True)
print("=" * 60, flush=True)
print(f"  Workspace: {BASE}", flush=True)
print("  Scanning Raw Material folders (03. RAW MATERIAL)...", flush=True)
rm_entries = {}
if os.path.exists(RM_DIR):
    for entry in os.listdir(RM_DIR):
        if os.path.isdir(os.path.join(RM_DIR, entry)):
            num = extract_num(entry)
            if num is not None:
                rm_entries[num] = clean_substation_name(entry)
print(f"  -> Found {len(rm_entries)} Raw Material substation folders.", flush=True)

# ── 2. Scan file folders ───────────────────────────────────────────────────────
print("  Scanning Full Report folder (01. FULL REPORT)...", flush=True)
fr_set = scan_full_report(FR_DIR)
print(f"  -> Found {len(fr_set)} completed Full Reports.", flush=True)

print("  Scanning Quick Report folder (02. QUICK REPORT)...", flush=True)
qr_set = scan_files(QR_DIR)
print(f"  -> Found {len(qr_set)} Quick Reports.", flush=True)

# ── 3. Read testsheet files (.xlsx 1st choice, .pdf 2nd choice fallback) ───────
scope_arg = "ALL"
format_arg = os.environ.get("TESTSHEET_FORMAT", "AUTO").upper().strip()

for a in sys.argv[1:]:
    a_clean = a.strip()
    if a_clean.upper() in ["EET", "TIANU", "AUTO"]:
        format_arg = a_clean.upper()
    elif a_clean.upper().startswith("--FORMAT="):
        format_arg = a_clean.split("=", 1)[1].strip().upper()
    elif a_clean.lower() == "--format" and sys.argv.index(a) + 1 < len(sys.argv):
        format_arg = sys.argv[sys.argv.index(a) + 1].strip().upper()
    elif not a_clean.startswith("-"):
        scope_arg = a_clean

if format_arg in ["EET", "TIANU"]:
    os.environ["TESTSHEET_FORMAT"] = format_arg

matches_scope, scope_label = parse_substation_filter(scope_arg)

print("=" * 60, flush=True)
print(f"  STEP 2: EXTRACTING TESTSHEET SPECIFICATIONS (04. TESTSHEET)", flush=True)
print(f"  TARGET SCOPE : {scope_label.upper()}", flush=True)
if format_arg in ["EET", "TIANU"]:
    print(f"  FORMAT ENGINE: {format_arg} (Manual Selection)", flush=True)
else:
    print(f"  FORMAT ENGINE: AUTO-DETECT (EET & TIANU Dual Engine)", flush=True)
print("=" * 60, flush=True)
ts_data = {}
ts_files = {}

xlsx_files = {}
pdf_files = {}

if os.path.exists(XLSX_DIR):
    for f in os.listdir(XLSX_DIR):
        if f.startswith('~'):
            continue
        f_lower = f.lower()
        num = extract_num(f)
        if num is not None:
            if f_lower.endswith('.xlsx'):
                xlsx_files[num] = os.path.join(XLSX_DIR, f)
            elif f_lower.endswith('.pdf'):
                pdf_files[num] = os.path.join(XLSX_DIR, f)

# Hierarchy: 1st choice = .xlsx, 2nd choice (fallback) = .pdf
all_candidate_nums = sorted(set(xlsx_files.keys()) | set(pdf_files.keys()))
for num in all_candidate_nums:
    if num in xlsx_files:
        ts_files[num] = (xlsx_files[num], 'xlsx')
    elif num in pdf_files:
        ts_files[num] = (pdf_files[num], 'pdf')

# Filter candidate testsheets according to user scope (exact, range, or ALL)
if matches_scope:
    all_ts_nums = sorted([n for n in ts_files.keys() if matches_scope(n)])
else:
    all_ts_nums = sorted(ts_files.keys())

total_ts = len(all_ts_nums)
xlsx_count = sum(1 for n in all_ts_nums if ts_files[n][1] == 'xlsx')
pdf_count = sum(1 for n in all_ts_nums if ts_files[n][1] == 'pdf')
print(f"  Found {total_ts} testsheet(s) matching scope [{scope_label}]: {xlsx_count} Excel (.xlsx) [1st Choice], {pdf_count} PDF (.pdf) [Fallback].", flush=True)

step_interval = 1 if total_ts <= 15 else (2 if total_ts <= 30 else 5)

for idx, num in enumerate(all_ts_nums, start=1):
    f_path, f_type = ts_files[num]
    if f_type == 'xlsx':
        ts_fields = read_ts_fields_excel(f_path, sub_num=num)
    else:
        ts_fields = read_ts_fields_pdf(f_path)
    ts_data[num] = ts_fields
    if idx % step_interval == 0 or idx == total_ts or idx == 1:
        sg_m = ts_fields.get("sg_manufacturer") or "-"
        sg_t = ts_fields.get("sg_type") or "-"
        sub_name = ts_fields.get("pe_name_erms") or ts_fields.get("pe_name_site") or rm_entries.get(num, "")
        percent = (idx / total_ts) * 100 if total_ts else 100
        fmt_label = ts_fields.get("_format", "XLSX")
        src_tag = f"[{fmt_label}]" if f_type == 'xlsx' else "[PDF]"
        print(f"  [{idx:03d}/{total_ts:03d} - {percent:5.1f}%] {src_tag} Substation {num} ({sub_name}) | SG: {sg_m} ({sg_t})", flush=True)

existing_excel_dict = {}
existing_panels_dict = {}
backup_data_rows = None

if os.path.exists(OUT):
    try:
        old_wb = openpyxl.load_workbook(OUT, data_only=True)
        # 1. Read 'Substation Checklist' first (active user working sheet)
        if 'Substation Checklist' in old_wb.sheetnames:
            ws_sc = old_wb['Substation Checklist']
            cols_sc = [ws_sc.cell(4, c).value for c in range(1, ws_sc.max_column + 1)]
            for r in range(5, ws_sc.max_row + 1):
                num_val = ws_sc.cell(r, 1).value
                if num_val is not None:
                    k_str = str(num_val).strip()
                    if k_str and k_str.lower() != 'number':
                        r_dict = {}
                        for c_idx, col_name in enumerate(cols_sc, start=1):
                            if col_name:
                                r_dict[col_name] = ws_sc.cell(r, c_idx).value
                        existing_excel_dict[k_str] = r_dict

        # 2. Read 'Substation Info' for GPS, ERMS/Site names, F/L No
        if 'Substation Info' in old_wb.sheetnames:
            ws_info = old_wb['Substation Info']
            cols_info = [ws_info.cell(3, c).value for c in range(1, ws_info.max_column + 1)]
            for r in range(4, ws_info.max_row + 1):
                num_val = ws_info.cell(r, 1).value
                if num_val is not None:
                    k_str = str(num_val).strip()
                    if k_str and k_str.lower() != 'no':
                        if k_str not in existing_excel_dict:
                            existing_excel_dict[k_str] = {}
                        for c_idx, col_name in enumerate(cols_info, start=1):
                            if col_name and ws_info.cell(r, c_idx).value is not None:
                                existing_excel_dict[k_str][col_name] = ws_info.cell(r, c_idx).value

        # 3. Read 'Backup' as fallback
        if 'Backup' in old_wb.sheetnames:
            ws_bk = old_wb['Backup']
            cols_bk = [ws_bk.cell(4, c).value for c in range(1, ws_bk.max_column + 1)]
            for r in range(5, ws_bk.max_row + 1):
                num_val = ws_bk.cell(r, 1).value
                if num_val is not None:
                    k_str = str(num_val).strip()
                    if k_str and k_str.lower() != 'number':
                        if k_str not in existing_excel_dict:
                            existing_excel_dict[k_str] = {}
                        for c_idx, col_name in enumerate(cols_bk, start=1):
                            if col_name and not existing_excel_dict[k_str].get(col_name):
                                existing_excel_dict[k_str][col_name] = ws_bk.cell(r, c_idx).value

        # 4. Read 'Panel Schedule'
        if 'Panel Schedule' in old_wb.sheetnames:
            ws_p = old_wb['Panel Schedule']
            for r in range(2, ws_p.max_row + 1):
                num_val = ws_p.cell(r, 1).value
                if num_val is not None:
                    k_str = str(num_val).strip()
                    if k_str and k_str.lower() != 'substation no':
                        p_row = {
                            "description": ws_p.cell(r, 3).value,
                            "link_number": ws_p.cell(r, 4).value,
                            "name": ws_p.cell(r, 5).value,
                            "cable_type": ws_p.cell(r, 6).value,
                            "breaker_status": ws_p.cell(r, 7).value,
                            "busbar_position": ws_p.cell(r, 8).value,
                            "load_amp": ws_p.cell(r, 9).value,
                            "heater_amp": ws_p.cell(r, 10).value,
                            "serial_number": ws_p.cell(r, 11).value,
                        }
                        if k_str not in existing_panels_dict:
                            existing_panels_dict[k_str] = []
                        existing_panels_dict[k_str].append(p_row)

        old_wb.close()
    except Exception as e:
        print(f"  [WARNING] Could not read existing Excel data: {e}", flush=True)

backup_data_dict = existing_excel_dict

def is_blank(v):
    if v is None:
        return True
    s = str(v).strip()
    return s in ["", "-", "N/A", "n/a", "None", "none", "nan", "NAN", "null", "NULL", ".", "/"]

def normalize_compare(v):
    if is_blank(v):
        return ""
    s = str(v).strip().lower()
    try:
        f = float(s)
        if f.is_integer():
            return str(int(f))
        return f"{f:.1f}"
    except (ValueError, TypeError):
        pass
    return s

def values_differ(v_existing, v_extracted):
    if is_blank(v_existing):
        return False
    if is_blank(v_extracted):
        return True
    return normalize_compare(v_existing) != normalize_compare(v_extracted)

preserved_edits_count = 0

def resolve_field(v_existing, v_extracted, field_name="", sub_num=""):
    """
    If data in existing Excel is not same with extracted data,
    it means it has already been edited, so leave that part (preserve existing Excel data).
    Only use extracted data if existing Excel cell is empty/blank/placeholder.
    """
    global preserved_edits_count
    if is_blank(v_existing):
        return v_extracted if not is_blank(v_extracted) else ""
    if field_name in ["FP1 Type", "FP2 Type"] and str(v_existing).strip().upper() in ["XLPE", "PILC", "PVC"]:
        return v_extracted if not is_blank(v_extracted) else ""
    if values_differ(v_existing, v_extracted):
        preserved_edits_count += 1
        return v_existing
    return v_existing

existing_keys_converted = set()
for k in existing_excel_dict.keys():
    m_k = re.match(r'^0*(\d+)', str(k).strip())
    if m_k:
        try:
            existing_keys_converted.add(int(m_k.group(1)))
        except ValueError:
            existing_keys_converted.add(k)
    else:
        existing_keys_converted.add(k)

all_nums = sorted(
    set(rm_entries.keys()) | fr_set | qr_set | set(ts_data.keys()) | existing_keys_converted,
    key=lambda x: (int(re.match(r'^\d+', str(x)).group(0)) if re.match(r'^\d+', str(x)) else 99999, str(x))
)


def guess_name(num):
    for folder in [XLSX_DIR, QR_DIR, FR_DIR]:
        if os.path.exists(folder):
            for f in os.listdir(folder):
                if extract_num(f) == num:
                    c_name = clean_substation_name(f)
                    if is_valid_substation_name(c_name):
                        return c_name
    return ""


# ── 4. Build rows ──────────────────────────────────────────────────────────────
print("=" * 60, flush=True)
print("  STEP 3: COMPILING MASTER CHECKLIST & PANEL SCHEDULES", flush=True)
print("=" * 60, flush=True)
print(f"  Compiling data rows for {len(all_nums)} total unique substations...", flush=True)
rows = []
for num in all_nums:
    bk = existing_excel_dict.get(str(num), {})
    f = ts_data.get(num, {})
    is_in_scope = True if matches_scope is None else matches_scope(num)

    if not is_in_scope and bk:
        # OUTSIDE SCOPE: 100% PRESERVE EXISTING EXCEL ROW (Never touch anything!)
        name = bk.get("SUBSTATION NAME") or f"Substation {num}"
        sub_type = bk.get("SUBSTATION TYPE", "")
        time_val = bk.get("Time", "")
        date = bk.get("Date", "")
        ambient = bk.get("Ambient Temp", "")
        humidity = bk.get("Humidity", "")
        sg_manuf = bk.get("Switchgear Manufacturer", "")
        sg_type = bk.get("Switchgear Type", "")
        sg_rating = bk.get("Switchgear Rating", "")
        sg_model = bk.get("Switchgear Model", "")
        tx1_manuf = bk.get("TX1 Manufacturer", "")
        tx1_model = bk.get("TX1 Model", "")
        tx1_rating = bk.get("TX1 Rating", "")
        tx1_serial = bk.get("TX1 Serial Number", "")
        tx2_manuf = bk.get("TX2 Manufacturer", "")
        tx2_model = bk.get("TX2 Model", "")
        tx2_rating = bk.get("TX2 Rating", "")
        tx2_serial = bk.get("TX2 Serial Number", "")
        lv1_manuf = bk.get("FP1 Manufacturer", "")
        lv1_model = ""  # User rule: FP1 & FP2 model always left blank
        lv1_rating = bk.get("FP1 Rating", "")
        lv1_serial = bk.get("FP1 Serial No", "")
        lv1_type = bk.get("FP1 Type", "")
        lv2_manuf = bk.get("FP2 Manufacturer", "")
        lv2_model = ""  # User rule: FP1 & FP2 model always left blank
        lv2_rating = bk.get("FP2 Rating", "")
        lv2_serial = bk.get("FP2 Serial No", "")
        lv2_type = bk.get("FP2 Type", "")
        bat1_manuf = bk.get("Battery Bank 1 Manufactrurer", "")
        bat1_model = bk.get("Battery Bank 1 Model", "")
        bat1_serial = bk.get("Battery Bank 1 Serial Number", "")
        bat2_manuf = bk.get("Battery Bank 2 Manufactrurer", "")
        bat2_model = bk.get("Battery Bank 2 Model", "")
        bat2_serial = bk.get("Battery Bank 2 Serial Number", "")
        bat3_manuf = bk.get("Battery Bank 3 Manufactrurer", "")
        bat3_model = bk.get("Battery Bank 3 Model", "")
        bat3_serial = bk.get("Battery Bank 3 Serial Number", "")
        bat4_manuf = bk.get("Battery Bank 4 Manufactrurer", "")
        bat4_model = bk.get("Battery Bank 4 Model", "")
        bat4_serial = bk.get("Battery Bank 4 Serial Number", "")
        has_fr = "YES" if num in fr_set else "NO"
        has_qr = "YES" if num in qr_set else "NO"
        has_rm = "YES" if num in rm_entries else "NO"
        has_ts = "YES" if num in ts_files or (isinstance(bk.get("TEST SHEET"), str) and bk.get("TEST SHEET").upper() == "YES") else "NO"
        vi_date = bk.get("DATE", date)
        gps_coord = bk.get("GPS COORDINATES", "")
        pe_name_erms = bk.get("SUBSTATION NAME (ERMS)", "")
        pe_name_site = bk.get("SUBSTATION NAME (SITE)", "")
        fl_no_erms = bk.get("F/L NO (ERMS)", "")
        fl_no_site = bk.get("F/L NO (SITE)", "")

        rows.append((num, name, sub_type, time_val, date, ambient, humidity,
                     sg_manuf, sg_type, sg_rating, sg_model,
                     tx1_manuf, tx1_model, tx1_rating, tx1_serial,
                     tx2_manuf, tx2_model, tx2_rating, tx2_serial,
                     lv1_manuf, lv1_model, lv1_rating, lv1_serial, lv1_type,
                     lv2_manuf, lv2_model, lv2_rating, lv2_serial, lv2_type,
                     bat1_manuf, bat1_model, bat1_serial,
                     bat2_manuf, bat2_model, bat2_serial,
                     bat3_manuf, bat3_model, bat3_serial,
                     bat4_manuf, bat4_model, bat4_serial,
                     has_fr, has_qr, has_rm, has_ts,
                     vi_date, gps_coord, pe_name_erms, pe_name_site,
                     fl_no_erms, fl_no_site))
        continue

    raw_rm = rm_entries.get(num)
    rm_name = raw_rm if is_valid_substation_name(raw_rm) else None

    ts_pe_erms = f.get("pe_name_erms")
    ts_pe_site = f.get("pe_name_site")
    valid_ts_name = ts_pe_erms if is_valid_substation_name(ts_pe_erms) else (ts_pe_site if is_valid_substation_name(ts_pe_site) else None)

    guessed = guess_name(num)
    guessed_name = guessed if is_valid_substation_name(guessed) else None

    extracted_name = rm_name or valid_ts_name or guessed_name or f"Substation {num}"

    bk_raw = bk.get("SUBSTATION NAME")
    if is_valid_substation_name(bk_raw) and not str(bk_raw).strip().lower().startswith("substation "):
        if values_differ(bk_raw, extracted_name):
            preserved_edits_count += 1
            name = str(bk_raw).strip()
        else:
            name = str(bk_raw).strip()
    else:
        name = extracted_name

    date = resolve_field(bk.get("Date"), f.get("date"), "Date", num)
    ambient_ext = format_temp(f.get("ambient")) if f.get("ambient") else ""
    ambient = resolve_field(bk.get("Ambient Temp"), ambient_ext, "Ambient Temp", num)
    humidity = resolve_field(bk.get("Humidity"), f.get("humidity"), "Humidity", num)
    time_in = resolve_field(bk.get("Time"), f.get("time_in"), "Time", num)
    sub_type = resolve_field(bk.get("SUBSTATION TYPE"), f.get("substation_type"), "SUBSTATION TYPE", num)

    pe_name_erms = resolve_field(bk.get("SUBSTATION NAME (ERMS)"), f.get("pe_name_erms"), "SUBSTATION NAME (ERMS)", num)
    pe_name_site = resolve_field(bk.get("SUBSTATION NAME (SITE)"), f.get("pe_name_site"), "SUBSTATION NAME (SITE)", num)
    fl_no_erms = resolve_field(bk.get("F/L NO (ERMS)"), f.get("fl_no_erms"), "F/L NO (ERMS)", num)
    fl_no_site = resolve_field(bk.get("F/L NO (SITE)"), f.get("fl_no_site"), "F/L NO (SITE)", num)

    # Rule: If F/L Number is blank, None, or dash, change to N/A (do not assume or fallback)
    if not fl_no_erms or is_blank(fl_no_erms) or fl_no_erms in ['0', '0.0']:
        fl_no_erms = "N/A"
    if not fl_no_site or is_blank(fl_no_site) or fl_no_site in ['0', '0.0']:
        fl_no_site = "N/A"

    gps_coord = resolve_field(bk.get("GPS COORDINATES"), f.get("gps_coordinate"), "GPS COORDINATES", num)
    vi_date = resolve_field(bk.get("DATE"), f.get("vi_date"), "DATE", num) or date

    sg_manuf_ext = f.get("sg_manufacturer")
    if not sg_manuf_ext or is_blank(sg_manuf_ext) or not any(c.isalnum() for c in str(sg_manuf_ext)):
        detected_mfr = detect_sg_manufacturer_from_dg(num, RM_DIR)
        if detected_mfr:
            sg_manuf_ext = detected_mfr
    sg_manuf = resolve_field(bk.get("Switchgear Manufacturer"), sg_manuf_ext, "Switchgear Manufacturer", num)

    sg_model = resolve_field(bk.get("Switchgear Model"), f.get("sg_model"), "Switchgear Model", num)
    sg_rating = resolve_field(bk.get("Switchgear Rating"), f.get("sg_rating"), "Switchgear Rating", num)
    sg_serial = resolve_field(bk.get("Switchgear Serial"), f.get("sg_serial"), "Switchgear Serial", num)
    sg_type = resolve_field(bk.get("Switchgear Type"), f.get("sg_type"), "Switchgear Type", num)

    sg_manufs_list = f.get("sg_manufs_list", [])
    sg_models_list = f.get("sg_models_list", [])
    sg_ratings_list = f.get("sg_ratings_list", [])
    sg_serials_list = f.get("sg_serials_list", [])

    is_split = False
    if len(sg_manufs_list) > 1 and sg_manufs_list[0] and sg_manufs_list[1]:
        if sg_manufs_list[0].upper() != sg_manufs_list[1].upper():
            is_split = True
    elif len(sg_serials_list) > 1 and sg_serials_list[0] and sg_serials_list[1]:
        if sg_serials_list[0].upper() != sg_serials_list[1].upper():
            is_split = True

    tx1_manuf = resolve_field(bk.get("TX1 Manufacturer"), f.get("tx1_manufacturer"), "TX1 Manufacturer", num)
    tx1_model = resolve_field(bk.get("TX1 Model"), f.get("tx1_model"), "TX1 Model", num)
    tx1_rating = resolve_field(bk.get("TX1 Rating"), f.get("tx1_rating"), "TX1 Rating", num)
    tx1_serial = resolve_field(bk.get("TX1 Serial Number"), f.get("tx1_serial"), "TX1 Serial Number", num)

    tx2_manuf = resolve_field(bk.get("TX2 Manufacturer"), f.get("tx2_manufacturer"), "TX2 Manufacturer", num)
    tx2_model = resolve_field(bk.get("TX2 Model"), f.get("tx2_model"), "TX2 Model", num)
    tx2_rating = resolve_field(bk.get("TX2 Rating"), f.get("tx2_rating"), "TX2 Rating", num)
    tx2_serial = resolve_field(bk.get("TX2 Serial Number"), f.get("tx2_serial"), "TX2 Serial Number", num)

    lv1_manuf = resolve_field(bk.get("FP1 Manufacturer"), f.get("lv1_manufacturer"), "FP1 Manufacturer", num)
    lv1_model = ""  # User rule: FP1 & FP2 model always left blank
    lv1_rating = resolve_field(bk.get("FP1 Rating"), f.get("lv1_rating"), "FP1 Rating", num)
    lv1_serial = resolve_field(bk.get("FP1 Serial No"), f.get("lv1_serial"), "FP1 Serial No", num)

    lv2_manuf = resolve_field(bk.get("FP2 Manufacturer"), f.get("lv2_manufacturer"), "FP2 Manufacturer", num)
    lv2_model = ""  # User rule: FP1 & FP2 model always left blank
    lv2_rating = resolve_field(bk.get("FP2 Rating"), f.get("lv2_rating"), "FP2 Rating", num)
    lv2_serial = resolve_field(bk.get("FP2 Serial No"), f.get("lv2_serial"), "FP2 Serial No", num)

    # Determine proper FP / LVDB convention names based on TX setup
    has_tx1_live = bool(tx1_manuf or tx1_serial or tx1_rating or tx1_model)
    has_tx2_live = bool(tx2_manuf or tx2_serial or tx2_rating or tx2_model)
    fp1_info = {
        'manufacturer': lv1_manuf, 'model': lv1_model, 'rating': lv1_rating,
        'serial': lv1_serial, 'type': f.get('lv1_type'), 'name': f.get('lv1_name')
    }
    fp2_info = {
        'manufacturer': lv2_manuf, 'model': lv2_model, 'rating': lv2_rating,
        'serial': lv2_serial, 'type': f.get('lv2_type'), 'name': f.get('lv2_name')
    }
    conv_fp1, conv_fp2 = determine_fp_lvdb_types(has_tx1_live, has_tx2_live, fp1_info, fp2_info)

    bk_fp1 = str(bk.get("FP1 Type") or "").strip()
    if is_blank(bk_fp1) or bk_fp1.upper() in ['XLPE', 'PILC', 'PVC']:
        lv1_type = conv_fp1
    else:
        lv1_type = resolve_field(bk.get("FP1 Type"), conv_fp1, "FP1 Type", num)

    bk_fp2 = str(bk.get("FP2 Type") or "").strip()
    if is_blank(bk_fp2) or bk_fp2.upper() in ['XLPE', 'PILC', 'PVC']:
        lv2_type = conv_fp2
    else:
        lv2_type = resolve_field(bk.get("FP2 Type"), conv_fp2, "FP2 Type", num)

    bat1_manuf = resolve_field(bk.get("Battery Bank 1 Manufactrurer"), f.get("bat1_manufacturer"), "Battery Bank 1 Manufactrurer", num)
    bat1_model = resolve_field(bk.get("Battery Bank 1 Model"), f.get("bat1_model"), "Battery Bank 1 Model", num)
    bat1_serial = resolve_field(bk.get("Battery Bank 1 Serial Number"), f.get("bat1_serial"), "Battery Bank 1 Serial Number", num)

    bat2_manuf = resolve_field(bk.get("Battery Bank 2 Manufactrurer"), f.get("bat2_manufacturer"), "Battery Bank 2 Manufactrurer", num)
    bat2_model = resolve_field(bk.get("Battery Bank 2 Model"), f.get("bat2_model"), "Battery Bank 2 Model", num)
    bat2_serial = resolve_field(bk.get("Battery Bank 2 Serial Number"), f.get("bat2_serial"), "Battery Bank 2 Serial Number", num)

    bat3_manuf = resolve_field(bk.get("Battery Bank 3 Manufactrurer"), f.get("bat3_manufacturer"), "Battery Bank 3 Manufactrurer", num)
    bat3_model = resolve_field(bk.get("Battery Bank 3 Model"), f.get("bat3_model"), "Battery Bank 3 Model", num)
    bat3_serial = resolve_field(bk.get("Battery Bank 3 Serial Number"), f.get("bat3_serial"), "Battery Bank 3 Serial Number", num)

    bat4_manuf = resolve_field(bk.get("Battery Bank 4 Manufactrurer"), f.get("bat4_manufacturer"), "Battery Bank 4 Manufactrurer", num)
    bat4_model = resolve_field(bk.get("Battery Bank 4 Model"), f.get("bat4_model"), "Battery Bank 4 Model", num)
    bat4_serial = resolve_field(bk.get("Battery Bank 4 Serial Number"), f.get("bat4_serial"), "Battery Bank 4 Serial Number", num)

    time_val = time_in
    has_fr = "YES" if num in fr_set else "NO"
    has_qr = "YES" if num in qr_set else "NO"
    has_rm = "YES" if num in rm_entries else "NO"
    has_ts = "YES" if num in ts_data else "NO"

    if is_split:
        sg_types_list = f.get("sg_types_list", [])

        # Switchgear 1 / Row A (Gets SWG 1 details + ALL ancillary data: TX1/2, LVDB, Batteries)
        num_a = f"{num}A"
        bk_a = existing_excel_dict.get(num_a, {})
        name_a = bk_a.get("SUBSTATION NAME") or (f"{name} (SWG 1)" if name else f"Substation {num}A")
        sg_mfr_a = resolve_field(bk_a.get("Switchgear Manufacturer"), sg_manufs_list[0] if len(sg_manufs_list) > 0 else sg_manuf, "Switchgear Manufacturer", num_a)
        sg_mdl_a = resolve_field(bk_a.get("Switchgear Model"), sg_models_list[0] if len(sg_models_list) > 0 else sg_model, "Switchgear Model", num_a)
        sg_rtg_a = resolve_field(bk_a.get("Switchgear Rating"), sg_ratings_list[0] if len(sg_ratings_list) > 0 else sg_rating, "Switchgear Rating", num_a)
        sg_typ_a = resolve_field(bk_a.get("Switchgear Type"), sg_types_list[0] if len(sg_types_list) > 0 and sg_types_list[0] else sg_type, "Switchgear Type", num_a)

        rows.append((num_a, name_a, sub_type, time_val, date, ambient, humidity,
                     sg_mfr_a, sg_typ_a, sg_rtg_a, sg_mdl_a,
                     tx1_manuf, tx1_model, tx1_rating, tx1_serial,
                     tx2_manuf, tx2_model, tx2_rating, tx2_serial,
                     lv1_manuf, lv1_model, lv1_rating, lv1_serial, lv1_type,
                     lv2_manuf, lv2_model, lv2_rating, lv2_serial, lv2_type,
                     bat1_manuf, bat1_model, bat1_serial,
                     bat2_manuf, bat2_model, bat2_serial,
                     bat3_manuf, bat3_model, bat3_serial,
                     bat4_manuf, bat4_model, bat4_serial,
                     has_fr, has_qr, has_rm, has_ts,
                     vi_date, gps_coord, pe_name_erms, pe_name_site,
                     fl_no_erms, fl_no_site))

        # Switchgear 2 / Row B (Gets SWG 2 details; ancillary fields left empty as they belong to 113A)
        num_b = f"{num}B"
        bk_b = existing_excel_dict.get(num_b, {})
        name_b = bk_b.get("SUBSTATION NAME") or (f"{name} (SWG 2)" if name else f"Substation {num}B")
        sg_mfr_b = resolve_field(bk_b.get("Switchgear Manufacturer"), sg_manufs_list[1] if len(sg_manufs_list) > 1 else sg_manuf, "Switchgear Manufacturer", num_b)
        sg_mdl_b = resolve_field(bk_b.get("Switchgear Model"), sg_models_list[1] if len(sg_models_list) > 1 else sg_model, "Switchgear Model", num_b)
        sg_rtg_b = resolve_field(bk_b.get("Switchgear Rating"), sg_ratings_list[1] if len(sg_ratings_list) > 1 else sg_rating, "Switchgear Rating", num_b)
        sg_typ_b = resolve_field(bk_b.get("Switchgear Type"), sg_types_list[1] if len(sg_types_list) > 1 and sg_types_list[1] else sg_type, "Switchgear Type", num_b)

        rows.append((num_b, name_b, sub_type, "", "", "", "",
                     sg_mfr_b, sg_typ_b, sg_rtg_b, sg_mdl_b,
                     "", "", "", "",
                     "", "", "", "",
                     "", "", "", "", "",
                     "", "", "", "", "",
                     "", "", "",
                     "", "", "",
                     "", "", "",
                     "", "", "",
                     has_fr, has_qr, has_rm, has_ts,
                     vi_date, gps_coord, pe_name_erms, pe_name_site,
                     fl_no_erms, fl_no_site))
    else:
        rows.append((num, name, sub_type, time_val, date, ambient, humidity,
                     sg_manuf, sg_type, sg_rating, sg_model,
                     tx1_manuf, tx1_model, tx1_rating, tx1_serial,
                     tx2_manuf, tx2_model, tx2_rating, tx2_serial,
                     lv1_manuf, lv1_model, lv1_rating, lv1_serial, lv1_type,
                     lv2_manuf, lv2_model, lv2_rating, lv2_serial, lv2_type,
                     bat1_manuf, bat1_model, bat1_serial,
                     bat2_manuf, bat2_model, bat2_serial,
                     bat3_manuf, bat3_model, bat3_serial,
                     bat4_manuf, bat4_model, bat4_serial,
                     has_fr, has_qr, has_rm, has_ts,
                     vi_date, gps_coord, pe_name_erms, pe_name_site,
                     fl_no_erms, fl_no_site))

print(f"  [SMART MERGE] Successfully preserved {preserved_edits_count} manually edited field(s) from existing Excel!", flush=True)

# ── 5. Create Master Checklist Workbook ────────────────────────────────────────
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Substation Checklist"

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
YES_FONT = Font(name="Calibri", bold=True, color="1E8449", size=11)
NO_FONT = Font(name="Calibri", bold=True, color="C0392B", size=11)
YES_FILL = PatternFill("solid", fgColor="D5F5E3")
NO_FILL = PatternFill("solid", fgColor="FADBD8")
ALT_FILL = PatternFill("solid", fgColor="D6E4F0")
PLAIN_FILL = PatternFill("solid", fgColor="FFFFFF")
INFO_FONT = Font(name="Calibri", size=10, color="1A5276")
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
thin = Side(style="thin", color="B0B0B0")
thick = Side(style="medium", color="1F4E79")
BORDER_INNER = Border(left=thin, right=thin, top=thin, bottom=thin)
BORDER_HEAD = Border(left=thick, right=thick, top=thick, bottom=thick)

po_title = os.path.basename(BASE)

col_widths = [13, 40, 20, 15, 13, 13, 12, 20, 15, 12, 15, 20, 15, 12, 18, 20, 15, 12, 18, 20, 15, 12, 18, 15, 20, 15,
              12, 18, 15, 25, 20, 18, 25, 20, 18, 25, 20, 18, 25, 20, 18, 13, 13, 13, 13]
for i, w in enumerate(col_widths, start=1):
    col_letter = openpyxl.utils.get_column_letter(i)
    ws.column_dimensions[col_letter].width = w

ws.row_dimensions[1].height = 32
ws.merge_cells("A1:AS1")
t = ws["A1"]
t.value = po_title + "  |  Substation Document Checklist"
t.font = Font(name="Calibri", bold=True, size=13, color="FFFFFF")
t.fill = PatternFill("solid", fgColor="0B2D6B")
t.alignment = Alignment(horizontal="center", vertical="center")

ws.row_dimensions[2].height = 18
ws.merge_cells("A2:AS2")
g = ws["A2"]
g.value = "Generated on: " + datetime.now().strftime("%d/%m/%Y  %H:%M")
g.font = Font(name="Calibri", italic=True, size=9, color="7F8C8D")
g.fill = PatternFill("solid", fgColor="EBF5FB")
g.alignment = Alignment(horizontal="right", vertical="center")

ws.row_dimensions[3].height = 16
ws.merge_cells("A3:G3")
ws.merge_cells("H3:K3")
ws.merge_cells("L3:O3")
ws.merge_cells("P3:S3")
ws.merge_cells("T3:X3")
ws.merge_cells("Y3:AC3")
ws.merge_cells("AD3:AF3")
ws.merge_cells("AG3:AI3")
ws.merge_cells("AJ3:AL3")
ws.merge_cells("AM3:AO3")
ws.merge_cells("AP3:AS3")
grp_labels = [
    ("A3", "SUBSTATION INFO", "1B4F72"), ("H3", "SWITCHGEAR DETAILS", "2E4053"),
    ("L3", "TRANSFORMER 1 DETAILS", "283747"), ("P3", "TRANSFORMER 2 DETAILS", "34495E"),
    ("T3", "LVDB / FEEDER PILLAR 1 DETAILS", "1B4F72"), ("Y3", "LVDB / FEEDER PILLAR 2 DETAILS", "1F618D"),
    ("AD3", "BATTERY BANK 1 DETAILS", "2E4053"), ("AG3", "BATTERY BANK 2 DETAILS", "283747"),
    ("AJ3", "BATTERY BANK 3 DETAILS", "34495E"), ("AM3", "BATTERY BANK 4 DETAILS", "1B4F72"),
    ("AP3", "DOCUMENT STATUS", "117A65")
]
for cell_ref, label, color in grp_labels:
    c = ws[cell_ref]
    c.value = label
    c.font = Font(name="Calibri", bold=True, size=9, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=color)
    c.alignment = CENTER

headers = [
    "NUMBER", "SUBSTATION NAME", "SUBSTATION TYPE", "Time", "Date", "Ambient Temp", "Humidity",
    "Switchgear Manufacturer", "Switchgear Type", "Switchgear Rating", "Switchgear Model",
    "TX1 Manufacturer", "TX1 Model", "TX1 Rating", "TX1 Serial Number",
    "TX2 Manufacturer", "TX2 Model", "TX2 Rating", "TX2 Serial Number",
    "FP1 Manufacturer", "FP1 Model", "FP1 Rating", "FP1 Serial No", "FP1 Type",
    "FP2 Manufacturer", "FP2 Model", "FP2 Rating", "FP2 Serial No", "FP2 Type",
    "Battery Bank 1 Manufactrurer", "Battery Bank 1 Model", "Battery Bank 1 Serial Number",
    "Battery Bank 2 Manufactrurer", "Battery Bank 2 Model", "Battery Bank 2 Serial Number",
    "Battery Bank 3 Manufactrurer", "Battery Bank 3 Model", "Battery Bank 3 Serial Number",
    "Battery Bank 4 Manufactrurer", "Battery Bank 4 Model", "Battery Bank 4 Serial Number",
    "FULL REPORT", "QUICK REPORT", "RAW MATERIAL", "TEST SHEET"
]
ws.row_dimensions[4].height = 28
for col, hdr in enumerate(headers, start=1):
    c = ws.cell(row=4, column=col, value=hdr)
    c.font = HEADER_FONT
    c.fill = HEADER_FILL
    c.alignment = CENTER
    c.border = BORDER_HEAD

for i, row_data in enumerate(rows):
    r = i + 5
    ws.row_dimensions[r].height = 20
    row_fill = ALT_FILL if i % 2 == 0 else PLAIN_FILL
    num, name, sub_type, time_val, date, ambient, humidity = row_data[:7]
    sg_manuf, sg_type, sg_rating, sg_model = row_data[7:11]
    tx1_manuf, tx1_model, tx1_rating, tx1_serial = row_data[11:15]
    tx2_manuf, tx2_model, tx2_rating, tx2_serial = row_data[15:19]
    lv1_manuf, lv1_model, lv1_rating, lv1_serial, lv1_type = row_data[19:24]
    lv2_manuf, lv2_model, lv2_rating, lv2_serial, lv2_type = row_data[24:29]
    bat1_manuf, bat1_model, bat1_serial = row_data[29:32]
    bat2_manuf, bat2_model, bat2_serial = row_data[32:35]
    bat3_manuf, bat3_model, bat3_serial = row_data[35:38]
    bat4_manuf, bat4_model, bat4_serial = row_data[38:41]
    fr, qr, rm, ts = row_data[41:45]


    def plain(col, val, bold=False, align=CENTER):
        c = ws.cell(row=r, column=col, value=val)
        c.font = Font(name="Calibri", size=10, bold=bold)
        c.alignment = align
        c.fill = row_fill
        c.border = BORDER_INNER


    def info(col, val):
        c = ws.cell(row=r, column=col, value=val)
        c.font = INFO_FONT
        c.alignment = CENTER
        c.fill = row_fill
        c.border = BORDER_INNER


    def yn(col, val):
        c = ws.cell(row=r, column=col, value=val)
        c.font = YES_FONT if val == "YES" else NO_FONT
        c.fill = YES_FILL if val == "YES" else NO_FILL
        c.alignment = CENTER
        c.border = BORDER_INNER


    plain(1, num, bold=True)
    c = ws.cell(row=r, column=2, value=name)
    c.font = Font(name="Calibri", size=10)
    c.alignment = LEFT
    c.fill = row_fill
    c.border = BORDER_INNER

    plain(3, sub_type)
    info(4, time_val)
    info(5, date)
    info(6, ambient)
    info(7, humidity)
    plain(8, sg_manuf, align=LEFT)
    plain(9, sg_type)
    plain(10, sg_rating)
    plain(11, sg_model, align=LEFT)
    plain(12, tx1_manuf, align=LEFT)
    plain(13, tx1_model, align=LEFT)
    plain(14, tx1_rating)
    plain(15, tx1_serial, align=LEFT)
    plain(16, tx2_manuf, align=LEFT)
    plain(17, tx2_model, align=LEFT)
    plain(18, tx2_rating)
    plain(19, tx2_serial, align=LEFT)
    plain(20, lv1_manuf, align=LEFT)
    plain(21, lv1_model, align=LEFT)
    plain(22, lv1_rating)
    plain(23, lv1_serial, align=LEFT)
    plain(24, lv1_type)
    plain(25, lv2_manuf, align=LEFT)
    plain(26, lv2_model, align=LEFT)
    plain(27, lv2_rating)
    plain(28, lv2_serial, align=LEFT)
    plain(29, lv2_type)
    plain(30, bat1_manuf, align=LEFT)
    plain(31, bat1_model, align=LEFT)
    plain(32, bat1_serial, align=LEFT)
    plain(33, bat2_manuf, align=LEFT)
    plain(34, bat2_model, align=LEFT)
    plain(35, bat2_serial, align=LEFT)
    plain(36, bat3_manuf, align=LEFT)
    plain(37, bat3_model, align=LEFT)
    plain(38, bat3_serial, align=LEFT)
    plain(39, bat4_manuf, align=LEFT)
    plain(40, bat4_model, align=LEFT)
    plain(41, bat4_serial, align=LEFT)
    yn(42, fr)
    yn(43, qr)
    yn(44, rm)
    yn(45, ts)

ws.freeze_panes = "A5"

# ── 6. Create "Panel Schedule" Sheet in SUBSTATION_CHECKLIST.xlsx ─────────────
print("  Generating 'Panel Schedule' sheet (extracting circuit panel details)...", flush=True)
ws_panel = wb.create_sheet(title="Panel Schedule")

col_widths_panel = [15, 30, 20, 18, 25, 20, 20, 20, 15, 15, 22]
for i, w in enumerate(col_widths_panel, start=1):
    ws_panel.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

ws_panel.row_dimensions[1].height = 28
panel_headers = [
    "SUBSTATION NO", "SUBSTATION NAME", "DESCRIPTION", "LINK NUMBER", "NAME", "CABLE TYPE",
    "BREAKER STATUS", "BUSBAR POSITION", "LOAD AMP",
    "HEATER AMP", "SERIAL NUMBER"
]

for col, hdr in enumerate(panel_headers, start=1):
    c = ws_panel.cell(row=1, column=col, value=hdr)
    c.font = HEADER_FONT
    c.fill = HEADER_FILL
    c.alignment = CENTER
    c.border = BORDER_HEAD

panel_row_idx = 2
for row_item in rows:
    num = str(row_item[0]).strip()
    sub_name = str(row_item[1]).strip()
    m_base = re.match(r'^(\d+)', num)
    base_num = int(m_base.group(1)) if m_base else num

    is_in_scope = True if matches_scope is None else matches_scope(base_num)
    existing_panels = existing_panels_dict.get(num, existing_panels_dict.get(str(base_num), []))

    if not is_in_scope and existing_panels:
        # OUTSIDE SCOPE: 100% PRESERVE EXISTING PANELS AS-IS
        for ep in existing_panels:
            p_vals = [
                num,
                sub_name,
                ep.get("description", ""),
                ep.get("link_number", ""),
                ep.get("name", ""),
                ep.get("cable_type", ""),
                ep.get("breaker_status", ""),
                ep.get("busbar_position", "-"),
                ep.get("load_amp", ""),
                ep.get("heater_amp", ""),
                ep.get("serial_number", "")
            ]
            ws_panel.row_dimensions[panel_row_idx].height = 20
            row_fill = ALT_FILL if panel_row_idx % 2 == 0 else PLAIN_FILL
            for c_idx, val in enumerate(p_vals, start=1):
                cell = ws_panel.cell(row=panel_row_idx, column=c_idx, value=val)
                cell.font = Font(name="Calibri", size=10)
                cell.fill = row_fill
                cell.border = BORDER_INNER
                cell.alignment = LEFT if c_idx in [1, 3, 4] else CENTER
            panel_row_idx += 1
        continue

    ts_fields = ts_data.get(base_num, ts_data.get(num, {}))

    if num.endswith('A'):
        extracted_panels = ts_fields.get("panels_1", ts_fields.get("panels", []))
    elif num.endswith('B'):
        extracted_panels = ts_fields.get("panels_2", [])
    else:
        extracted_panels = ts_fields.get("panels", [])

    # Merge panels: if existing panel has edited values, keep them!
    max_panels = max(len(extracted_panels), len(existing_panels))
    merged_panels = []
    for p_idx in range(max_panels):
        tp = extracted_panels[p_idx] if p_idx < len(extracted_panels) else {}
        ep = existing_panels[p_idx] if p_idx < len(existing_panels) else {}

        p_desc = resolve_field(ep.get("description"), tp.get("description", ""), "Panel Desc", num)
        p_link = resolve_field(ep.get("link_number"), tp.get("link_number", ""), "Panel Link", num)
        p_name = resolve_field(ep.get("name"), tp.get("name", ""), "Panel Name", num)
        p_cable = resolve_field(ep.get("cable_type"), tp.get("cable_type", ""), "Panel Cable", num)
        p_breaker = resolve_field(ep.get("breaker_status"), tp.get("breaker_status", ""), "Panel Breaker", num)
        p_busbar = resolve_field(ep.get("busbar_position"), tp.get("busbar_position", "-"), "Panel Busbar", num)
        p_load = resolve_field(ep.get("load_amp"), tp.get("load_amp", ""), "Panel Load", num)
        p_heater = resolve_field(ep.get("heater_amp"), tp.get("heater_amp", ""), "Panel Heater", num)
        p_serial = resolve_field(ep.get("serial_number"), tp.get("serial_number", ""), "Panel Serial", num)

        merged_panels.append([
            num,
            sub_name,
            p_desc,
            p_link,
            p_name,
            p_cable,
            p_breaker,
            p_busbar,
            p_load,
            p_heater,
            p_serial
        ])

    for p_vals in merged_panels:
        ws_panel.row_dimensions[panel_row_idx].height = 20
        row_fill = ALT_FILL if panel_row_idx % 2 == 0 else PLAIN_FILL

        for c_idx, val in enumerate(p_vals, start=1):
            cell = ws_panel.cell(row=panel_row_idx, column=c_idx, value=val)
            cell.font = Font(name="Calibri", size=10)
            cell.fill = row_fill
            cell.border = BORDER_INNER
            cell.alignment = LEFT if c_idx in [1, 3, 4] else CENTER

        panel_row_idx += 1

ws_panel.freeze_panes = "A2"
print(f"  -> Mapped {panel_row_idx - 2} circuit panels across all substations.", flush=True)

# ── 7. Create "Substation Info" sheet ──────────────────────────────────────────
print("  Generating 'Substation Info' sheet (GPS, ERMS/Site names, F/L No)...", flush=True)
ws2 = wb.create_sheet(title="Substation Info")

col_widths2 = [10, 40, 20, 18, 30, 40, 40, 28, 28]
for i, w in enumerate(col_widths2, start=1):
    ws2.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

ws2.row_dimensions[1].height = 32
ws2.merge_cells("A1:I1")
t2 = ws2["A1"]
t2.value = po_title + "  |  Substation Info (Date, Coordinates, Names, F/L No)"
t2.font = Font(name="Calibri", bold=True, size=13, color="FFFFFF")
t2.fill = PatternFill("solid", fgColor="0B2D6B")
t2.alignment = Alignment(horizontal="center", vertical="center")

ws2.row_dimensions[2].height = 18
ws2.merge_cells("A2:I2")
g2 = ws2["A2"]
g2.value = "Generated on: " + datetime.now().strftime("%d/%m/%Y  %H:%M")
g2.font = Font(name="Calibri", italic=True, size=9, color="7F8C8D")
g2.fill = PatternFill("solid", fgColor="EBF5FB")
g2.alignment = Alignment(horizontal="right", vertical="center")

headers2 = [
    "NO", "SUBSTATION NAME", "SUBSTATION TYPE", "DATE", "GPS COORDINATES",
    "SUBSTATION NAME (ERMS)", "SUBSTATION NAME (SITE)", "F/L NO (ERMS)", "F/L NO (SITE)"
]
ws2.row_dimensions[3].height = 28
for col, hdr in enumerate(headers2, start=1):
    c = ws2.cell(row=3, column=col, value=hdr)
    c.font = HEADER_FONT
    c.fill = HEADER_FILL
    c.alignment = CENTER
    c.border = BORDER_HEAD

for i, row in enumerate(rows):
    r = i + 4
    ws2.row_dimensions[r].height = 20
    row_fill = ALT_FILL if i % 2 == 0 else PLAIN_FILL
    data2 = [row[0], row[1], row[2], row[45], row[46], row[47], row[48], row[49], row[50]]
    for col_idx, val in enumerate(data2, start=1):
        c = ws2.cell(row=r, column=col_idx, value=val)
        c.font = Font(name="Calibri", size=10)
        if col_idx == 1:
            c.font = Font(name="Calibri", size=10, bold=True)
            c.alignment = CENTER
        elif col_idx in (3, 4):
            c.alignment = CENTER
        else:
            c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        c.fill = row_fill
        c.border = BORDER_INNER

ws2.freeze_panes = "A4"

# Create Backup sheet with identical table formatting
ws_bk = wb.copy_worksheet(ws)
ws_bk.title = "Backup"

create_timestamped_backup(OUT)
print("  Saving SUBSTATION_CHECKLIST.xlsx...", flush=True)
try:
    wb.save(OUT)
    print(f"  [OK] Checklist saved to: {OUT}", flush=True)
except PermissionError:
    print(f"  [CRITICAL ERROR] Cannot save: {os.path.basename(OUT)} is currently open in Excel!", flush=True)
    print(f"  --> Please close {os.path.basename(OUT)} and run extraction again.", flush=True)
    sys.exit(1)

# ── 8. Update Master List.xlsx via Unified Sync Engine ────────────────────────
print("=" * 60, flush=True)
print("  STEP 4: SYNCING DATA & CSV MAPPINGS TO MASTER LIST.XLSX", flush=True)
print("=" * 60, flush=True)
try:
    import sync_master_list
    sync_master_list.sync_and_merge_to_master(scope_arg)
except Exception as e:
    print(f"  [ERROR] Master List sync encountered error: {e}", flush=True)

total_elapsed = time.time() - start_overall_time
print("")
print("=" * 60, flush=True)
print(f"  ALL TASKS COMPLETED SUCCESSFULLY IN {total_elapsed:.2f} SECONDS", flush=True)
print("=" * 60, flush=True)