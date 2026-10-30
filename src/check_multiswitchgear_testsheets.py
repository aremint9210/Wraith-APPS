import os
import openpyxl
from pathlib import Path

DEFAULT_BASE = r"C:\Users\aremi\Desktop\REPORT CBM\PO 42234207 - JOHOR - JBU - 400 PE IR US TEV"
BASE_DIR = Path(os.environ.get("BASE_DIR", DEFAULT_BASE))
TESTSHEET_DIR = BASE_DIR / "04. TESTSHEET"

def inspect_testsheet_multiswitchgears(filepath):
    """
    Inspects a test sheet Excel file to check if it contains multiple test sheets.
    Determines whether it represents:
    - Case A: 2 Different Switchgears -> Split into e.g. 113A and 113B
    - Case B: > 4 Panels (Same Switchgear) -> Single 113.csv with extra pages
    """
    filepath = Path(filepath)
    if not filepath.exists():
        return None

    try:
        wb = openpyxl.load_workbook(filepath, data_only=True, read_only=True)
        sheet_names = wb.sheetnames
        ts_sheets = [s for s in sheet_names if 'pce testsheet' in s.lower()]
        
        if len(ts_sheets) <= 1:
            wb.close()
            return {"multiple_sheets": False, "reason": "Single test sheet"}

        # Multiple test sheets present, check PCE VI for Switchgear 1 vs Switchgear 2
        vi_sheet = next((s for s in sheet_names if 'pce vi' in s.lower() or 'vi' in s.lower()), None)
        
        swg1_mfr, swg1_sn, swg1_type = "", "", ""
        swg2_mfr, swg2_sn, swg2_type = "", "", ""

        if vi_sheet:
            ws_vi = wb[vi_sheet]
            for r in range(1, 30):
                row_vals = [str(ws_vi.cell(r, c).value or '').strip() for c in range(1, 16)]
                line_str = " | ".join(row_vals).upper()

                if "SWITCEGAR TYPE 1" in line_str or "SWITCHGEAR TYPE 1" in line_str:
                    swg1_type = line_str
                elif "SWITCEGAR TYPE 2" in line_str or "SWITCHGEAR TYPE 2" in line_str:
                    swg2_type = line_str
                
                # Check Manufacturer & S/N rows
                if "MANUFACTURER" in line_str:
                    if not swg1_mfr:
                        swg1_mfr = line_str
                    else:
                        swg2_mfr = line_str
                
                if "S/N" in line_str or "SERIAL" in line_str:
                    if not swg1_sn:
                        swg1_sn = line_str
                    else:
                        swg2_sn = line_str

        wb.close()

        def clean_sn(val):
            if val is None:
                return ""
            s = str(val).strip()
            if s.endswith(".0"):
                return s[:-2]
            try:
                # Handle floats represented as numbers or scientific notation
                f = float(s)
                if f.is_integer():
                    return str(int(f))
            except ValueError:
                pass
            return s

        wb_full = openpyxl.load_workbook(filepath, data_only=True)
        if vi_sheet and vi_sheet in wb_full.sheetnames:
            ws = wb_full[vi_sheet]
            mfr1 = str(ws.cell(12, 3).value or ws.cell(12, 6).value or '').strip()
            mfr2 = str(ws.cell(15, 3).value or ws.cell(15, 6).value or '').strip()
            sn1  = clean_sn(ws.cell(13, 15).value)
            sn2  = clean_sn(ws.cell(16, 15).value)
        else:
            mfr1, mfr2, sn1, sn2 = "", "", "", ""
        wb_full.close()

        different_mfr = (mfr1.upper() != mfr2.upper()) if (mfr1 and mfr2) else False
        different_sn  = (sn1.upper() != sn2.upper()) if (sn1 and sn2) else False

        if different_mfr or different_sn:
            case_type = "CASE_A_SPLIT_SWITCHGEAR"
            desc = f"2 Different Switchgears detected: SWG1=({mfr1}, S/N:{sn1}) vs SWG2=({mfr2}, S/N:{sn2})"
        else:
            case_type = "CASE_B_MORE_THAN_4_PANELS"
            desc = f">4 Panels on single switchgear: SWG=({mfr1}, S/N:{sn1})"

        return {
            "multiple_sheets": True,
            "sheets": ts_sheets,
            "case_type": case_type,
            "description": desc,
            "swg1": {"mfr": mfr1, "sn": sn1},
            "swg2": {"mfr": mfr2, "sn": sn2}
        }

    except Exception as e:
        return {"multiple_sheets": False, "error": str(e)}

if __name__ == "__main__":
    if TESTSHEET_DIR.exists():
        print(f"Scanning test sheets in: {TESTSHEET_DIR}")
        for f in sorted(os.listdir(TESTSHEET_DIR)):
            if f.endswith(".xlsx") and not f.startswith("~$"):
                res = inspect_testsheet_multiswitchgears(TESTSHEET_DIR / f)
                if res and res.get("multiple_sheets"):
                    print(f"\n[MULTIPLE SWITCHGEARS/PANELS DETECTED] {f}:")
                    print(f"   Case: {res['case_type']}")
                    print(f"   Details: {res['description']}")
