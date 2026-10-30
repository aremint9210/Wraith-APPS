import os
import sys
import glob
import re
import time
import argparse

# Force UTF-8 output
if sys.stdout and hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

def log(msg, level="INFO"):
    print(f"[{level}] {msg}", flush=True)

def find_testsheet_dir(base_dir):
    candidates = [
        os.path.join(base_dir, "04. TESTSHEET", "test sheet PJY"),
        os.path.join(base_dir, "04. TESTSHEET"),
        os.path.join(base_dir, "TESTSHEET"),
        base_dir
    ]
    for c in candidates:
        if os.path.exists(c):
            xlsm_files = [f for f in glob.glob(os.path.join(c, "*.xlsm")) if not os.path.basename(f).startswith('~$')]
            if xlsm_files:
                return c, xlsm_files
    return None, []

def get_known_names(base_dir):
    known = {}
    
    # 1. From 02. QUICK REPORT
    qr_dir = os.path.join(base_dir, "02. QUICK REPORT")
    if os.path.exists(qr_dir):
        for f in os.listdir(qr_dir):
            m = re.match(r"^(\d+)\.\s*(\d{8})\s*(.*?)(?:\s*\((?:VI|CBM&VI|CBM & VI)\))?\.pdf$", f, re.IGNORECASE)
            if m:
                known[int(m.group(1))] = m.group(3).strip()

    # 2. From 03. RAW MATERIAL
    rm_dir = os.path.join(base_dir, "03. RAW MATERIAL")
    if os.path.exists(rm_dir):
        for d in os.listdir(rm_dir):
            m = re.match(r"^(\d+)\.\s*(.*)$", d)
            if m:
                idx = int(m.group(1))
                if idx not in known:
                    known[idx] = m.group(2).strip()

    # 3. From SUBSTATION_CHECKLIST.xlsx
    cl_path = os.path.join(base_dir, "SUBSTATION_CHECKLIST.xlsx")
    if os.path.exists(cl_path):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(cl_path, read_only=True, data_only=True)
            s = wb.active
            for r in range(5, s.max_row + 1):
                num = s.cell(r, 1).value
                name = s.cell(r, 2).value
                if num is not None and name:
                    try:
                        n_int = int(num)
                        if n_int not in known:
                            known[n_int] = str(name).strip()
                    except ValueError:
                        pass
            wb.close()
        except Exception as e:
            log(f"Warning reading checklist: {e}", "WARN")

    return known

def split_testsheets(base_dir, target_files=None):
    ts_dir, all_files = find_testsheet_dir(base_dir)
    if not ts_dir or not all_files:
        log(f"No .xlsm test sheet files found under: {base_dir}", "ERROR")
        return False

    files_to_process = target_files if target_files else all_files
    out_dir = os.path.join(base_dir, "04. TESTSHEET", "INDIVIDUAL EXCEL")
    os.makedirs(out_dir, exist_ok=True)

    log("=" * 65, "INFO")
    log("   TIANU TESTSHEET SPLITTER (SUBSTATION DATA + VI)   ", "INFO")
    log("=" * 65, "INFO")
    log(f"Base Directory   : {base_dir}", "INFO")
    log(f"Source Folder    : {ts_dir}", "INFO")
    log(f"Output Directory : {out_dir}", "INFO")
    log(f"Daily Workbooks  : {len(files_to_process)} file(s) found", "INFO")

    known_names = get_known_names(base_dir)

    import openpyxl
    file_tasks = {}
    total_substations = 0

    log("Scanning workbooks for valid substation test sheet pairs...", "INFO")
    for f in sorted(files_to_process):
        fname = os.path.basename(f)
        date_str = fname.replace('.xlsm', '')
        try:
            wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        except Exception as e:
            log(f"Could not open {fname}: {e}", "WARN")
            continue

        file_tasks[f] = []
        for i in range(1, 16):
            cba = f"PE {i}"
            vi = f"PE {i} (VI)"
            if cba in wb.sheetnames and vi in wb.sheetnames:
                s = wb[cba]
                ref_val = s['C4'].value
                c3_val = s['C3'].value
                c176_val = s['C176'].value

                # Filter out empty template sheets
                if ref_val in [None, 0, '0', '', '-']:
                    continue
                if (c3_val in [None, 0, '0', '', 'PE NAME:'] and c176_val in [None, 0, '0', '', 'None']):
                    continue

                try:
                    ref_int = int(ref_val)
                except ValueError:
                    continue

                if ref_int in known_names:
                    name = known_names[ref_int]
                else:
                    name = str(c3_val or c176_val or '').strip()
                    name = re.sub(r'[\\/*?:"<>|]', '-', name)
                    name = re.sub(r'\s+', ' ', name).strip()

                out_filename = f"{ref_int:02d}. {date_str} {name}.xlsx"
                file_tasks[f].append({
                    'ref_no': ref_int,
                    'cba': cba,
                    'vi': vi,
                    'out_path': os.path.join(out_dir, out_filename),
                    'filename': out_filename
                })
                total_substations += 1
        wb.close()

    log(f"Total valid substations to export: {total_substations}", "INFO")
    if total_substations == 0:
        log("No valid substations found to extract.", "WARN")
        return True

    # Use Excel COM for native 100% lossless copy
    try:
        import win32com.client
        has_win32 = True
    except ImportError:
        has_win32 = False

    t0 = time.time()
    processed_count = 0

    if has_win32:
        log("Launching Microsoft Excel automation engine...", "INFO")
        excel = win32com.client.Dispatch("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        excel.ScreenUpdating = False

        try:
            for src_file, tasks in file_tasks.items():
                if not tasks:
                    continue
                fname = os.path.basename(src_file)
                log(f"Processing {fname} ({len(tasks)} substations)...", "INFO")
                wb_src = excel.Workbooks.Open(src_file, ReadOnly=True)

                for t in tasks:
                    wb_src.Sheets([t['cba'], t['vi']]).Copy()
                    wb_new = excel.ActiveWorkbook

                    # Rename sheets cleanly: Substation Data & VI
                    wb_new.Sheets(1).Name = "Substation Data"
                    wb_new.Sheets(2).Name = "VI"

                    # Save as standard .xlsx (xlOpenXMLWorkbook = 51)
                    wb_new.SaveAs(t['out_path'], FileFormat=51)
                    wb_new.Close(SaveChanges=False)

                    processed_count += 1
                    log(f"[{processed_count:02d}/{total_substations}] Generated: {t['filename']}", "OK")

                wb_src.Close(SaveChanges=False)
        finally:
            excel.Quit()
    else:
        log("Excel COM not available, using openpyxl fallback...", "WARN")
        # openpyxl fallback implementation if ever needed
        for src_file, tasks in file_tasks.items():
            wb_src = openpyxl.load_workbook(src_file)
            for t in tasks:
                wb_new = openpyxl.Workbook()
                wb_new.remove(wb_new.active)
                # Note: copy sheets
                s1 = wb_new.create_sheet(title="Substation Data")
                s2 = wb_new.create_sheet(title="VI")
                # copy rows
                for r in wb_src[t['cba']].iter_rows(values_only=False):
                    s1.append([cell.value for cell in r])
                for r in wb_src[t['vi']].iter_rows(values_only=False):
                    s2.append([cell.value for cell in r])
                wb_new.save(t['out_path'])
                processed_count += 1
                log(f"[{processed_count:02d}/{total_substations}] Generated: {t['filename']}", "OK")

    elapsed = time.time() - t0
    log("=" * 65, "OK")
    log(f"SUCCESS: {processed_count} individual substation test sheets created!", "OK")
    log(f"Duration: {elapsed:.1f}s | Saved in: {out_dir}", "OK")
    log("=" * 65, "OK")
    return True

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Split daily test sheets by substation (Substation Data + VI) - TIANU")
    parser.add_argument("--dir", help="Base directory of the PO / workspace", default=None)
    args = parser.parse_args()

    base_dir = args.dir or os.environ.get("BASE_DIR") or r"C:\Users\aremi\Desktop\REPORT CBM\PO 42390496 - TNBD SELANGOR - PETALING JAYA -210PCE"
    split_testsheets(base_dir)
