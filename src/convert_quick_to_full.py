import os
import sys
import re
import tempfile
import shutil

try:
    import docx
    import fitz
    import win32com.client
    HAS_REQUIRED_LIBS = True
except ImportError:
    HAS_REQUIRED_LIBS = False

DEFAULT_BASE = r"C:\Users\aremi\Desktop\REPORT CBM\PO 42234207 - JOHOR - JBU - 400 PE IR US TEV"
BASE_DIR = os.environ.get("BASE_DIR", DEFAULT_BASE)
NOT_DONE_DIR = os.path.join(BASE_DIR, "01. FULL REPORT", "NOT DONE")
QR_DIR = os.path.join(BASE_DIR, "02. QUICK REPORT")


def extract_prefix_num(filename):
    m = re.match(r'^(\d+)', filename)
    return int(m.group(1)) if m else None


def is_substation_file(filename):
    name_lower = filename.lower()
    if "group 0" in name_lower or "group 1" in name_lower:
        return False
    return extract_prefix_num(filename) is not None


def process_substation_swap(num, word_app, target_dir):
    """
    1. Reads original .docx and .pdf from 02. QUICK REPORT without editing originals.
    2. Modifies 'QUICK' -> 'FULL' in a temp .docx copy.
    3. Exports temp .docx to a temp PDF via Word COM.
    4. Swaps Page 1 of original PDF with Page 1 of temp PDF.
    5. Saves output PDF into target_dir (01. FULL REPORT\\NOT DONE).
    """
    if not os.path.exists(QR_DIR):
        print(f"[SKIP] QUICK REPORT directory not found: {QR_DIR}")
        return False

    qr_docx_list = [f for f in os.listdir(QR_DIR) if f.lower().endswith('.docx') and is_substation_file(f) and extract_prefix_num(f) == num]
    qr_pdf_list  = [f for f in os.listdir(QR_DIR) if f.lower().endswith('.pdf') and is_substation_file(f) and extract_prefix_num(f) == num]

    if not qr_docx_list or not qr_pdf_list:
        print(f"[SKIP] Substation {num:03d}: Missing matching .docx or .pdf in QUICK REPORT.")
        return False

    src_docx = os.path.join(QR_DIR, qr_docx_list[0])
    orig_pdf = os.path.join(QR_DIR, qr_pdf_list[0])

    prefix = f"{num:03d}"
    parts = qr_pdf_list[0].split(". ", 1)
    rest_name = parts[1] if len(parts) > 1 else qr_pdf_list[0]
    target_pdf = os.path.join(target_dir, f"{prefix}. {rest_name}")

    # Check if PDF already exists in target_dir (01. FULL REPORT\NOT DONE)
    existing_pdfs = [f for f in os.listdir(target_dir) if f.lower().endswith('.pdf') and not f.startswith("~$") and is_substation_file(f) and extract_prefix_num(f) == num]
    if existing_pdfs or os.path.exists(target_pdf):
        exist_file = existing_pdfs[0] if existing_pdfs else os.path.basename(target_pdf)
        print(f"[SKIP] Substation {num:03d}: PDF already exists in NOT DONE folder ({exist_file}).")
        return False

    try:
        # 1. Update text in temporary .docx copy
        doc = docx.Document(src_docx)
        for p in doc.paragraphs:
            if "QUICK" in p.text:
                for run in p.runs:
                    if "QUICK" in run.text:
                        run.text = run.text.replace("QUICK", "FULL")
                if "QUICK" in p.text:
                    p.text = p.text.replace("QUICK", "FULL")

        for t in doc.tables:
            for r in t.rows:
                for c in r.cells:
                    for p in c.paragraphs:
                        if "QUICK" in p.text:
                            for run in p.runs:
                                if "QUICK" in run.text:
                                    run.text = run.text.replace("QUICK", "FULL")
                            if "QUICK" in p.text:
                                p.text = p.text.replace("QUICK", "FULL")

        temp_docx = os.path.join(tempfile.gettempdir(), f"temp_docx_{num}_{os.getpid()}.docx")
        temp_pdf  = os.path.join(tempfile.gettempdir(), f"temp_pdf_{num}_{os.getpid()}.pdf")

        doc.save(temp_docx)

        # 2. Export updated docx to temp PDF using Word COM
        wdoc = word_app.Documents.Open(temp_docx)
        try:
            wdoc.DoNotCompressPictures = True
        except Exception:
            pass
        wdoc.ExportAsFixedFormat(temp_pdf, 17) # 17 = wdExportFormatPDF
        wdoc.Close(False)

        if os.path.exists(temp_docx):
            os.remove(temp_docx)

        # 3. Swap Page 1 of orig_pdf with Page 1 of temp_pdf
        pdf_orig_doc = fitz.open(orig_pdf)
        pdf_page1_doc = fitz.open(temp_pdf)

        out_doc = fitz.open()
        # Insert Word-generated Page 1
        out_doc.insert_pdf(pdf_page1_doc, from_page=0, to_page=0)

        # Append Pages 2..N of original PDF untouched
        if len(pdf_orig_doc) > 1:
            out_doc.insert_pdf(pdf_orig_doc, from_page=1, to_page=len(pdf_orig_doc) - 1)

        out_doc.save(target_pdf)
        out_doc.close()
        pdf_orig_doc.close()
        pdf_page1_doc.close()

        if os.path.exists(temp_pdf):
            os.remove(temp_pdf)

        print(f"[PAGE 1 SWAPPED] Substation {num:03d} -> {os.path.basename(target_pdf)}")
        return True

    except Exception as e:
        print(f"[ERROR] Failed substation {num:03d}: {e}")
        return False


def process_directory(target_dir):
    if not os.path.exists(target_dir):
        print(f"[ERROR] Target directory does not exist: {target_dir}")
        return 0

    print("============================================================")
    print(" Scan NOT DONE First -> Page 1 Swap via Word & Original PDF")
    print(f" Target Folder (NOT DONE): {target_dir}")
    print(f" Source Folder (QUICK REPORT): {QR_DIR}")
    print("============================================================")

    # STEP 1: Scan NOT DONE directory first to get present substation numbers
    not_done_files = os.listdir(target_dir)
    target_nums = sorted(list(set(
        extract_prefix_num(f) for f in not_done_files if extract_prefix_num(f) is not None
    )))

    if not target_nums:
        print("No files with substation numbers found in NOT DONE folder.")
        return 0

    print(f"Found {len(target_nums)} substation number(s) in NOT DONE folder to process:")
    print(f"Numbers: {target_nums}\n")

    word_app = win32com.client.DispatchEx("Word.Application")
    word_app.Visible = False
    word_app.DisplayAlerts = False

    count_updated = 0
    try:
        for num in target_nums:
            if process_substation_swap(num, word_app, target_dir):
                count_updated += 1
    finally:
        try:
            word_app.Quit()
        except Exception:
            pass

    print("============================================================")
    print(f" Completed! Processed {count_updated} / {len(target_nums)} substations found in NOT DONE.")
    print("============================================================")
    return count_updated


if __name__ == "__main__":
    target = sys.argv[1] if len(sys.argv) > 1 else NOT_DONE_DIR
    process_directory(target)
