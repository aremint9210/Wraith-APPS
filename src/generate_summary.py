import os
import sys
import pandas as pd
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

DEFAULT_BASE = Path(r"C:\Users\aremi\Desktop\REPORT CBM\PO 42234207 - JOHOR - JBU - 400 PE IR US TEV")
BASE_DIR = Path(os.environ.get("BASE_DIR", DEFAULT_BASE))

MAPPING_DIR = Path(os.environ.get("MAPPING_DIR", BASE_DIR / "06. EXCELL MAPPING"))
TEMPLATE_DIR = Path(os.environ.get("TEMPLATE_DIR", BASE_DIR / "05. TEMPLATE"))
OUTPUT_DIR = Path(os.environ.get("OUTPUT_DIR", BASE_DIR / "07. OUTPUT FILE"))
RM_DIR = Path(os.environ.get("RM_DIR", BASE_DIR / "03. RAW MATERIAL"))
import re

def resolve_rm_subfolder_for_report(rm_dir, csv_prefix):
    if not rm_dir or not Path(rm_dir).exists():
        return None
    val_str = str(csv_prefix).strip()
    m = re.match(r'^0*(\d+)([a-zA-Z]*)', val_str)
    if m:
        clean_num = m.group(1) + m.group(2).upper()
        base_digits = m.group(1)
        suffix = m.group(2).upper()
    else:
        clean_num = val_str
        base_digits = val_str
        suffix = ''

    try:
        dir_entries = [d for d in Path(rm_dir).iterdir() if d.is_dir()]
    except Exception:
        return None

    # 1. Exact match
    for d in dir_entries:
        if d.name.lower() == val_str.lower() or d.name.lower() == clean_num.lower():
            return d

    # 2. Strict match with suffix
    if suffix:
        for d in dir_entries:
            m_d = re.match(r'^0*(\d+[a-zA-Z]?)[\.\s\-_]', d.name)
            if m_d and m_d.group(1).upper() == clean_num:
                return d

    # 3. Numeric prefix match
    for d in dir_entries:
        m_d = re.match(r'^0*(\d+[a-zA-Z]?)[\.\s\-_]', d.name)
        if m_d and m_d.group(1).upper() == clean_num:
            return d

    # 4. Fallback startswith
    for d in dir_entries:
        if (d.name.startswith(f"{clean_num}.") or d.name.startswith(f"{clean_num} ") or 
            d.name.startswith(f"{val_str}.") or d.name.startswith(f"{val_str} ") or
            d.name.startswith(f"{base_digits}.") or d.name.startswith(f"{base_digits} ")):
            return d

    return None


def set_cell_shading(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def set_table_borders(table):
    tblPr = table._element.xpath('w:tblPr')
    if not tblPr:
        return
    tblBorders = tblPr[0].find(qn('w:tblBorders'))
    if tblBorders is None:
        tblBorders = OxmlElement('w:tblBorders')
        tblPr[0].append(tblBorders)
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        existing = tblBorders.find(qn(f'w:{border_name}'))
        if existing is not None:
            tblBorders.remove(existing)
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), 'auto')
        tblBorders.append(border)


def get_category(type_v):
    tv = str(type_v).upper().strip()
    if "OVERVIEW" in tv or "PANEL" in tv:
        return "SWG"
    if "TX1" in tv or "POWER TX1" in tv:
        return "TX1"
    if "TX2" in tv or "POWER TX2" in tv:
        return "TX2"
    if "FP" in tv or "FEEDER PILLAR" in tv:
        fp_digits = "".join(filter(str.isdigit, tv))
        return f"FP_{fp_digits if fp_digits else '1'}"
    if "BATTERY" in tv:
        b_digits = "".join(filter(str.isdigit, tv))
        return f"BATTERY_{b_digits if b_digits else '1'}"
    return tv


def is_substation_in_range(substation_val, range_str):
    if not range_str or not str(range_str).strip():
        return False
    range_str = str(range_str).strip().upper()
    if range_str in ['ALL', '*', 'YES', 'TRUE', '1']:
        return True
    if range_str in ['NONE', 'NO', 'FALSE', '0', '-']:
        return False
    
    m = re.search(r'(\d+)', str(substation_val))
    if not m:
        return False
    sub_num = int(m.group(1))
    
    parts = re.split(r'[,;\s]+', range_str)
    for part in parts:
        part = part.strip()
        if not part:
            continue
        if '-' in part:
            bounds = part.split('-')
            if len(bounds) == 2:
                try:
                    low = int(bounds[0].strip())
                    high = int(bounds[1].strip())
                    if low <= sub_num <= high:
                        return True
                except ValueError:
                    pass
        else:
            try:
                if int(part) == sub_num:
                    return True
            except ValueError:
                pass
    return False


def resolve_scopes_for_substation(csv_prefix, base_dir=None):
    env_ir = os.environ.get("IR_RANGE", "")
    env_us = os.environ.get("US_RANGE", "")
    env_tev = os.environ.get("TEV_RANGE", "")
    
    config_ir, config_us, config_tev = "", "", ""
    if base_dir:
        scope_cfg_path = Path(base_dir) / "scope_config.json"
        if scope_cfg_path.exists():
            try:
                import json
                with open(scope_cfg_path, "r", encoding="utf-8") as f:
                    cfg = json.load(f)
                    config_ir = cfg.get("ir_range", "")
                    config_us = cfg.get("us_range", "")
                    config_tev = cfg.get("tev_range", "")
            except Exception:
                pass

    ir_range = env_ir or config_ir
    us_range = env_us or config_us
    tev_range = env_tev or config_tev

    if not ir_range:
        has_ir = True
    else:
        has_ir = is_substation_in_range(csv_prefix, ir_range)

    has_us = is_substation_in_range(csv_prefix, us_range)
    has_tev = is_substation_in_range(csv_prefix, tev_range)

    if not ir_range and not us_range and not tev_range and base_dir:
        base_name = Path(base_dir).name.upper()
        if "TEV" in base_name:
            has_tev = True
        if "US" in base_name or "UE" in base_name:
            has_us = True
        has_ir = True

    return has_ir, has_us, has_tev


def resolve_template_dir_for_substation(csv_prefix, base_dir=None, default_template_dir=None):
    if default_template_dir is None:
        default_template_dir = TEMPLATE_DIR
    default_template_dir = Path(default_template_dir)

    has_ir, has_us, has_tev = resolve_scopes_for_substation(csv_prefix, base_dir)

    t_ir = os.environ.get("TEMPLATE_IR_DIR", "")
    t_us = os.environ.get("TEMPLATE_IR_UE_DIR", "")
    t_tev = os.environ.get("TEMPLATE_IR_UE_TEV_DIR", "")

    if base_dir:
        scope_cfg_path = Path(base_dir) / "scope_config.json"
        if scope_cfg_path.exists():
            try:
                import json
                with open(scope_cfg_path, "r", encoding="utf-8") as f:
                    sc = json.load(f)
                    t_ir = t_ir or sc.get("template_ir", "")
                    t_us = t_us or sc.get("template_ir_ue", "")
                    t_tev = t_tev or sc.get("template_ir_ue_tev", "")
            except Exception:
                pass

    if has_tev and t_tev and Path(t_tev).exists():
        return Path(t_tev)
    if has_us and t_us and Path(t_us).exists():
        return Path(t_us)
    if t_ir and Path(t_ir).exists():
        return Path(t_ir)

    # Dynamic auto-routing inside default_template_dir if subdirectories exist
    if default_template_dir.exists():
        subdirs = [d for d in default_template_dir.iterdir() if d.is_dir() and not d.name.startswith(".")]
        if subdirs:
            if has_tev:
                for d in subdirs:
                    if "TEV" in d.name.upper():
                        return d
            if has_us:
                for d in subdirs:
                    name_u = d.name.upper()
                    if "UE" in name_u or "US" in name_u:
                        return d
            for d in subdirs:
                name_u = d.name.upper()
                if "IR" in name_u and not any(x in name_u for x in ["UE", "US", "TEV"]):
                    return d

    return default_template_dir


def generate_summary(csv_input):
    print("\n--- Starting Executive Summary Generation ---")

    if os.path.isabs(csv_input) and os.path.exists(csv_input):
        csv_file = Path(csv_input)
    else:
        csv_file = MAPPING_DIR / Path(csv_input).name

    if not csv_file.exists():
        print(f"Error: CSV file not found at {csv_file}")
        return

    csv_prefix = csv_file.stem
    # FLAT INDIVIDUAL directory: 07. OUTPUT FILE\INDIVIDUAL\
    indiv_dir = OUTPUT_DIR / "INDIVIDUAL"
    indiv_dir.mkdir(parents=True, exist_ok=True)
    output_file = indiv_dir / f"{csv_prefix}. Group 0.docx"

    print(f"Reading data from: {csv_file.name}")
    try:
        df = pd.read_csv(csv_file, na_filter=False)
        print(f"Successfully read {len(df)} rows.")
    except Exception as e:
        print(f"Error reading CSV: {e}")
        return

    swg_type = str(df.iloc[0].get('{{ swg.type }}', '')).strip()
    if swg_type in ['6', '6.0']:
        swg_type = 'RM6'
    swg_manuf = str(df.iloc[0].get('{{ swg.manufacturer }}', '')).strip()

    sub_tmpl_dir = resolve_template_dir_for_substation(csv_prefix, BASE_DIR, TEMPLATE_DIR)
    print(f"[{csv_prefix}] Template scope folder: {sub_tmpl_dir.name}")

    if "VCB" in swg_type.upper():
        template_doc = sub_tmpl_dir / "00. EXECUTIVE SUMMARY - VCB.docx"
    elif "H-POLE" in swg_type.upper() or "POLE" in swg_type.upper():
        template_doc = sub_tmpl_dir / "00. EXECUTIVE SUMMARY - H-POLE.docx"
    else:
        template_doc = sub_tmpl_dir / "00. EXECUTIVE SUMMARY - RMU.docx"

    if not template_doc.exists():
        if (TEMPLATE_DIR / template_doc.name).exists():
            template_doc = TEMPLATE_DIR / template_doc.name
        else:
            print(f"Error: Template not found at {template_doc}")
            return

    print(f"Using template: {template_doc.name}")
    try:
        doc = Document(template_doc)
    except Exception as e:
        print(f"Error loading template: {e}")
        return

    for section in doc.sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(1.125)
        section.right_margin = Cm(1.125)

    for t in doc.tables:
        t._element.getparent().remove(t._element)

    has_ir, has_us, has_tev = resolve_scopes_for_substation(csv_prefix, BASE_DIR)
    print(f"Scope detected for {csv_prefix}: IR={has_ir}, US={has_us}, TEV={has_tev}")

    headers = ["NO.", "EQUIPMENT", "DEFECT AREA"]
    if has_ir:
        headers.extend(["IR (Abs.T)", "IR (ΔT)"])
    if has_us:
        headers.append("US (dB)")
    if has_tev:
        headers.append("TEV (dB)")
    headers.append("DEFECT")

    total_cols = len(headers)
    table = doc.add_table(rows=1, cols=total_cols)
    set_table_borders(table)
    try:
        table.style = 'Table Grid'
    except KeyError:
        pass

    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        set_cell_shading(cell, "D9D9D9")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if p.runs:
            p.runs[0].font.bold = True

    processed_combinations = set()
    item_sequence = 0
    last_category = None
    last_item_sequence = None
    last_equip_desc = None
    last_equip_cell = None
    last_no_cell = None

    # Load saved thermal coordinates & values if present (BETA)
    coords_db = {}
    try:
        from thermal_tagger_beta import ThermalCoordinatesManager
        coords_db = ThermalCoordinatesManager.load_coordinates(str(BASE_DIR), csv_prefix)
    except Exception:
        pass

    for _, row in df.iterrows():
        type_val = str(row.get('Type', '')).strip()
        if not type_val or type_val.lower() in ['nan', 'none']:
            continue

        fn = str(row.get('Image File Name', '')).strip()
        folder = str(row.get('Folder', '')).strip()
        if not fn or fn.lower() in ['nan', 'none', '0', '0.0']:
            continue

        image_exists = False
        if folder and os.path.exists(os.path.join(folder, fn)):
            image_exists = True
        else:
            rm_root = RM_DIR
            local_base = resolve_rm_subfolder_for_report(rm_root, csv_prefix)
            if local_base:
                if (local_base / "IR" / fn).exists() or (local_base / fn).exists():
                    image_exists = True

        if not image_exists:
            continue

        current_cat = get_category(type_val)

        area_main = str(row.get('{{ swg.area }}', '')).strip()
        if not area_main or area_main.lower() == 'nan':
            area_main = str(row.get('Attribute', '')).strip()

        area = area_main

        # Rule for INDKOM RMU SF6 TX Panels: CABLE COMPARTMENT -> FUSE COMPARTMENT
        is_indkom = 'INDKOM' in swg_manuf.upper()
        is_rmu_sf6 = any(x in swg_type.upper() for x in ['RMU', 'MRMU', 'SF6'])
        if 'panel' in type_val.lower() and is_indkom and is_rmu_sf6:
            p_name = str(row.get('{{ panel.name }}', '')).strip().upper()
            p_no = str(row.get('{{ panel.linknumber }}', '')).strip().upper()
            if 'TX' in p_name or 'TX' in p_no or 'TRANSFORMER' in p_name:
                if area == 'CABLE COMPARTMENT':
                    area = 'FUSE COMPARTMENT'

        if type_val.startswith('TX') or 'POWER TX' in type_val or 'LOCALT TX' in type_val:
            side = str(row.get('{{ panel.name }}', '')).strip()
            if side in ['HV SIDE', 'LV SIDE']:
                prefix = 'HV ' if side == 'HV SIDE' else 'LV '
                if area_main in ['BUSHING', 'CABLE', 'CABLE SPLIT']:
                    area = f"{prefix}{area_main}"

        equip_desc = ""
        if "Overview" in type_val:
            equip_desc = f"{swg_type}, {swg_manuf}"
        elif 'POWER TX' in type_val or 'LOCALT TX' in type_val or type_val.startswith('TX'):
            tx_num = "".join(filter(str.isdigit, type_val)) or "1"
            manuf = str(row.get(f'{{{{ tx{tx_num}.manufacturer }}}}', '')).strip()
            r = str(row.get(f'{{{{ tx{tx_num}.rating }}}}', '')).strip()
            if manuf.lower() in ['0.0', '0', 'nan', 'none', '-']: manuf = ''
            if r.lower() in ['0.0', '0', 'nan', 'none', '-']: r = ''
            parts = [p for p in [manuf, r] if p]
            if parts:
                equip_desc = f"TX{tx_num} - {' '.join(parts)}"
            else:
                equip_desc = f"TX{tx_num}"
        elif 'panel' in type_val.lower():
            p_no = str(row.get('{{ panel.linknumber }}', '')).strip()
            p_name = str(row.get('{{ panel.name }}', '')).strip()
            if p_no and p_name and p_no != '-' and p_name != '-':
                equip_desc = f"PANEL {p_no}\n{p_name}"
            elif p_no and p_no != '-':
                equip_desc = f"PANEL {p_no}"
            else:
                equip_desc = p_name if p_name and p_name != '-' else type_val.upper()
        elif 'BATTERY' in type_val.upper():
            b_num = "".join(filter(str.isdigit, type_val)) or "1"
            equip_desc = f"BATTERY BANK {b_num}"
        elif 'FP' in type_val.upper() or 'FEEDER PILLAR' in type_val.upper() or 'LVDB' in type_val.upper():
            fp_digits = "".join(filter(str.isdigit, type_val)) or "1"
            label_col = f'{{{{ fp{fp_digits}.labelsource }}}}'
            custom_label = str(row.get(label_col, '')).strip()
            if custom_label and custom_label.lower() not in ['nan', 'none', '0', '0.0', 'xlpe', 'pilc', '-', '']:
                equip_desc = custom_label.upper()
            else:
                has_tx1 = bool(str(row.get('{{ tx1.manufacturer }}', '')).strip() not in ['', '-', 'nan', 'none', '0', '0.0'])
                has_tx2 = bool(str(row.get('{{ tx2.manufacturer }}', '')).strip() not in ['', '-', 'nan', 'none', '0', '0.0'])
                is_lvdb = 'LVDB' in type_val.upper()
                kind = 'LVDB' if is_lvdb else 'FP'
                tx_count = (1 if has_tx1 else 0) + (1 if has_tx2 else 0)
                if tx_count >= 2:
                    equip_desc = f"{kind} TX{fp_digits}"
                else:
                    equip_desc = f"{kind} {fp_digits}"
        else:
            equip_desc = type_val.upper()

        equip_desc = equip_desc.replace('nan', '').strip()

        combo = (equip_desc, area)
        if combo in processed_combinations:
            continue
        processed_combinations.add(combo)

        if current_cat != last_category:
            item_sequence += 1
            last_category = current_cat

        row_obj = table.add_row()
        cells = row_obj.cells

        cells[2].text = area

        # Check saved thermal coordinates for this page
        p_name_out = str(row.get('Output File Name', '')).strip()
        abs_t_val = '-'
        delta_t_val = '-'
        if p_name_out in coords_db:
            p_dat = coords_db[p_name_out]
            s1 = str(p_dat.get('sp1', {}).get('temp', '')).strip()
            a1 = str(p_dat.get('ar1', {}).get('temp_avg', '')).strip()
            dt = str(p_dat.get('delta_t', '')).strip()
            if s1: abs_t_val = f"{s1} °C" if "°" not in s1 else s1
            elif a1: abs_t_val = f"{a1} °C" if "°" not in a1 else a1
            if dt: delta_t_val = f"{dt} °C" if "°" not in dt else dt
        
        c_idx = 3
        if has_ir:
            cells[c_idx].text = abs_t_val
            cells[c_idx + 1].text = delta_t_val
            c_idx += 2
        if has_us:
            cells[c_idx].text = '-'
            c_idx += 1
        if has_tev:
            cells[c_idx].text = '-'
            c_idx += 1

        defect_col_idx = c_idx
        if "OVERVIEW" in area.upper():
            cells[defect_col_idx].text = '-'
        else:
            cells[defect_col_idx].text = ''
            set_cell_shading(cells[defect_col_idx], "00B050")

        if item_sequence == last_item_sequence and last_no_cell is not None and "BATTERY" not in current_cat:
            last_no_cell.merge(cells[0])
        else:
            cells[0].text = f"{item_sequence}."
            last_no_cell = cells[0]
            last_item_sequence = item_sequence

        if equip_desc == last_equip_desc and last_equip_cell is not None:
            last_equip_cell.merge(cells[1])
        else:
            cells[1].text = equip_desc
            last_equip_cell = cells[1]
            last_equip_desc = equip_desc

    # Column Widths based on total columns (A4 printable width ~18.75 cm)
    if total_cols == 8:  # IR (2) + US (1) + TEV (1) + Base (4)
        col_widths = [Cm(1.0), Cm(4.0), Cm(4.25), Cm(1.8), Cm(1.8), Cm(1.8), Cm(1.8), Cm(2.3)]
    elif total_cols == 7:  # IR (2) + US/TEV (1) + Base (4)
        col_widths = [Cm(1.0), Cm(4.5), Cm(4.75), Cm(2.0), Cm(2.0), Cm(2.0), Cm(2.5)]
    elif total_cols == 6:  # IR (2) + Base (4)
        col_widths = [Cm(1.0), Cm(5.0), Cm(5.25), Cm(2.25), Cm(2.25), Cm(3.0)]
    else:
        w_rest = max(1.5, (18.75 - 1.0 - 4.5 - 4.5 - 2.5) / max(1, total_cols - 4))
        col_widths = [Cm(1.0), Cm(4.5), Cm(4.5)] + [Cm(w_rest)] * (total_cols - 4) + [Cm(2.5)]

    for row in table.rows:
        row.height = Cm(0.65)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for i, cell in enumerate(row.cells):
            if i < len(col_widths):
                cell.width = col_widths[i]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.name = 'Calibri'
                    run.font.size = Pt(9.0)

    for p in doc.paragraphs:
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)

    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                for p in c.paragraphs:
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)

    doc.save(output_file)
    print(f"Summary generated successfully at: {output_file}")


if __name__ == "__main__":
    if len(sys.argv) > 1:
        generate_summary(sys.argv[1])
    else:
        print("Usage: python generate_summary.py <csv_file_or_path>")