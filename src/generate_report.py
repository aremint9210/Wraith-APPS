import os
import sys
import re
import json
import pandas as pd
import subprocess
from pathlib import Path
from datetime import datetime
from docx import Document
from docx.shared import Inches, Cm, Pt
from docxtpl import DocxTemplate, InlineImage, RichText

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")


def format_time_string(val):
    val = str(val).strip().upper().replace('.', ':')
    if re.match(r'^\d{4}$', val):
        val = val[:2] + ":" + val[2:]
    for fmt in ["%I:%M %p", "%I:%M%p", "%H:%M", "%H%M", "%H:%M:%S"]:
        try:
            return datetime.strptime(val, fmt).strftime("%I:%M %p")
        except ValueError:
            pass
    return val.upper()


def format_date_string(val):
    val = str(val).strip()
    if not val:
        return val
    if re.match(r'^\d{2}/\d{2}/\d{4}$', val):
        return val
    match = re.match(r'^(\d{4})-(\d{2})-(\d{2})$', val)
    if match:
        year, month, day = match.groups()
        return f"{day}/{month}/{year}"
    for fmt in ["%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"]:
        try:
            return datetime.strptime(val, fmt).strftime("%d/%m/%Y")
        except ValueError:
            pass
    return val


def page_sort_key(name):
    numbers = re.findall(r'\d+', str(name))
    return int(numbers[0]) if numbers else 999999


def format_heater_amp_val(val):
    if val is None:
        return "-"
    s = str(val).strip()
    if not s or s in ['-', 'N/A', '', '.', '/'] or s.lower() in ['n/a', 'nan', 'none', 'null']:
        return "-"
    m = re.search(r'[-+]?\d*\.?\d+', s)
    if m:
        try:
            f = float(m.group(0))
            return f"{f:.2f}A"
        except ValueError:
            pass
    return s


# ── Auto-Detect Workspace Directory Dynamically ─────────────────────────────
SCRIPT_PATH = Path(__file__).resolve()
SCRIPT_DIR = SCRIPT_PATH.parent

def find_active_workspace():
    env_base = os.environ.get("BASE_DIR")
    if env_base and (Path(env_base) / "06. EXCELL MAPPING").exists():
        return Path(env_base)

    if (Path.cwd() / "06. EXCELL MAPPING").exists():
        return Path.cwd()

    parent_dir = SCRIPT_PATH.parents[1] if len(SCRIPT_PATH.parents) > 1 else SCRIPT_PATH.parent
    if (parent_dir / "06. EXCELL MAPPING").exists():
        return parent_dir

    if parent_dir.exists():
        for p in parent_dir.iterdir():
            if p.is_dir() and (p / "06. EXCELL MAPPING").exists():
                return p

    default_base = Path(r"C:\Users\aremi\Desktop\REPORT CBM\PO 42234207 - JOHOR - JBU - 400 PE IR US TEV")
    if (default_base / "06. EXCELL MAPPING").exists():
        return default_base

    return parent_dir


BASE_DIR = Path(os.environ.get("BASE_DIR", find_active_workspace()))
TEMPLATE_DIR = Path(os.environ.get("TEMPLATE_DIR", BASE_DIR / "05. TEMPLATE"))
MAPPING_DIR = Path(os.environ.get("MAPPING_DIR", BASE_DIR / "06. EXCELL MAPPING"))
OUTPUT_DIR = Path(os.environ.get("OUTPUT_DIR", BASE_DIR / "07. OUTPUT FILE"))
OUTPUT_INDIVIDUAL = OUTPUT_DIR / "INDIVIDUAL"
RM_DIR = Path(os.environ.get("RM_DIR", BASE_DIR / "03. RAW MATERIAL"))

HELPER_SCRIPT_PATH = SCRIPT_DIR / "flir_word_report_helper.ps1"
MERGE_SCRIPT_PATH = SCRIPT_DIR / "dynamic_merge_tool.ps1"
POWERSHELL_32_PATH = Path(r"C:\Windows\SysWOW64\WindowsPowerShell\v1.0\powershell.exe")

OUTPUT_INDIVIDUAL.mkdir(parents=True, exist_ok=True)


def resolve_scopes_for_substation(csv_prefix, base_dir=None):
    env_ir = os.environ.get("IR_RANGE", "")
    env_us = os.environ.get("US_RANGE", "")
    env_tev = os.environ.get("TEV_RANGE", "")
    
    config_ir, config_us, config_tev = "", "", ""
    if base_dir:
        scope_cfg_path = Path(base_dir) / "scope_config.json"
        if scope_cfg_path.exists():
            try:
                import json
                with open(scope_cfg_path, "r", encoding="utf-8") as f:
                    cfg = json.load(f)
                    config_ir = cfg.get("ir_range", "")
                    config_us = cfg.get("us_range", "")
                    config_tev = cfg.get("tev_range", "")
            except Exception:
                pass

    ir_range = env_ir or config_ir
    us_range = env_us or config_us
    tev_range = env_tev or config_tev

    def _is_in_range(val, r_str):
        if not r_str or not str(r_str).strip():
            return False
        r_str = str(r_str).strip().upper()
        if r_str in ['ALL', '*', 'YES', 'TRUE', '1']:
            return True
        if r_str in ['NONE', 'NO', 'FALSE', '0', '-']:
            return False
        import re
        m = re.search(r'(\d+)', str(val))
        if not m:
            return False
        sub_n = int(m.group(1))
        for part in re.split(r'[,;\s]+', r_str):
            part = part.strip()
            if not part: continue
            if '-' in part:
                bounds = part.split('-')
                if len(bounds) == 2:
                    try:
                        if int(bounds[0].strip()) <= sub_n <= int(bounds[1].strip()):
                            return True
                    except ValueError: pass
            else:
                try:
                    if int(part) == sub_n:
                        return True
                except ValueError: pass
        return False

    has_ir = True if not ir_range else _is_in_range(csv_prefix, ir_range)
    has_us = _is_in_range(csv_prefix, us_range)
    has_tev = _is_in_range(csv_prefix, tev_range)

    if not ir_range and not us_range and not tev_range and base_dir:
        base_name = Path(base_dir).name.upper()
        if "TEV" in base_name: has_tev = True
        if "US" in base_name or "UE" in base_name: has_us = True
        has_ir = True

    return has_ir, has_us, has_tev


def resolve_template_dir_for_substation(csv_prefix, base_dir=None, default_template_dir=None):
    if default_template_dir is None:
        default_template_dir = TEMPLATE_DIR
    default_template_dir = Path(default_template_dir)

    has_ir, has_us, has_tev = resolve_scopes_for_substation(csv_prefix, base_dir)

    t_ir = os.environ.get("TEMPLATE_IR_DIR", "")
    t_us = os.environ.get("TEMPLATE_IR_UE_DIR", "")
    t_tev = os.environ.get("TEMPLATE_IR_UE_TEV_DIR", "")

    if base_dir:
        scope_cfg_path = Path(base_dir) / "scope_config.json"
        if scope_cfg_path.exists():
            try:
                import json
                with open(scope_cfg_path, "r", encoding="utf-8") as f:
                    sc = json.load(f)
                    t_ir = t_ir or sc.get("template_ir", "")
                    t_us = t_us or sc.get("template_ir_ue", "")
                    t_tev = t_tev or sc.get("template_ir_ue_tev", "")
            except Exception:
                pass

    if has_tev and t_tev and Path(t_tev).exists():
        return Path(t_tev)
    if has_us and t_us and Path(t_us).exists():
        return Path(t_us)
    if t_ir and Path(t_ir).exists():
        return Path(t_ir)

    if default_template_dir.exists():
        subdirs = [d for d in default_template_dir.iterdir() if d.is_dir() and not d.name.startswith(".")]
        if subdirs:
            if has_tev:
                for d in subdirs:
                    if "TEV" in d.name.upper():
                        return d
            if has_us:
                for d in subdirs:
                    name_u = d.name.upper()
                    if "UE" in name_u or "US" in name_u:
                        return d
            for d in subdirs:
                name_u = d.name.upper()
                if "IR" in name_u and not any(x in name_u for x in ["UE", "US", "TEV"]):
                    return d

    return default_template_dir

def list_csv_files():
    if not MAPPING_DIR.exists():
        return []
    return sorted(list(MAPPING_DIR.glob("*.csv")))


def resolve_rm_subfolder_for_report(rm_dir, csv_prefix):
    if not rm_dir or not Path(rm_dir).exists():
        return None
    val_str = str(csv_prefix).strip()
    m = re.match(r'^0*(\d+)([a-zA-Z]*)', val_str)
    if m:
        clean_num = m.group(1) + m.group(2).upper()
        base_digits = m.group(1)
        suffix = m.group(2).upper()
    else:
        clean_num = val_str
        base_digits = val_str
        suffix = ''

    try:
        dir_entries = [d for d in Path(rm_dir).iterdir() if d.is_dir()]
    except Exception:
        return None

    # 1. Exact match
    for d in dir_entries:
        if d.name.lower() == val_str.lower() or d.name.lower() == clean_num.lower():
            return d

    # 2. Strict match with suffix
    if suffix:
        for d in dir_entries:
            m_d = re.match(r'^0*(\d+[a-zA-Z]?)[\.\s\-_]', d.name)
            if m_d and m_d.group(1).upper() == clean_num:
                return d

    # 3. Numeric prefix match
    for d in dir_entries:
        m_d = re.match(r'^0*(\d+[a-zA-Z]?)[\.\s\-_]', d.name)
        if m_d and m_d.group(1).upper() == clean_num:
            return d

    # 4. Fallback startswith
    for d in dir_entries:
        if (d.name.startswith(f"{clean_num}.") or d.name.startswith(f"{clean_num} ") or 
            d.name.startswith(f"{val_str}.") or d.name.startswith(f"{val_str} ") or
            d.name.startswith(f"{base_digits}.") or d.name.startswith(f"{base_digits} ")):
            return d

    return None


def find_local_image(base_dir, filename):
    if not base_dir or not filename:
        return None
    base_path = Path(base_dir)
    options = [
        base_path / filename,
        base_path / "IR" / filename,
        base_path / "DIGITAL" / filename,
        base_path / "DG" / filename
    ]
    return next((p for p in options if p.exists()), None)


def build_context_from_row(df_columns, row_visual):
    context = {}
    for col in df_columns:
        raw_key = str(col).strip()
        val = str(row_visual[col].iloc[0]).strip() if hasattr(row_visual[col], 'iloc') else str(row_visual[col]).strip()
        if val.lower() in ['nan', 'none', 'null', '', '0', '0.0', '0.', '-', 'n/a', '.', '/']:
            if raw_key.strip() == '{{ panel.db_level }}':
                val = "0dB"
            elif raw_key.strip() == '{{ panel.us.char }}':
                val = "Normal"
            elif 'heateramp' in raw_key.lower() and val.lower() in ['0', '0.0', '0.']:
                val = "0.00A"
            else:
                val = "-"

        # Clean trailing .0 from serial numbers or float representation of integers
        if val != "-" and val.endswith('.0') and 'heateramp' not in raw_key.lower():
            try:
                f_val = float(val)
                if f_val.is_integer():
                    val = str(int(f_val))
            except ValueError:
                pass

        if 'time' in raw_key.lower() and val != "-":
            val = format_time_string(val)
        elif 'date' in raw_key.lower() and val != "-":
            val = format_date_string(val)
        elif 'humidity' in raw_key.lower() and val != "-":
            if '%' not in val: val = f"{val}%"
        elif 'ambient' in raw_key.lower() and val != "-":
            if '°' not in val: val = f"{val} °C"
        elif 'heateramp' in raw_key.lower() and val != "-":
            val = format_heater_amp_val(val)
        elif raw_key.strip() == '{{ panel.db_level }}' and (not val or val == "-"):
            val = "0dB"
        if str(val).strip().upper() == "BERKAT INSAF":
            rt = RichText()
            rt.add("BERKAT INSAF", size=18)  # size 18 half-points = 9pt font
            val = rt

        context[raw_key] = val

        # Clean key without braces for docxtpl nested jinja context
        clean_key = raw_key.replace('{', '').replace('}', '').strip()
        context[clean_key] = val

        if '.' in clean_key:
            parts = clean_key.split('.')
            parent = parts[0]
            child = parts[1]
            if parent not in context or not isinstance(context[parent], dict):
                context[parent] = {}
            context[parent][child] = val

    if 'panel' in context and isinstance(context['panel'], dict):
        if 'serialnumber' not in context['panel'] or not context['panel']['serialnumber']:
            swg_sn = context.get('swg', {}).get('serialnumber', '') if isinstance(context.get('swg'), dict) else ''
            context['panel']['serialnumber'] = swg_sn

    # Rules for Switchgear / Panels
    swg_type = str(context.get('{{ swg.type }}', '') or context.get('swg', {}).get('type', '')).upper()
    swg_manuf = str(context.get('{{ swg.manufacturer }}', '') or context.get('swg', {}).get('manufacturer', '')).upper()
    type_val = str(context.get('Type', '')).lower()
    p_name = str(context.get('{{ panel.name }}', '') or context.get('panel', {}).get('name', '')).upper()
    p_no = str(context.get('{{ panel.linknumber }}', '') or context.get('panel', {}).get('linknumber', '')).upper()

    # Rule: VCB + TRANSITION PANEL / TRANSITION -> cable type '-', link number '-'
    is_vcb = "VCB" in swg_type
    is_transition = "TRANSITION" in p_name
    is_bus_section = "BUS SECTION" in p_name

    if is_vcb and is_transition:
        context['{{ panel.cabletype }}'] = '-'
        context['{{ panel.linknumber }}'] = '-'
        context['panel.cabletype'] = '-'
        context['panel.linknumber'] = '-'
        if 'panel' in context and isinstance(context['panel'], dict):
            context['panel']['cabletype'] = '-'
            context['panel']['linknumber'] = '-'
    elif is_bus_section:
        context['{{ panel.cabletype }}'] = '-'
        context['panel.cabletype'] = '-'
        if 'panel' in context and isinstance(context['panel'], dict):
            context['panel']['cabletype'] = '-'

    # Rule for INDKOM RMU SF6 TX Panels: CABLE COMPARTMENT -> FUSE COMPARTMENT
    is_indkom = 'INDKOM' in swg_manuf
    is_rmu_sf6 = any(x in swg_type for x in ['RMU', 'MRMU', 'SF6'])
    is_tx_panel = 'TX' in p_name or 'TX' in p_no or 'TRANSFORMER' in p_name or 'tx' in type_val

    if 'panel' in type_val and is_indkom and is_rmu_sf6 and is_tx_panel:
        area_val = str(context.get('{{ swg.area }}', '') or context.get('swg', {}).get('area', ''))
        if area_val == 'CABLE COMPARTMENT':
            context['{{ swg.area }}'] = 'FUSE COMPARTMENT'
            context['swg.area'] = 'FUSE COMPARTMENT'
            if 'swg' in context and isinstance(context['swg'], dict):
                context['swg']['area'] = 'FUSE COMPARTMENT'

    return context


def generate_executive_summary(df, csv_prefix):
    print("\n[STEP 1] Generating Executive Summary (page0)...")
    try:
        from generate_summary import generate_summary
        csv_file = MAPPING_DIR / f"{csv_prefix}.csv"
        generate_summary(str(csv_file))
    except Exception as e:
        print(f" Error generating summary: {e}")


def process_csv(csv_path, target_pages_input=None, no_merge=False):
    csv_path = Path(csv_path)
    print("=" * 60)
    print(f" Processing Substation CSV: {csv_path.name}")
    print("=" * 60)

    try:
        df = pd.read_csv(csv_path, encoding='utf-8')
    except UnicodeDecodeError:
        df = pd.read_csv(csv_path, encoding='latin1')

    print(f"Successfully read {len(df)} rows.")

    csv_prefix = csv_path.name.split('.')[0].strip()
    available_pages = sorted([p for p in df['Output File Name'].dropna().unique() if str(p).startswith('page')], key=page_sort_key)

    # Always ensure page0 is included in available_pages
    if "page0" not in available_pages:
        available_pages.insert(0, "page0")

    print(f"Total pages listed in CSV: {len(available_pages)}")

    if target_pages_input:
        user_input = str(target_pages_input).strip()
    else:
        user_input = "all"

    if user_input.lower() == 'all':
        target_pages = available_pages
    else:
        target_pages = [p.strip() for p in user_input.replace(',', ' ').split()]
        target_pages = [p for p in target_pages if p in available_pages]
        if not target_pages:
            print("No valid pages selected.")
            return

    # ── STEP 1: Executive Summary ─────────────────────────────────────────────
    if "page0" in target_pages or user_input.lower() == 'all':
        os.environ["BASE_DIR"] = str(BASE_DIR)
        os.environ["MAPPING_DIR"] = str(MAPPING_DIR)
        os.environ["TEMPLATE_DIR"] = str(TEMPLATE_DIR)
        os.environ["OUTPUT_DIR"] = str(OUTPUT_DIR)
        generate_executive_summary(df, csv_prefix)

    df_selected = df[df['Output File Name'].isin(target_pages)]

    # ── STEP 2: Pure Python Text & Visual Image Replacement ───────────────────
    print("\n[STEP 2] Replacing text and visual images page-by-page...")
    pages_group = df_selected.groupby('Output File Name')
    sorted_page_names = sorted(pages_group.groups.keys(), key=page_sort_key)

    valid_pages = []
    skipped_pages = []

    # Load saved thermal coordinates & values if present (BETA)
    coords_db = {}
    try:
        from thermal_tagger_beta import ThermalCoordinatesManager
        coords_db = ThermalCoordinatesManager.load_coordinates(str(BASE_DIR), csv_prefix)
    except Exception:
        pass

    for output_name in sorted_page_names:
        if output_name == "page0":
            continue

        group_data = pages_group.get_group(output_name)
        row_visual = group_data[group_data['Attribute'].str.contains('Visual', case=False, na=False)]
        row_ir = group_data[group_data['Attribute'].str.contains('IR', case=False, na=False)]

        if row_visual.empty or row_ir.empty:
            skipped_pages.append((output_name, "Missing Visual or IR row in CSV"))
            continue

        template_name = ""
        try:
            template_name = str(row_visual.iloc[0].iloc[3]).strip()
        except Exception:
            pass

        if not template_name or template_name.lower() in ['nan', 'none']:
            template_name = str(row_visual['Word Template'].iloc[0]).strip()

        if not template_name.lower().endswith('.docx'):
            template_name += '.docx'

        sub_tmpl_dir = resolve_template_dir_for_substation(csv_prefix, BASE_DIR, TEMPLATE_DIR)
        template_path = sub_tmpl_dir / Path(template_name).name
        if not template_path.exists():
            if (TEMPLATE_DIR / Path(template_name).name).exists():
                template_path = TEMPLATE_DIR / Path(template_name).name
            else:
                skipped_pages.append((output_name, f"Template not found: {template_name} (checked {sub_tmpl_dir})"))
                continue

        visual_filename = str(row_visual['Image File Name'].iloc[0]).strip()
        ir_filename = str(row_ir['Image File Name'].iloc[0]).strip()

        csv_folder = Path(str(row_visual['Folder'].iloc[0])) if 'Folder' in row_visual.columns else None
        visual_path = find_local_image(csv_folder, visual_filename) if csv_folder else None
        ir_path = find_local_image(csv_folder, ir_filename) if csv_folder else None

        if not visual_path or not ir_path:
            local_base = resolve_rm_subfolder_for_report(RM_DIR, csv_prefix)
            if local_base:
                if not visual_path:
                    visual_path = find_local_image(local_base, visual_filename)
                if not ir_path:
                    ir_path = find_local_image(local_base, ir_filename)

        if not visual_path or not visual_path.exists():
            skipped_pages.append((output_name, f"Missing Visual image: {visual_filename}"))
            continue
        if not ir_path or not ir_path.exists():
            skipped_pages.append((output_name, f"Missing IR image: {ir_filename}"))
            continue

        output_path = OUTPUT_INDIVIDUAL / f"{csv_prefix}. {output_name}.docx"

        try:
            doc = DocxTemplate(template_path)
            context = build_context_from_row(df.columns, row_visual)
            context['visual_image'] = InlineImage(doc, str(visual_path), width=Cm(7.5), height=Cm(5.5))
            context['ir_image'] = InlineImage(doc, str(ir_path), width=Cm(7.5), height=Cm(5.5))
            doc.render(context)

            # If saved thermal measurements exist for this page (BETA), populate table cells
            if output_name in coords_db:
                p_coords = coords_db[output_name]
                sp1_val = str(p_coords.get('sp1', {}).get('temp', '')).strip()
                sp2_val = str(p_coords.get('sp2', {}).get('temp', '')).strip()
                ar1_val = str(p_coords.get('ar1', {}).get('temp_avg', '')).strip()
                dt_val = str(p_coords.get('delta_t', '')).strip()

                for t in doc.tables:
                    for row in t.rows:
                        for idx_c in range(len(row.cells) - 1):
                            c_text = row.cells[idx_c].text.strip()
                            if 'Spot Temp.1' in c_text and sp1_val:
                                row.cells[idx_c + 1].text = f"{sp1_val} °C" if "°" not in sp1_val else sp1_val
                            elif 'Spot Temp. 2' in c_text and sp2_val:
                                row.cells[idx_c + 1].text = f"{sp2_val} °C" if "°" not in sp2_val else sp2_val
                            elif 'Average Temp.' in c_text and ar1_val:
                                row.cells[idx_c + 1].text = f"{ar1_val} °C" if "°" not in ar1_val else ar1_val
                            elif 'Delta T' in c_text and dt_val:
                                row.cells[idx_c + 1].text = f"{dt_val} °C" if "°" not in dt_val else dt_val

            for p in doc.paragraphs:
                for run in p.runs:
                    txt = run.text.upper()
                    if "HERMETICALLY SEAL" in txt or "CONSERVATOR TANK" in txt:
                        run.font.size = Pt(7)
                    elif "BERKAT INSAF" in txt:
                        run.font.size = Pt(9)

            for t in doc.tables:
                for row in t.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            for run in p.runs:
                                txt = run.text.upper()
                                if "HERMETICALLY SEAL" in txt or "CONSERVATOR TANK" in txt:
                                    run.font.size = Pt(7)
                                elif "BERKAT INSAF" in txt:
                                    run.font.size = Pt(9)

            doc.save(output_path)
            print(f" Page {output_name} text & visual rendered: {output_path.name}")

            valid_pages.append({
                'doc_path': str(output_path),
                'ir_path': str(ir_path),
                'visual_path': str(visual_path),
                'template_path': str(template_path),
            })
        except Exception as render_err:
            skipped_pages.append((output_name, f"Render error: {render_err}"))

    if skipped_pages:
        print(f"\nSkipped {len(skipped_pages)} pages:")
        for name, reason in skipped_pages:
            print(f"  - {name}: {reason}")

    print(f"\nTotal valid pages saved for FLIR pass: {len(valid_pages)}")

    # ── STEP 3: FLIR IR Injection (Word opens visibly, add-in loads, injects IR) ──
    if valid_pages:
        print(f"\n[STEP 3] Running FLIR SDK Engine for interactive IR controls...")
        batch_tasks = []
        for page in valid_pages:
            batch_tasks.append({
                "template_path": page['doc_path'],   # inject into rendered docx
                "output_path":   page['doc_path'],   # save back to same file
                "pages": [
                    {
                        "template_page_no": 1,
                        "ir_path":     page['ir_path'],
                        "visual_path": page['visual_path'],
                        "field_overrides": {}
                    }
                ]
            })

        batch_json_path = OUTPUT_INDIVIDUAL / f"temp_batch_{csv_prefix}.json"
        with open(batch_json_path, 'w', encoding='utf-8') as f:
            json.dump(batch_tasks, f, indent=2)

        try:
            cmd = [
                str(POWERSHELL_32_PATH) if POWERSHELL_32_PATH.exists() else "powershell.exe",
                "-NoProfile", "-ExecutionPolicy", "Bypass",
                "-File", str(HELPER_SCRIPT_PATH),
                "-ManifestPath", str(batch_json_path)
            ]
            process = subprocess.Popen(cmd, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True)
            for line in process.stdout:
                print(f"    {line.strip()}")
            process.wait()
        finally:
            if batch_json_path.exists():
                os.remove(batch_json_path)

    print("\nPage generation and FLIR IR injection finished.")

    if no_merge:
        return

    # ── STEP 4: Merging Reports into Final Group Document ────────────────────
    print("\n[STEP 4] Merging all Word pages into Group document...")
    if MERGE_SCRIPT_PATH.exists():
        try:
            subprocess.run([
                "powershell.exe",
                "-NoProfile",
                "-ExecutionPolicy", "Bypass",
                "-File", str(MERGE_SCRIPT_PATH),
                "-csvName", csv_path.name,
                "-targetGroupsRaw", "all"
            ], check=True)
        except subprocess.CalledProcessError as merge_err:
            print(f"Error during Word report merging: {merge_err}")


def main():
    if len(sys.argv) > 1:
        csv_path_arg = Path(sys.argv[1])
        if not csv_path_arg.is_absolute():
            csv_path_arg = MAPPING_DIR / csv_path_arg.name

        if not csv_path_arg.exists():
            print(f"Error: CSV file not found at {csv_path_arg}")
            sys.exit(1)

        no_merge = "--no-merge" in sys.argv

        if len(sys.argv) > 2 and sys.argv[2] != "--no-merge":
            target_pages_arg = sys.argv[2]
            process_csv(csv_path_arg, target_pages_arg, no_merge)
        else:
            process_csv(csv_path_arg, no_merge=no_merge)
        return

    csv_files = list_csv_files()
    if not csv_files:
        print(f"No CSV files found in {MAPPING_DIR}")
        return

    print("\n=== Select CSV File ===")
    for i, f in enumerate(csv_files):
        print(f"{i + 1}. {f.name}")
    print("0. Exit")

    choice = input("\nEnter choice (number): ").strip()
    if choice != '0':
        try:
            idx = int(choice) - 1
            if 0 <= idx < len(csv_files):
                process_csv(csv_files[idx])
            else:
                print("Invalid choice.")
        except ValueError:
            print("Please enter a number.")

    print("\nProcess finished.")


if __name__ == "__main__":
    main()